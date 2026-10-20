from __future__ import annotations

import io
import math
import re
import zipfile
from datetime import datetime
from typing import Any

from openpyxl.utils.cell import column_index_from_string, range_boundaries

from ..contracts import DocumentPage, ExtractedValue, ExtractionMethod
from ..ids import utcnow
from . import ExtractionOutput, parse_number_parts
from ._common import (
    ASSOCIATION, EMPTY, LABEL, NA, NUM, analyze_grid, build_value, cell_number_text, core_properties,
    derived_change_flag, finalize_values, flag, iso_date_text, issue, number_text, text_cell_kind,
)

ADAPTER_VERSION = "1"
EXCEL_ERRORS = frozenset({"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A",
                          "#GETTING_DATA", "#SPILL!", "#CALC!"})
_PLAIN_FORMATS = frozenset({"General", "0", "@", None})
_YEAR = "year"
_YEAR_HEADER = re.compile(r"\b(?:years?|yr|fy|cy|period|calendar)\b", re.I)
_CELL_SCOPED_CODES = frozenset({"MISSING_FORMULA_RESULT", "WORKSHEET_ERROR_VALUE"})


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value))


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if _is_number(value):
        return number_text(value)
    if isinstance(value, dict):
        return f"{value.get('seconds')}s"
    return iso_date_text(value)


def _kind(value: Any, fmt: str | None) -> str:
    if value is None or value == "":
        return EMPTY
    if isinstance(value, bool) or isinstance(value, dict):
        return LABEL
    if _is_number(value):
        if (float(value).is_integer() and 1990 <= value <= 2100 and fmt in _PLAIN_FORMATS):
            return _YEAR
        return NUM
    if isinstance(value, str) and value.strip() in EXCEL_ERRORS:
        return NA
    return text_cell_kind(value)[0]


def _resolve_years(texts: list[list[str]], kinds: list[list[str]], max_header_rows: int = 3) -> list[list[str]]:
    out = [list(row) for row in kinds]
    header_rows = 0
    for r, row in enumerate(kinds):
        years = [c for c, k in enumerate(row) if k == _YEAR]
        above = texts[r - 1] if r else []
        stacked = r == 0 or len(years) > 1 or any(not (above[c] if c < len(above) else "").strip() for c in years)
        if (r == header_rows and r < max_header_rows and NUM not in row and NA not in row
                and (LABEL in row or len(years) > 1) and stacked):
            out[r] = [LABEL if k == _YEAR else k for k in row]
            header_rows += 1
    for r in range(header_rows, len(kinds)):
        filled = [c for c, k in enumerate(kinds[r]) if k != EMPTY]
        for c, k in enumerate(kinds[r]):
            if k != _YEAR:
                continue
            header = " ".join(texts[i][c] for i in range(header_rows) if c < len(texts[i])).strip()
            row_label = c == filled[0] and len(filled) > 1 and (not header or bool(_YEAR_HEADER.search(header)))
            out[r][c] = LABEL if row_label else NUM
    return out


def _hidden_columns(data: dict[str, Any]) -> set[int]:
    cols: set[int] = set()
    for rng in data.get("hidden_column_ranges") or []:
        cols.update(range(int(rng["min"]), int(rng["max"]) + 1))
    return cols


def _merged_spans(data: dict[str, Any], first_row: int, first_col: int) -> list[tuple[int, int, int]]:
    spans = []
    for ref in data.get("merged_ranges") or []:
        try:
            min_col, min_row, max_col, _max_row = range_boundaries(ref)
        except (TypeError, ValueError):
            continue
        spans.append((min_row - first_row, min_col - first_col, max_col - min_col + 1))
    return spans


def worksheet_issues(data: dict[str, Any], *, page: int | None, part: str | None) -> list[dict[str, Any]]:
    """Formula cells without a saved result and error values."""
    out = []
    sheet = data.get("sheet_name")
    for f in data.get("formulas") or []:
        cached = f.get("cached_value")
        if cached is None:
            out.append(issue("FORMULA_RESULT_MISSING", f"{sheet}!{f.get('cell')} {f.get('formula')} has no saved result",
                             page=page, part=part))
        elif isinstance(cached, str) and cached.strip() in EXCEL_ERRORS:
            out.append(issue("WORKSHEET_FORMULA_ERROR", f"{sheet}!{f.get('cell')} {f.get('formula')} = {cached}",
                             page=page, part=part))
    for m in data.get("cell_metadata") or []:
        if m.get("data_type") == "e":
            out.append(issue("WORKSHEET_ERROR_VALUE", f"{sheet}!{m.get('cell')} holds an Excel error value",
                             page=page, part=part))
    return out


def worksheet_values(data: dict[str, Any], *, doc_id: str, method: ExtractionMethod, extractor: str,
                     extractor_version: str, base_locator: dict[str, Any], base_context: dict[str, Any],
                     base_flags: list[dict[str, Any]], run_id: str | None, created_at: datetime,
                     consumed: set[str] | None = None, source_record_id: str | None = None) -> list[ExtractedValue]:
    """Numeric cells of one worksheet (vendored ``_worksheet_data`` payload) -> values, like a table."""
    rows: list[list[Any]] = data.get("rows") or []
    if not rows:
        return []
    letters: list[str] = data["columns"]
    first_row = int(data.get("first_excel_row") or 1)
    first_col = column_index_from_string(letters[0])
    meta = {m["cell"]: m for m in data.get("cell_metadata") or []}
    formulas = {f["cell"]: f for f in data.get("formulas") or []}
    hidden_rows = set(data.get("hidden_rows") or [])
    hidden_cols = _hidden_columns(data)
    consumed = consumed or set()
    spans = _merged_spans(data, first_row, first_col)
    sheet = data.get("sheet_name")
    upstream = [f for f in base_flags if f.get("code") not in _CELL_SCOPED_CODES]

    blocks: list[list[int]] = [[]]
    for i, row in enumerate(rows):
        if all(v is None or v == "" for v in row):
            if blocks[-1]:
                blocks.append([])
        else:
            blocks[-1].append(i)
    out: list[ExtractedValue] = []
    for block in (b for b in blocks if b):
        texts = [[_cell_text(v) for v in rows[i]] for i in block]
        kinds = []
        for i in block:
            excel_row = first_row + i
            kinds.append([_kind(v, (meta.get(f"{letters[j]}{excel_row}") or {}).get("number_format"))
                          for j, v in enumerate(rows[i])])
        kinds = _resolve_years(texts, kinds)
        block_spans = [(r - block[0], c, n) for r, c, n in spans if block[0] <= r <= block[-1]]
        layout = analyze_grid(texts, kinds, spans=block_spans)
        for bi, i in enumerate(block):
            excel_row = first_row + i
            for j, value in enumerate(rows[i]):
                if kinds[bi][j] != NUM or (bi, j) in layout.header_cells:
                    continue
                coord = f"{letters[j]}{excel_row}"
                if coord in consumed:
                    continue
                m = meta.get(coord) or {}
                col_index = first_col + j
                mark_flags: list[dict[str, Any]] = []
                if _is_number(value):
                    raw_number, number_source = float(value), "stored"
                    raw_text = m.get("raw_xml_value") or (formulas.get(coord) or {}).get("cached_raw") or number_text(value)
                else:
                    parts = parse_number_parts(str(value))
                    raw_text, mark_flags = cell_number_text(str(value))
                    raw_number, number_source = parts.number, "text"
                hidden = excel_row in hidden_rows or col_index in hidden_cols
                cell_flags = list(upstream) + mark_flags
                if hidden:
                    cell_flags.append(flag("HIDDEN_CELL", ASSOCIATION, f"{sheet}!{coord} is in a hidden row or column."))
                col_header = layout.col_headers[j] if j < len(layout.col_headers) else None
                row_label = layout.row_labels[bi]
                cell_flags += derived_change_flag(col_header, row_label)
                left_labels = layout.labels_left_of(texts, kinds, bi, j)
                context = {
                    **base_context, "object_kind": "worksheet", "object_title": sheet,
                    "row_label": row_label, "col_header": col_header, "header_row": layout.header_row(),
                    "row_labels": left_labels if len(left_labels) > 1 else None,
                    "number_source": number_source,
                    "formula": (formulas.get(coord) or {}).get("formula"),
                    "hidden": True if hidden else base_context.get("hidden"),
                }
                locator = {**base_locator, "sheet": sheet, "cell": coord, "row": excel_row - 1, "col": col_index - 1}
                out.append(build_value(
                    doc_id=doc_id, method=method, extractor=extractor, extractor_version=extractor_version,
                    locator=locator, context=context, raw_text=str(raw_text), raw_number=raw_number,
                    number_format=m.get("number_format") or "General", flags=cell_flags, confidence=1.0,
                    run_id=run_id, source_record_id=source_record_id, created_at=created_at))
    return out


def extractor_versions() -> dict[str, str]:
    from ._vendor import pptx_extract

    return {"xlsx_adapter": f"kpi_evidence.extract.xlsx/{ADAPTER_VERSION}",
            "pptx_extract": pptx_extract.__version__}


def _sheet_text(data: dict[str, Any], limit: int = 20000) -> str:
    lines = []
    for row in data.get("rows") or []:
        cells = [_cell_text(v) for v in row]
        if any(cells):
            lines.append("\t".join(cells).rstrip())
    return "\n".join(lines)[:limit]


def extract_xlsx_values(doc_id: str, blob: bytes, file_name: str, run_id: str | None = None) -> ExtractionOutput:
    """Standalone workbook -> one value per numeric cell (header/row labels like tables) + one page per sheet."""
    from openpyxl import load_workbook

    from ._vendor import pptx_extract as px

    out = ExtractionOutput(extractor_versions=extractor_versions())
    if not blob.startswith(b"PK\x03\x04") or not zipfile.is_zipfile(io.BytesIO(blob)):
        out.issues.append(issue("NOT_AN_XLSX", f"{file_name} is not an OOXML workbook (ZIP) package", severity="error"))
        return out
    created_at = utcnow()
    version = f"openpyxl+kpe.{ADAPTER_VERSION}"
    config = px.ExtractionConfig()
    try:
        archive = px._checked_zip(blob, config)
    except Exception as exc:
        out.issues.append(issue("XLSX_PARSE_FAILED", f"{type(exc).__name__}: {exc}", severity="error"))
        return out
    with archive:
        if "xl/workbook.xml" not in archive.namelist():
            out.issues.append(issue("NOT_AN_XLSX", "package has no xl/workbook.xml", severity="error"))
            return out
        try:
            out.document_properties = core_properties(blob)
        except Exception as exc:
            out.issues.append(issue("CORE_PROPERTIES_UNREADABLE", f"{type(exc).__name__}: {exc}", severity="info"))
        try:
            workbook_root = px._xml(archive, "xl/workbook.xml")
            relationships = px._rels(archive, "xl/workbook.xml")
            formulas_wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=False, keep_links=False)
            values_wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True, keep_links=False)
        except Exception as exc:
            out.issues.append(issue("XLSX_PARSE_FAILED", f"{type(exc).__name__}: {exc}", severity="error"))
            return out
        try:
            sheets = workbook_root.findall("x:sheets/x:sheet", px.NS)
            out.document_properties["page_count"] = len(sheets)
            for index, sheet in enumerate(sheets, 1):
                name = sheet.get("name")
                state = sheet.get("state", "visible")
                rel = relationships.get(sheet.get(px.RID))
                hidden = state != "visible"
                page = DocumentPage(doc_id=doc_id, page_number=index, title=name, text="", notes=None, hidden=hidden,
                                    image_count=0, page_width=None, page_height=None, shapes=[], created_at=created_at)
                out.pages.append(page)
                if not rel or rel["type"] != "worksheet" or rel["external"]:
                    out.issues.append(issue("UNSUPPORTED_SHEET", f"sheet {name!r} is not a worksheet (chartsheet or "
                                            "external); not read", severity="info", page=index))
                    continue
                try:
                    data, _warnings = px._worksheet_data(archive, px._xml(archive, rel["target"]), formulas_wb[name],
                                                         values_wb[name], config)
                except Exception as exc:
                    out.issues.append(issue("SHEET_EXTRACTION_FAILED", f"{name!r}: {type(exc).__name__}: {exc}",
                                            page=index, part=rel["target"]))
                    continue
                data["sheet_name"] = name
                page.text = _sheet_text(data)
                base_flags = [flag("HIDDEN_SHEET", ASSOCIATION, f"Sheet {name!r} is hidden.")] if hidden else []
                out.values += worksheet_values(
                    data, doc_id=doc_id, method=ExtractionMethod.WORKBOOK, extractor="openpyxl",
                    extractor_version=version, base_locator={"page": index, "part": rel["target"]},
                    base_context={"page_title": name, "hidden": hidden}, base_flags=base_flags,
                    run_id=run_id, created_at=created_at)
                out.issues += worksheet_issues(data, page=index, part=rel["target"])
        finally:
            formulas_wb.close()
            values_wb.close()
    out.values = finalize_values(out.values, out.issues)
    return out
