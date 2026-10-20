from __future__ import annotations

import io
import json
import posixpath
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import unquote

from openpyxl.utils.cell import get_column_letter, range_boundaries

from ..contracts import DocumentPage, ExtractedValue, ExtractionMethod
from ..ids import sha256_bytes, utcnow
from . import ExtractionOutput
from ._common import (
    ASSOCIATION, EXTRACTION, MEANING, NUM, analyze_grid, bbox_list, build_value, cell_number_text, core_properties,
    derived_change_flag, finalize_values, find_labeled_numbers, flag, footnote_mark_flag, hidden_slide_flag,
    is_footnote, is_note_line, issue, number_text, serial_to_iso, text_cell_kind,
)
from ._vendor import pptx_extract as px
from .xlsx import worksheet_issues, worksheet_values

ADAPTER_VERSION = "1"
EXTRACTOR = "pptx_extract"
TEXT_CONFIDENCE = 0.7
LARGE_IMAGE_FRACTION = 0.15
LOGO_FRACTION = 0.03
_VENDOR_INFO = json.loads((Path(px.__file__).with_name("VENDOR.json")).read_text(encoding="utf-8"))
_SKIP_TEXT_PLACEHOLDERS = frozenset({"dt", "ftr", "sldNum"})
_POINT_SCOPED_CODES = frozenset({"NONNUMERIC_CHART_VALUE", "MISSING_CHART_POINTS"})
_FORMULA_AREA = re.compile(r"^(?:'(?P<q>(?:[^']|'')+)'|(?P<s>[^'!]+))!(?P<ref>\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?)$")
_MAX_FORMULA_CELLS = 250_000


def extractor_versions(report: dict[str, Any] | None = None) -> dict[str, str]:
    out = {"pptx_extract": px.__version__, "pptx_extract_commit": _VENDOR_INFO.get("commit", ""),
           "pptx_extract_sha256": _VENDOR_INFO.get("sha256", ""), "pptx_adapter": f"kpi_evidence.extract.pptx/{ADAPTER_VERSION}"}
    if report:
        out["pptx_extract_schema"] = str(report.get("schema_version", ""))
    return out


def sniff_problem(blob: bytes) -> str | None:
    """Why these bytes are not a PPTX/PPTM package, or None."""
    if blob[:4] == b"\xd0\xcf\x11\xe0":
        return ("OLE compound file: a legacy .ppt or an encrypted (sensitivity-labeled) package. Save as .pptx, "
                "or remove the encryption / export to PDF.")
    if blob[:4] == b"%PDF":
        return "this is a PDF, not a PowerPoint package"
    if blob[:4] != b"PK\x03\x04" or not zipfile.is_zipfile(io.BytesIO(blob)):
        return "not a ZIP/OOXML package"
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        if "ppt/presentation.xml" not in archive.namelist():
            return "ZIP package without ppt/presentation.xml (not a PowerPoint deck)"
    return None


def _slide_parts(archive: zipfile.ZipFile) -> dict[int, str]:
    root = px._xml(archive, "ppt/presentation.xml")
    rels_path = "ppt/_rels/presentation.xml.rels"
    targets: dict[str, str] = {}
    if rels_path in archive.namelist():
        for rel in px._xml(archive, rels_path):
            if rel.get("TargetMode") == "External":
                continue
            target = unquote(rel.get("Target", ""))
            resolved = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("ppt", target))
            if not resolved.startswith("..") and "\\" not in resolved:
                targets[rel.get("Id")] = resolved
    return {n: targets[s.get(px.RID)] for n, s in enumerate(root.findall("p:sldIdLst/p:sldId", px.NS), 1)
            if s.get(px.RID) in targets}


@dataclass
class _TextShape:
    shape_id: str | None
    name: str
    placeholder: str | None
    text: str
    bbox: list[float] | None
    position: str
    nearby: list[str]
    marks: frozenset[tuple[int, int]] = frozenset()


@dataclass
class _SlideScan:
    number: int
    part: str | None = None
    hidden: bool = False
    title: str | None = None
    text_shapes: list[_TextShape] = field(default_factory=list)
    shapes: list[dict[str, Any]] = field(default_factory=list)
    footnotes: list[str] = field(default_factory=list)
    error: str | None = None
    tables: dict[str | None, list[list[tuple[str, bool]]] | None] = field(default_factory=dict)


_A = "{" + px.NS["a"] + "}"
_MARK_TEXT = re.compile(r"[\d*†‡§]{1,3}(?:\s?,\s?[\d*†‡§]{1,3})*|[a-z]")
_NUMBER_END = re.compile(r"\d(?:\s?(?:bn|mm|mn|bps|pts?|pp)|[kmbx])?[%)]*$", re.I)


def _is_raised(props) -> bool:
    try:
        return props is not None and int(props.get("baseline") or 0) > 0
    except ValueError:
        return False


def _marked_text(node) -> tuple[str, frozenset[tuple[int, int]]]:
    if node is None:
        return "", frozenset()
    paragraphs: list[str] = []
    found: list[tuple[int, int]] = []
    for index, paragraph in enumerate(node.findall(".//a:p", px.NS)):
        chunks: list[str] = []
        raised = False
        for child in paragraph.iter():
            if child.tag in (_A + "r", _A + "fld"):
                raised = _is_raised(child.find("a:rPr", px.NS))
            elif child.tag == _A + "t":
                text = child.text or ""
                before = "".join(chunks)
                if raised and _MARK_TEXT.fullmatch(text.strip()) and _NUMBER_END.search(before):
                    found.append((index, len(before)))
                    chunks.append(text[len(text.rstrip()):])
                else:
                    chunks.append(text)
            elif child.tag == _A + "br":
                chunks.append("\n")
            elif child.tag == _A + "tab":
                chunks.append("\t")
        paragraphs.append("".join(chunks))
    full = "\n".join(paragraphs)
    text = full.strip()
    lead = len(full) - len(full.lstrip())
    starts, position = [], 0
    for paragraph_text in paragraphs:
        starts.append(position)
        position += len(paragraph_text) + 1
    marks = set()
    for index, offset in found:
        at = starts[index] + offset - lead
        if 0 <= at <= len(text):
            marks.add((text.count("\n", 0, at), at - (text.rfind("\n", 0, at) + 1)))
    return text, frozenset(marks)


def _table_cells(table) -> list[list[tuple[str, bool]]]:
    rows = []
    for tr in table.findall("a:tr", px.NS):
        cells = []
        for tc in tr.findall("a:tc", px.NS):
            text, marks = _marked_text(tc)
            cells.append((text, bool(marks)))
        rows.append(cells)
    return rows


def _shape_kind(shape, info: dict[str, Any]) -> str:
    tag = px._tag(shape)
    if tag == "pic":
        return "picture"
    if tag == "graphicFrame":
        if shape.find(".//a:tbl", px.NS) is not None:
            return "table"
        if any(px._tag(n) == "chart" for n in shape.iter()):
            return "chart"
        if any(px._tag(n) == "oleObj" for n in shape.iter()):
            return "ole_object"
        if shape.find(".//dgm:relIds", px.NS) is not None:
            return "diagram"
        return "frame"
    if info.get("placeholder_type") in {"title", "ctrTitle"}:
        return "title"
    return "text" if px._text(shape) else "shape"


def scan_slide(archive: zipfile.ZipFile, number: int, part: str, slide_size: dict[str, Any]) -> _SlideScan:
    """Per-shape text with ids, placeholder types and bounding boxes (uses the vendored geometry helpers)."""
    scan = _SlideScan(number=number, part=part)
    root = px._xml(archive, part)
    scan.hidden = root.get("show") in {"0", "false"}
    tree = root.find("p:cSld/p:spTree", px.NS)
    shapes = list(px._shapes(tree)) if tree is not None else []
    try:
        layouts = px._shape_layouts(tree, slide_size.get("width"), slide_size.get("height"))
    except (TypeError, ValueError, KeyError):
        layouts = {}
    footnotes: list[str] = []
    for shape in shapes:
        tag = px._tag(shape)
        if tag in {"cxnSp", "contentPart"}:
            continue
        info = px._shape_info(shape)
        layout = layouts.get(id(shape), {})
        box = bbox_list(layout.get("bounding_box"))
        scan.shapes.append({"shape_id": info["shape_id"], "kind": _shape_kind(shape, info),
                            "name": info["shape_name"], "bbox": box})
        table = shape.find(".//a:tbl", px.NS)
        if table is not None:
            key = info["shape_id"]
            scan.tables[key] = None if key in scan.tables else _table_cells(table)
        if tag != "sp":
            continue
        text, marks = _marked_text(shape)
        if not text:
            continue
        if scan.title is None and info["placeholder_type"] in {"title", "ctrTitle"}:
            scan.title = text
        footnotes += [line.strip() for line in text.split("\n") if is_footnote(line)]
        try:
            nearby = [c["text"] for c in px._nearby_labels(shape, shapes, layouts)]
        except (TypeError, ValueError, KeyError, ZeroDivisionError):
            nearby = []
        scan.text_shapes.append(_TextShape(info["shape_id"], info["shape_name"], info["placeholder_type"], text, box,
                                           layout.get("position_on_slide") or "unknown", nearby, marks))
    scan.footnotes = list(dict.fromkeys(footnotes))
    return scan


def _scan_all(blob: bytes, report: dict[str, Any], issues: list[dict[str, Any]]) -> dict[int, _SlideScan]:
    scans: dict[int, _SlideScan] = {}
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        try:
            parts = _slide_parts(archive)
        except Exception as exc:
            issues.append(issue("SLIDE_SCAN_FAILED", f"slide list unreadable: {type(exc).__name__}: {exc}"))
            return scans
        for number in range(1, int(report.get("slide_count") or 0) + 1):
            part = parts.get(number)
            if part is None or part not in archive.namelist():
                scans[number] = _SlideScan(number=number, error="slide part not found")
                continue
            try:
                scans[number] = scan_slide(archive, number, part, report.get("slide_size_emu") or {})
            except Exception as exc:
                scans[number] = _SlideScan(number=number, part=part, error=f"{type(exc).__name__}: {exc}")
                issues.append(issue("SLIDE_SCAN_FAILED", scans[number].error, page=number, part=part))
    return scans


def _split_areas(text: str) -> list[str]:
    areas, buf, quoted = [], [], False
    for ch in text:
        if ch == "'":
            quoted = not quoted
        if ch == "," and not quoted:
            areas.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
    areas.append("".join(buf))
    return [a.strip() for a in areas if a.strip()]


def formula_cells(formula: str | None) -> list[tuple[str, str]]:
    """``"Sheet1!$B$2:$B$4"`` -> ``[("Sheet1","B2"), ("Sheet1","B3"), ("Sheet1","B4")]``; [] if unparseable."""
    if not formula:
        return []
    text = formula.strip()
    if text.startswith("(") and text.endswith(")"):
        text = text[1:-1]
    cells: list[tuple[str, str]] = []
    for area in _split_areas(text):
        m = _FORMULA_AREA.match(area)
        if not m:
            return []
        sheet = m["q"].replace("''", "'") if m["q"] else m["s"]
        try:
            min_col, min_row, max_col, max_row = range_boundaries(m["ref"].replace("$", ""))
        except (TypeError, ValueError):
            return []
        if (max_row - min_row + 1) * (max_col - min_col + 1) + len(cells) > _MAX_FORMULA_CELLS:
            return []
        cells += [(sheet, f"{get_column_letter(c)}{r}") for r in range(min_row, max_row + 1)
                  for c in range(min_col, max_col + 1)]
    return cells


@dataclass
class _Sheet:
    record: dict[str, Any]
    part: str
    name: str
    cells: dict[str, Any]
    consumed: set[str] = field(default_factory=set)


def _sheet_cells(data: dict[str, Any]) -> dict[str, Any]:
    letters, first = data.get("columns") or [], int(data.get("first_excel_row") or 1)
    return {f"{letters[j]}{first + i}": v for i, row in enumerate(data.get("rows") or []) for j, v in enumerate(row)
            if j < len(letters)}


def _same_number(a: Any, b: Any) -> bool:
    numeric = (int, float)
    if isinstance(a, bool) or isinstance(b, bool) or not isinstance(a, numeric) or not isinstance(b, numeric):
        return False
    return abs(float(a) - float(b)) <= 1e-9 * max(1.0, abs(float(a)), abs(float(b)))


@dataclass
class _Run:
    doc_id: str
    run_id: str | None
    created_at: datetime
    version: str
    scans: dict[int, _SlideScan]
    slide_area: float | None
    include_hidden: bool = True
    groupings: dict[str, list[str | None]] = field(default_factory=dict)

    def value(self, method: ExtractionMethod, **kw: Any) -> ExtractedValue:
        return build_value(doc_id=self.doc_id, method=method, extractor=EXTRACTOR, extractor_version=self.version,
                           run_id=self.run_id, created_at=self.created_at, **kw)

    def footnotes(self, page: int | None) -> list[str]:
        scan = self.scans.get(page) if page else None
        return scan.footnotes if scan else []

    def table_cells(self, rec: dict[str, Any]) -> list[list[tuple[str, bool]]] | None:
        scan = self.scans.get(rec.get("slide_number")) if rec.get("slide_number") else None
        cells = scan.tables.get(rec.get("shape_id")) if scan else None
        rows = (rec.get("data") or {}).get("rows") or []
        if not cells or [len(r) for r in cells] != [len(r) for r in rows]:
            return None
        return cells if any(marked for row in cells for _t, marked in row) else None


def _chart_groupings(blob: bytes, parts: list[str]) -> dict[str, list[str | None]]:
    out: dict[str, list[str | None]] = {}
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        names = set(archive.namelist())
        for part in parts:
            if part not in names:
                continue
            try:
                plot = px._xml(archive, part).find("c:chart/c:plotArea", px.NS)
            except Exception:
                continue
            if plot is None:
                continue
            groupings: list[str | None] = []
            for group in plot:
                if px._tag(group).endswith("Chart"):
                    groupings += [px._value(group, "c:grouping")] * len(group.findall("c:ser", px.NS))
            out[part] = groupings
    return out


def _base_context(rec: dict[str, Any], run: _Run) -> dict[str, Any]:
    evidence = rec.get("description_evidence") or {}
    page = rec.get("slide_number")
    return {
        "page_title": rec.get("slide_title") or None,
        "nearby_text": [c.get("text") for c in evidence.get("nearby_text_candidates") or [] if c.get("text")],
        "footnotes": run.footnotes(page),
        "hidden": bool(rec.get("hidden_slide")),
        "position": rec.get("position_on_slide") or "unknown",
        "alt_text": evidence.get("alt_text"),
    }


def _base_flags(rec: dict[str, Any], skip: frozenset[str] = frozenset()) -> list[dict[str, Any]]:
    flags = [dict(f) for f in rec.get("uncertainty_reasons") or [] if f.get("code") not in skip]
    if rec.get("hidden_slide"):
        flags.append(hidden_slide_flag())
    return flags


def _shape_locator(rec: dict[str, Any]) -> dict[str, Any]:
    return {"page": rec.get("slide_number"), "part": rec.get("source_part"), "shape_id": rec.get("shape_id"),
            "shape_name": rec.get("shape_name") or None, "bbox": bbox_list(rec.get("bounding_box"))}


def _series_format(series: dict[str, Any], axes: list[dict[str, Any]], vec: str) -> tuple[str | None, str | None]:
    fmt = (series["sources"].get(vec) or {}).get("number_format")
    if fmt and fmt != "General":
        return fmt, "series"
    for axis in axes:
        if axis.get("axis_id") in series.get("axis_ids", []) and axis.get("axis_type") == "valAx":
            if axis.get("number_format") and axis["number_format"] != "General":
                return axis["number_format"], "axis"
    return None, None


def _is_percent_format(fmt: str | None) -> bool:
    return "%" in re.sub(r'"[^"]*"|\\.', "", fmt or "")


def _counts_in_percent_stack(series_list: list[dict[str, Any]], rows: list[dict[str, Any]],
                             groupings: list[str | None]) -> set[int]:
    stacked = {s["series_number"] for s in series_list
               if 0 < s["series_number"] <= len(groupings) and groupings[s["series_number"] - 1] == "percentStacked"}
    counts = any(isinstance(r.get("value"), (int, float)) and not isinstance(r.get("value"), bool)
                 and abs(r["value"]) > 1.5 for r in rows if r.get("series_number") in stacked)
    return stacked if counts else set()


def _value_axis_title(series: dict[str, Any], axes: list[dict[str, Any]]) -> str | None:
    return next((a.get("title") for a in axes if a.get("axis_id") in series.get("axis_ids", [])
                 and a.get("axis_type") == "valAx" and a.get("title")), None)


def _is_date_category(series: dict[str, Any], axes: list[dict[str, Any]]) -> bool:
    from openpyxl.styles.numbers import is_date_format

    fmt = (series["sources"].get("cat") or {}).get("number_format")
    if fmt and is_date_format(fmt):
        return True
    return any(a.get("axis_type") == "dateAx" and a.get("axis_id") in series.get("axis_ids", []) for a in axes)


def chart_values(rec: dict[str, Any], run: _Run, sheets: dict[tuple[str, str], _Sheet],
                 chart_embeddings: dict[str, list[str]]) -> list[ExtractedValue]:
    """One value per chart data point (series x category), with the freshness check against the workbook."""
    data = rec.get("data") or {}
    axes = data.get("axes") or []
    series_list = data.get("series") or []
    record_flags = {f["code"]: f for f in rec.get("uncertainty_reasons") or []}
    base_flags = _base_flags(rec, skip=_POINT_SCOPED_CODES)
    evidence = rec.get("description_evidence") or {}
    title = data.get("chart_title") or evidence.get("alt_title") or evidence.get("alt_text") or None
    out: list[ExtractedValue] = []
    embeddings = chart_embeddings.get(rec.get("source_part"), [])
    count_stacks = _counts_in_percent_stack(series_list, data.get("rows") or [], run.groupings.get(rec.get("source_part"), []))
    for series in series_list:
        vec = "val" if series["sources"]["val"]["present"] else "yVal"
        field_name = "value" if vec == "val" else "y"
        fmt, fmt_source = _series_format(series, axes, vec)
        series_flags: list[dict[str, Any]] = []
        if fmt_source == "axis" and series["series_number"] in count_stacks and _is_percent_format(fmt):
            fmt, fmt_source = None, None
            series_flags.append(flag(
                "PERCENT_AXIS_NOT_APPLIED", EXTRACTION,
                "This 100% stacked chart stores counts. Its percent value axis shows each series' share of the "
                "stack, so that format was not applied to the stored numbers."))
        date_cats = _is_date_category(series, axes)
        cells = formula_cells(series["sources"][vec].get("formula"))
        for key in ("tx", "cat", "val", "xVal", "yVal", "bubbleSize"):
            for sheet_name, cell in formula_cells((series["sources"].get(key) or {}).get("formula")):
                for part in embeddings:
                    if (part, sheet_name) in sheets:
                        sheets[(part, sheet_name)].consumed.add(cell)
        rows = sorted((r for r in data.get("rows") or [] if r["series_number"] == series["series_number"]),
                      key=lambda r: r["point_index"])
        last_levels: list[Any] = []
        for row in rows:
            levels = list(row.get("category_levels") or [])
            filled = [lv if lv is not None or i == 0 else (last_levels[i] if i < len(last_levels) else None)
                      for i, lv in enumerate(levels)]
            last_levels = filled
            category = row.get("category")
            if isinstance(category, (int, float)) and not isinstance(category, bool):
                category = serial_to_iso(category, data.get("date_system", "1900")) if date_cats else number_text(category)
            value = row.get(field_name)
            if value is None:
                continue
            numeric = isinstance(value, (int, float)) and not isinstance(value, bool)
            flags = list(base_flags) + series_flags
            if not numeric:
                flags.append(record_flags.get("NONNUMERIC_CHART_VALUE") or flag(
                    "NONNUMERIC_CHART_VALUE", EXTRACTION, "The chart cache holds a non-numeric value at this point."))
            if row.get("category") is None and series["sources"]["cat"]["present"]:
                flags.append(record_flags.get("MISSING_CHART_POINTS") or flag(
                    "MISSING_CHART_POINTS", EXTRACTION, "This point has no saved category label."))
            flags += derived_change_flag(series.get("series_name"))
            workbook_ref = None
            k = row["point_index"]
            if series["sources"][vec].get("formula"):
                if k < len(cells):
                    sheet_name, cell = cells[k]
                    sheet = next((sheets[(p, sheet_name)] for p in embeddings if (p, sheet_name) in sheets), None)
                    if sheet is not None and cell in sheet.cells:
                        wb_value = sheet.cells[cell]
                        workbook_ref = {"part": sheet.part, "sheet": sheet_name, "cell": cell, "value": wb_value}
                        if numeric and _same_number(value, wb_value):
                            flags.append(flag("CACHE_MATCHES_WORKBOOK", EXTRACTION,
                                              f"Chart cache equals the embedded workbook cell {sheet_name}!{cell}."))
                        else:
                            flags.append(flag("CACHE_DIFFERS_FROM_WORKBOOK", EXTRACTION,
                                              f"Chart cache {row.get(field_name + '_raw')!r} differs from embedded "
                                              f"workbook {sheet_name}!{cell} = {wb_value!r}; the chart may be stale."))
                    else:
                        flags.append(flag("CACHE_NOT_COMPARED", EXTRACTION,
                                          f"Series formula cell {sheet_name}!{cell} is not in an extracted embedded workbook."))
                else:
                    flags.append(flag("CACHE_NOT_COMPARED", EXTRACTION, "Series formula could not be resolved to a cell."))
            context = {
                **_base_context(rec, run), "object_kind": "chart", "object_title": title,
                "series_name": series.get("series_name"), "category": category,
                "category_levels": ["" if lv is None else str(lv) for lv in reversed(filled)] if len(filled) > 1 else None,
                "axis_title": _value_axis_title(series, axes), "number_source": "stored",
                "chart_type": series.get("chart_type"), "series_count": len(series_list),
                "category_count": (series["sources"].get("cat") or {}).get("declared_count") or len(rows),
                "number_format_source": fmt_source, "workbook_ref": workbook_ref,
                "x": row.get("x") if vec == "yVal" else None,
            }
            locator = {**_shape_locator(rec), "chart_index": rec.get("chart_index_on_slide"),
                       "series_index": series["series_number"] - 1, "point_index": k}
            raw_text = row.get(field_name + "_raw")
            out.append(run.value(
                ExtractionMethod.NATIVE_CHART, locator=locator, context=context,
                raw_text=raw_text if raw_text is not None else str(value),
                raw_number=float(value) if numeric else None, number_format=fmt, flags=flags,
                confidence=1.0 if numeric else 0.0, source_record_id=rec.get("record_id")))
    return out


def table_values(rec: dict[str, Any], run: _Run) -> list[ExtractedValue]:
    """One value per numeric-looking table cell, labeled by its row label and (detected) column header."""
    data = rec.get("data") or {}
    texts = [[c or "" for c in row] for row in data.get("rows") or []]
    superscripts: set[tuple[int, int]] = set()
    cells = run.table_cells(rec)
    if cells is not None:
        texts = [[text for text, _m in row] for row in cells]
        superscripts = {(r, c) for r, row in enumerate(cells) for c, (_t, marked) in enumerate(row) if marked}
    classified = [[text_cell_kind(t) for t in row] for row in texts]
    kinds = [[k for k, _p in row] for row in classified]
    spans = [(int(m["row"]) - 1, int(m["column"]) - 1, int(m["gridSpan"]))
             for m in data.get("merge_metadata") or [] if str(m.get("gridSpan", "")).isdigit()]
    layout = analyze_grid(texts, kinds, spans=spans)
    evidence = rec.get("description_evidence") or {}
    base = _base_context(rec, run)
    title = evidence.get("alt_title") or (base["nearby_text"][0] if base["nearby_text"] else None)
    base_flags = _base_flags(rec)
    out: list[ExtractedValue] = []
    for r, row in enumerate(classified):
        for c, (kind, parts) in enumerate(row):
            if kind != NUM or (r, c) in layout.header_cells:
                continue
            col_header = layout.col_headers[c] if c < len(layout.col_headers) else None
            left = layout.labels_left_of(texts, kinds, r, c)
            raw_text, mark_flags = cell_number_text(texts[r][c])
            if (r, c) in superscripts:
                mark_flags.append(footnote_mark_flag())
            context = {
                **base, "object_kind": "table", "object_title": title, "row_label": layout.row_labels[r],
                "col_header": col_header, "header_row": layout.header_row(),
                "row_labels": left if len(left) > 1 else None, "number_source": "text",
            }
            out.append(run.value(
                ExtractionMethod.NATIVE_TABLE, locator={**_shape_locator(rec), "row": r, "col": c}, context=context,
                raw_text=raw_text, raw_number=parts.number, number_format=None,
                flags=base_flags + derived_change_flag(col_header, layout.row_labels[r]) + mark_flags, confidence=1.0,
                source_record_id=rec.get("record_id")))
    return out


def text_values(lines_source: list[tuple[_TextShape, str]], *, page: int, page_title: str | None, hidden: bool,
                footnotes: list[str], run: _Run, part: str | None, extra_flags: list[dict[str, Any]] | None = None,
                source_kind: str | None = None, source_record_id: str | None = None) -> list[ExtractedValue]:
    """Labeled numbers in text shapes (``SLIDE_TEXT``, confidence 0.7)."""
    out: list[ExtractedValue] = []
    for shape, text in lines_source:
        lines = text.split("\n")
        for li, line in enumerate(lines):
            if not line.strip() or is_note_line(line):
                continue
            for hit in find_labeled_numbers(line):
                label, flags = hit.label, list(extra_flags or [])
                if (li, hit.end) in shape.marks:
                    flags.append(footnote_mark_flag())
                if hit.label_source == "inherited":
                    flags.append(flag("ROW_LABEL_INHERITED", MEANING,
                                      "No text next to this number; the label of the previous number on the line was used."))
                if label is None:
                    label, how = _adjacent_label(lines, li, shape)
                    if label is None:
                        continue
                    flags.append(how)
                if hidden:
                    flags.append(hidden_slide_flag())
                flags += derived_change_flag(hit.label_before, hit.label_after)
                context = {
                    "object_kind": "slide_text", "page_title": page_title, "text_line": line.strip(),
                    "row_label": label, "label_before": hit.label_before, "label_after": hit.label_after,
                    "nearby_text": shape.nearby, "footnotes": footnotes, "hidden": hidden,
                    "position": shape.position, "number_source": "text", "placeholder_type": shape.placeholder,
                    "source_kind": source_kind,
                }
                locator = {"page": page, "part": part, "shape_id": shape.shape_id, "shape_name": shape.name or None,
                           "bbox": shape.bbox, "line": li, "span": [hit.start, hit.end]}
                out.append(run.value(ExtractionMethod.SLIDE_TEXT, locator=locator, context=context,
                                     raw_text=hit.raw_text, raw_number=hit.parts.number, number_format=None,
                                     flags=flags, confidence=TEXT_CONFIDENCE, source_record_id=source_record_id))
    return out


def _adjacent_label(lines: list[str], index: int, shape: _TextShape) -> tuple[str | None, dict[str, str]]:
    for j in (index - 1, index + 1):
        if 0 <= j < len(lines):
            neighbor = lines[j].strip()
            if neighbor and not find_labeled_numbers(neighbor) and re.search(r"[^\W\d_]{2,}", neighbor):
                return neighbor, flag("LABEL_FROM_ADJACENT_LINE", MEANING,
                                       f"The label {neighbor!r} is on the neighboring line of the same text box.")
    if shape.nearby:
        return shape.nearby[0], flag("NEARBY_LABEL_UNVERIFIED", MEANING,
                                     "Nearby text is a candidate label; spatial proximity does not prove it belongs to this number.")
    return None, {}


def image_placeholder(rec: dict[str, Any], run: _Run, area: float | None, media_sha: str | None) -> ExtractedValue:
    data = rec.get("data") or {}
    evidence = rec.get("description_evidence") or {}
    context = {
        **_base_context(rec, run), "object_kind": "image", "object_title": evidence.get("alt_title") or None,
        "area_fraction": round(area, 4) if area is not None else None, "media_part": data.get("media_part"),
        "media_sha256": media_sha, "external": bool(data.get("external")) or None,
    }
    flags = _base_flags(rec) + [flag("IMAGE_NOT_PARSED", EXTRACTION,
                                     "This image may contain numbers; it was not read (OCR / ai_parse_document needed).")]
    locator = {**_shape_locator(rec), "image_index": data.get("image_index")}
    return run.value(ExtractionMethod.IMAGE_PLACEHOLDER, locator=locator, context=context, raw_text=None,
                     raw_number=None, number_format=None, flags=flags, confidence=0.0,
                     source_record_id=rec.get("record_id"))


def extract_pptx_values(doc_id: str, blob: bytes, file_name: str, run_id: str | None = None, *,
                        config: px.ExtractionConfig | None = None) -> ExtractionOutput:
    """PPTX/PPTM bytes -> values (charts, tables, embedded worksheets, slide text, image placeholders) + pages."""
    out = ExtractionOutput(extractor_versions=extractor_versions())
    problem = sniff_problem(blob)
    if problem:
        out.issues.append(issue("NOT_A_PPTX", f"{file_name}: {problem}", severity="error"))
        return out
    try:
        report = px.extract_pptx(blob, source_name=file_name, config=config)
    except Exception as exc:
        out.issues.append(issue("PPTX_PARSE_FAILED", f"{type(exc).__name__}: {exc}", severity="error"))
        return out
    out.extractor_versions = extractor_versions(report)
    try:
        out.document_properties = core_properties(blob)
    except Exception as exc:
        out.issues.append(issue("CORE_PROPERTIES_UNREADABLE", f"{type(exc).__name__}: {exc}", severity="info"))
    out.document_properties["page_count"] = int(report.get("slide_count") or 0)
    size = report.get("slide_size_emu") or {}
    width, height = size.get("width"), size.get("height")
    records: list[dict[str, Any]] = report.get("records") or []
    run = _Run(doc_id=doc_id, run_id=run_id, created_at=utcnow(),
               version=f"{report.get('extractor_version')}+kpe.{ADAPTER_VERSION}",
               scans=_scan_all(blob, report, out.issues), slide_area=float(width * height) if width and height else None,
               include_hidden=(config or px.ExtractionConfig()).include_hidden_slides,
               groupings=_chart_groupings(blob, [r["source_part"] for r in records
                                                 if r["kind"] == "chart" and r.get("source_part")]))

    sheets: dict[tuple[str, str], _Sheet] = {}
    chart_embeddings: dict[str, list[str]] = {}
    for rec in records:
        if rec["kind"] != "embedded_worksheet":
            continue
        data = rec.get("data") or {}
        sheets[(rec["source_part"], data.get("sheet_name"))] = _Sheet(rec, rec["source_part"], data.get("sheet_name"),
                                                                     _sheet_cells(data))
        for ref in data.get("source_references") or []:
            if str(ref.get("referencing_part", "")).startswith("ppt/charts/"):
                chart_embeddings.setdefault(ref["referencing_part"], [])
                if rec["source_part"] not in chart_embeddings[ref["referencing_part"]]:
                    chart_embeddings[ref["referencing_part"]].append(rec["source_part"])

    native_slides: set[int] = set()
    for rec in records:
        if rec["kind"] in {"chart", "table"} and rec.get("slide_number"):
            native_slides.add(rec["slide_number"])
        if rec["kind"] == "embedded_worksheet":
            native_slides.update(s["slide_number"] for s in (rec.get("data") or {}).get("source_references") or []
                                 if s.get("slide_number"))

    by_kind: dict[str, list[dict[str, Any]]] = {}
    for rec in records:
        by_kind.setdefault(rec["kind"], []).append(rec)
    for rec in by_kind.get("chart", []):
        out.values += _guarded(rec, out.issues, chart_values, rec, run, sheets, chart_embeddings)
    for rec in by_kind.get("table", []):
        out.values += _guarded(rec, out.issues, table_values, rec, run)
    for rec in by_kind.get("embedded_worksheet", []):
        out.values += _guarded(rec, out.issues, _worksheet_record_values, rec, run, sheets)
        out.issues += worksheet_issues(rec.get("data") or {}, page=rec.get("slide_number"), part=rec.get("source_part"))
    out.values += _slide_text_values(records, run)
    out.values += _image_values(by_kind.get("image_reference", []), run, blob, native_slides)
    out.issues += _record_issues(records, report)
    out.pages = _pages(doc_id, report, records, run)
    out.values = finalize_values(out.values, out.issues)
    return out


def _guarded(rec: dict[str, Any], issues: list[dict[str, Any]], produce, *args: Any) -> list[ExtractedValue]:
    try:
        return produce(*args)
    except Exception as exc:
        name = rec.get("shape_name") or rec.get("shape_id")
        what = f"{rec['kind']} {name}" if name else rec["kind"]
        issues.append(issue("OBJECT_EXTRACTION_FAILED", f"{what}: {type(exc).__name__}: {exc}",
                            page=rec.get("slide_number"), part=rec.get("source_part"), shape_id=rec.get("shape_id")))
        return []


def _worksheet_record_values(rec: dict[str, Any], run: _Run, sheets: dict[tuple[str, str], _Sheet]) -> list[ExtractedValue]:
    data = rec.get("data") or {}
    sheet = sheets.get((rec["source_part"], data.get("sheet_name")))
    refs = data.get("source_references") or []
    flags = _base_flags(rec)
    if any(str(r.get("referencing_part", "")).startswith("ppt/charts/") for r in refs):
        flags.append(flag("WORKBOOK_UNPLOTTED", ASSOCIATION,
                          "Cell of a chart's embedded workbook that the chart does not plot; may not be shown on the slide."))
    if any(str(r.get("referencing_part", "")).startswith("ppt/slides/") for r in refs):
        flags.append(flag("OLE_VISIBLE_RANGE_UNKNOWN", ASSOCIATION,
                          "Embedded worksheet object: only part of the stored range may be visible on the slide."))
    if not data:
        return []
    return worksheet_values(
        data, doc_id=run.doc_id, method=ExtractionMethod.EMBEDDED_WORKSHEET, extractor=EXTRACTOR,
        extractor_version=run.version, base_locator=_shape_locator(rec), base_context=_base_context(rec, run),
        base_flags=flags, run_id=run.run_id, created_at=run.created_at,
        consumed=sheet.consumed if sheet else set(), source_record_id=rec.get("record_id"))


def _slide_text_values(records: list[dict[str, Any]], run: _Run) -> list[ExtractedValue]:
    out: list[ExtractedValue] = []
    text_records = {r["slide_number"]: r for r in records if r["kind"] == "slide_text"}
    for number in sorted(set(run.scans) | {n for n in text_records if n is not None}):
        scan = run.scans.get(number) or _SlideScan(number=number, error="not scanned")
        rec = text_records.get(number)
        hidden = scan.hidden if scan.error is None else bool(rec and rec.get("hidden_slide"))
        if hidden and not run.include_hidden:
            continue
        title = scan.title or (rec.get("slide_title") if rec else None) or None
        if scan.error is None:
            shapes = [(s, s.text) for s in scan.text_shapes if s.placeholder not in _SKIP_TEXT_PLACEHOLDERS]
            out += text_values(shapes, page=number, page_title=title, hidden=hidden, footnotes=scan.footnotes,
                               run=run, part=scan.part)
        elif rec is not None:
            fallback = _TextShape(None, "", None, (rec.get("data") or {}).get("text") or "", None, "unknown", [])
            out += text_values([(fallback, fallback.text)], page=number, page_title=title, hidden=hidden,
                               footnotes=[], run=run, part=rec.get("source_part"),
                               extra_flags=[flag("TEXT_SHAPE_UNRESOLVED", ASSOCIATION,
                                                 "Slide text was read as one block; the text box is unknown.")])
    for rec in (r for r in records if r["kind"] == "diagram_text"):
        page = rec.get("slide_number")
        shape = _TextShape(rec.get("shape_id"), rec.get("shape_name") or "", None, (rec.get("data") or {}).get("text") or "",
                           bbox_list(rec.get("bounding_box")), rec.get("position_on_slide") or "unknown", [])
        out += text_values([(shape, shape.text)], page=page, page_title=rec.get("slide_title") or None,
                           hidden=bool(rec.get("hidden_slide")), footnotes=run.footnotes(page), run=run,
                           part=rec.get("source_part"), extra_flags=[dict(f) for f in rec.get("uncertainty_reasons") or []],
                           source_kind="smartart", source_record_id=rec.get("record_id"))
    return out


def _image_values(images: list[dict[str, Any]], run: _Run, blob: bytes, native_slides: set[int]) -> list[ExtractedValue]:
    out: list[ExtractedValue] = []
    if not images:
        return out
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        names = set(archive.namelist())
        for rec in images:
            box = rec.get("bounding_box")
            area = (float(box["width"]) * float(box["height"]) / run.slide_area) if box and run.slide_area else None
            page = rec.get("slide_number")
            has_native = page in native_slides
            large = area is not None and area >= LARGE_IMAGE_FRACTION
            if not large and (has_native or (area is not None and area < LOGO_FRACTION)):
                continue
            media = (rec.get("data") or {}).get("media_part") or ""
            if media.rsplit(".", 1)[-1].lower() in {"emf", "wmf"} and has_native:
                continue
            sha = sha256_bytes(archive.read(media)) if media in names else None
            out.append(image_placeholder(rec, run, area, sha))
    return out


def _record_issues(records: list[dict[str, Any]], report: dict[str, Any]) -> list[dict[str, Any]]:
    out = []
    for rec in records:
        detail = "; ".join(rec.get("warnings") or []) or rec.get("what_data_represents") or rec["kind"]
        common = {"page": rec.get("slide_number"), "part": rec.get("source_part"), "shape_id": rec.get("shape_id")}
        if rec["kind"] == "extraction_issue":
            out.append(issue("EXTRACTION_ISSUE", detail, **common))
        elif rec["kind"] == "unsupported_embedding":
            out.append(issue("UNSUPPORTED_EMBEDDING", detail, **common))
        elif rec["kind"] == "external_link" and (rec.get("data") or {}).get("type") != "hyperlink":
            out.append(issue("EXTERNAL_DATA_LINK", f"{(rec.get('data') or {}).get('type')}: "
                                                    f"{(rec.get('data') or {}).get('target')} (not fetched)",
                             severity="info", **common))
    out += [issue("EXTRACTOR_WARNING", w, severity="info") for w in report.get("warnings") or []]
    return out


def _pages(doc_id: str, report: dict[str, Any], records: list[dict[str, Any]], run: _Run) -> list[DocumentPage]:
    size = report.get("slide_size_emu") or {}
    text = {r["slide_number"]: r for r in records if r["kind"] == "slide_text"}
    notes: dict[int, list[str]] = {}
    images: dict[int, int] = {}
    record_shapes: dict[int, list[dict[str, Any]]] = {}
    for r in records:
        page = r.get("slide_number")
        if page is None:
            continue
        if r["kind"] == "speaker_notes" and (r.get("data") or {}).get("text"):
            notes.setdefault(page, []).append(r["data"]["text"])
        if r["kind"] == "image_reference":
            images[page] = images.get(page, 0) + 1
        if r["kind"] in {"chart", "table", "image_reference"} and r.get("shape_id"):
            kind = {"image_reference": "picture"}.get(r["kind"], r["kind"])
            record_shapes.setdefault(page, []).append({"shape_id": r["shape_id"], "kind": kind,
                                                       "name": r.get("shape_name") or "", "bbox": bbox_list(r.get("bounding_box"))})
    pages = []
    for number in range(1, int(report.get("slide_count") or 0) + 1):
        scan = run.scans.get(number) or _SlideScan(number=number, error="not scanned")
        rec = text.get(number)
        body = (rec.get("data") or {}).get("text") if rec else None
        if body is None:
            body = "\n".join(s.text for s in scan.text_shapes)
        hidden = scan.hidden if scan.part and scan.error is None else bool(rec and rec.get("hidden_slide"))
        pages.append(DocumentPage(
            doc_id=doc_id, page_number=number, title=(rec.get("slide_title") if rec else None) or scan.title or None,
            text=body or "", notes="\n".join(notes[number]) if number in notes else None, hidden=hidden,
            image_count=images.get(number, 0), page_width=float(size["width"]) if size.get("width") else None,
            page_height=float(size["height"]) if size.get("height") else None,
            shapes=scan.shapes if scan.error is None and scan.shapes else record_shapes.get(number, []),
            created_at=run.created_at))
    return pages
