from __future__ import annotations

import json
import math
import numbers
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .report import FLAG_INFO, FLAG_ORDER, TILE_NAMES, VERDICT_INFO, VERDICTS, CheckResult, clean, status_label

SHEETS: tuple[tuple[str, str | None], ...] = (
    ("Summary", None), ("Comparisons", "comparisons"), ("KPI verdicts", "kpis"), ("SQL", "sql"),
    ("Platform series", "platform_values"), ("Documents", "documents"), ("Extracted values", "extraction"),
    ("Issues", "issues"),
)
_MAX_CELL = 32_000
_DATE_COLUMNS = {"period_start", "period_end", "first_period", "last_period"}
_UNIT_VALUE_COLUMNS = {"report_value", "platform_value", "abs_diff", "alternative_value", "value"}
_PERCENT_COLUMNS = {"rel_diff"}
_INTEGER_COLUMNS = {"page", "version", "count", "comparisons", "pages", "extracted_values", "periods_compared",
                    "periods_validated", "verified", "review", "blocked", "report_only", "hint_row", "priority"}
_WRAP_COLUMNS = {
    "SQL": {"definition", "value_sql", "components_sql", "checks_sql", "function_sql"},
    "Issues": {"message", "remedy"},
    "KPI verdicts": {"verdict_detail", "definition", "error"},
}
_WIDE_COLUMNS = {"sql": 70, "value_sql": 80, "components_sql": 70, "checks_sql": 60, "function_sql": 70,
                 "definition": 50, "reason": 60, "message": 60, "remedy": 50, "verdict_detail": 60, "error": 50,
                 "flags": 50, "file_path": 50}

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="F1F1ED")
_HEADER_BORDER = Border(bottom=Side(style="thin", color="C3C2B7"))
_TOP = Alignment(vertical="top")
_WRAP = Alignment(vertical="top", wrap_text=True)
_TITLE_FONT = Font(bold=True, size=14)
_MUTED_FONT = Font(color="6B6964")
_FORMATS = {"date": "yyyy-mm-dd", "datetime": "yyyy-mm-dd hh:mm:ss", "percent": "0.00%",
            "percent1": "+0.0%;-0.0%;0.0%", "number": "#,##0.00##", "integer": "0"}


def _text(value: str) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", value)[:_MAX_CELL]


def _set(ws: Any, row: int, col: int, value: Any, number_format: str | None = None, *, wrap: bool = False) -> None:
    value = clean(value)
    cell = ws.cell(row=row, column=col)
    if value is None:
        pass
    elif isinstance(value, bool) or type(value).__name__ == "bool_":
        cell.value = bool(value)
    elif isinstance(value, numbers.Number):
        number = float(value)
        if math.isfinite(number):
            cell.value = int(value) if isinstance(value, numbers.Integral) else number
            if number_format:
                cell.number_format = number_format
    elif isinstance(value, datetime):
        cell.value = value.replace(tzinfo=None)
        cell.number_format = _FORMATS["datetime"]
    elif isinstance(value, date):
        cell.value = value
        cell.number_format = _FORMATS["date"]
    else:
        text = json.dumps(value, default=str, sort_keys=True) if isinstance(value, (dict, list, tuple)) else str(value)
        cell.value = _text(text)
        cell.data_type = "s"
    cell.alignment = _WRAP if wrap else _TOP


def _width(name: str, values: list[Any]) -> int:
    if name in _WIDE_COLUMNS:
        return _WIDE_COLUMNS[name]
    longest = max([len(str(name))] + [len(str(v)) for v in values[:300] if clean(v) is not None])
    return int(min(max(longest + 2, 8), 48))


def _number_format(column: str, row: dict[str, Any]) -> str | None:
    if column in _PERCENT_COLUMNS:
        return _FORMATS["percent1"]
    if column in _UNIT_VALUE_COLUMNS:
        return _FORMATS["percent"] if str(row.get("unit") or "").lower() == "fraction" else _FORMATS["number"]
    if column in _INTEGER_COLUMNS:
        return _FORMATS["integer"]
    return None


def _header(ws: Any, row: int, names: list[str]) -> None:
    for c, name in enumerate(names, start=1):
        cell = ws.cell(row=row, column=c, value=_text(str(name)))
        cell.data_type = "s"
        cell.font, cell.fill, cell.border, cell.alignment = _HEADER_FONT, _HEADER_FILL, _HEADER_BORDER, _WRAP


def _frame_sheet(wb: Workbook, name: str, frame: pd.DataFrame) -> None:
    ws = wb.create_sheet(name)
    columns = [str(c) for c in frame.columns]
    _header(ws, 1, columns)
    records = frame.to_dict(orient="records") if len(frame) else []
    wrap_columns = _WRAP_COLUMNS.get(name, set())
    for r, row in enumerate(records, start=2):
        for c, column in enumerate(columns, start=1):
            _set(ws, r, c, row.get(column), _number_format(column, row), wrap=column in wrap_columns)
    for c, column in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(c)].width = _width(column, [row.get(column) for row in records])
    ws.freeze_panes = "A2"
    if columns:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(records) + 1, 2)}"


def _block(ws: Any, row: int, title: str, head: list[str], rows: list[list[Any]]) -> int:
    ws.cell(row=row, column=1, value=title).font = _HEADER_FONT
    _header(ws, row + 1, head)
    for i, values in enumerate(rows, start=row + 2):
        for c, value in enumerate(values, start=1):
            _set(ws, i, c, value)
    return row + len(rows) + 3


def _summary_sheet(wb: Workbook, result: CheckResult) -> None:
    ws = wb.active
    ws.title = "Summary"
    s = result.summary or {}
    ws.cell(row=1, column=1, value="KPI check report").font = _TITLE_FONT
    ws.cell(row=2, column=1, value="Report values compared with platform values; nothing was published.").font = _MUTED_FONT
    facts = [
        ("Run id", s.get("run_id")), ("Generated at (UTC)", s.get("generated_at")), ("Platform", s.get("platform")),
        ("KPI definitions", s.get("kpi_source")), ("Brand", ", ".join(s.get("brand") or [])),
        ("Reports", len(s.get("reports") or [])), ("Comparisons", len(result.comparisons)),
        ("KPIs", len(result.kpis)), ("Issues", len(result.issues)),
        ("Reused comparisons", int(s.get("reused") or 0)), ("Focus", ", ".join(s.get("focus") or []) or "all KPIs"),
        ("AI mode", s.get("ai_mode")), ("Failed model calls", int(s.get("ai_failures") or 0)),
    ]
    row = 4
    for label, value in facts:
        ws.cell(row=row, column=1, value=label).font = _HEADER_FONT
        _set(ws, row, 2, value)
        row += 1
    row += 1
    tiles = s.get("tiles") or {}
    row = _block(ws, row, "At a glance", ["Tile", "Count"], [[name, tiles.get(key, 0)] for key, name in TILE_NAMES])
    by_flag = s.get("by_flag") or {}
    row = _block(ws, row, "Comparisons by flag", ["Flag", "Count", "Meaning"],
                 [[FLAG_INFO[f][0], by_flag.get(f, 0), FLAG_INFO[f][1]] for f in FLAG_ORDER])
    by_status = s.get("by_status") or {}
    row = _block(ws, row, "Comparisons by status", ["Status", "Label", "Count"],
                 [[k, status_label(k), v] for k, v in by_status.items()])
    verdicts = s.get("verdicts") or {}
    row = _block(ws, row, "KPI verdicts", ["Verdict", "KPIs", "Meaning"],
                 [[v, verdicts.get(v, 0), VERDICT_INFO[v][1]] for v in VERDICTS if verdicts.get(v)])
    by_doc = s.get("by_document") or {}
    row = _block(ws, row, "By report file", ["Report", "Status", "Needs review", "Blocked", "Good", "Other", "Total"],
                 [[name, d.get("status"), d.get("REVIEW", 0), d.get("BLOCKED", 0), d.get("GOOD", 0), d.get("INFO", 0),
                   d.get("total", 0)] for name, d in by_doc.items()])
    timings = s.get("timings") or {}
    if timings:
        _block(ws, row, "Timings", ["Step", "Seconds"], [[k, v] for k, v in timings.items()])
    for letter, width in (("A", 44), ("B", 28), ("C", 70), ("D", 12), ("E", 12), ("F", 12), ("G", 12)):
        ws.column_dimensions[letter].width = width


def write(result: CheckResult, path: str | Path) -> str:
    """Write the workbook to ``path`` (parent folders are created) and return the path."""
    import shutil
    import tempfile

    from .store.state_file import state_location

    target = state_location(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _summary_sheet(wb, result)
    for name, attr in SHEETS[1:]:
        _frame_sheet(wb, name, getattr(result, attr))
    with tempfile.TemporaryDirectory() as tmp:
        local = Path(tmp) / target.name
        wb.save(str(local))
        shutil.copyfile(local, target)
    return str(target)
