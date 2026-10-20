from __future__ import annotations

import io
import json
import random
import re
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from ..ids import scope_key, sha256_bytes
from . import truth

DATASET = "kestrel_synthetic_reports"
GENERATOR_VERSION = "1.0"
LABELS_FILE = "labels.json"
README_FILE = "README.md"
BRAND = "KESTREL"
AUTHOR = "Synthetic"
SEED = 20260610

DECK1 = "Kestrel_Monthly_Brand_Review_2026-06.pptx"
DECK1_V2 = "Kestrel_Monthly_Brand_Review_2026-06_v2.pptx"
EXEC_PDF = "Kestrel_Executive_Summary_Q2-2026.pdf"
UPDATE_DECK = "kestrel-update.pptx"
SCAN_PDF = "Kestrel_Field_Update_scan.pdf"
REPORT_FILES = (DECK1, EXEC_PDF, UPDATE_DECK, SCAN_PDF, DECK1_V2)

MAY = date(2026, 5, 31)
JUNE = date(2026, 6, 30)
Q1_END = date(2026, 3, 31)
PY_MAY = date(2025, 5, 31)
TREND = [("Dec-25", date(2025, 12, 31)), ("Jan-26", date(2026, 1, 31)), ("Feb-26", date(2026, 2, 28)),
         ("Mar-26", date(2026, 3, 31)), ("Apr-26", date(2026, 4, 30)), ("May-26", date(2026, 5, 31))]
REGIONS = ("East", "North", "South", "West")

_ZIP_TIME = (2026, 6, 10, 9, 0, 0)
_CORE_STAMP = "2026-06-10T09:00:00Z"
_CORE_DATE = re.compile(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)")

BRAND_AWARENESS, SHARE_OF_VOICE, PERSISTENCY = "62%", "18%", "71%"
SAMPLES_INFLATION = 1.06
CONFLICT_FACTOR = 0.96
SPEAKER_LIFT_TEXT = "Speaker programs drove +8% TRx lift in the East"


@dataclass
class Label:
    """One labeled number: where it is, what it is, what the pipeline should read and conclude."""

    file: str
    page: int
    locator_hint: dict[str, Any]
    kpi_id: str | None
    period_end: str | None
    period_grain: str | None
    scope: dict[str, Any]
    report_value: float | None
    display_half_unit: float
    expected_status: str
    expected_reason: str
    note: str
    display_text: str | None = None
    unit: str | None = None
    also_acceptable: list[str] = field(default_factory=list)
    label_id: str = ""

    def to_dict(self) -> dict[str, Any]:
        d = asdict(self)
        d["scope_key"] = scope_key(self.scope)
        return d


@dataclass
class DocumentLabel:
    """Document-level metadata (what intake should propose, or what a person confirms)."""

    file: str
    format: str
    brand: str
    report_type: str
    period_label: str
    period_start: str
    period_end: str
    period_grain: str
    publication_date: str
    version: int
    supersedes: str | None
    needs_confirmation: bool
    initial_status: list[str]
    final_status: str
    note: str
    sha256: str = ""


def fmt_int(value: float) -> str:
    return f"{int(round(value)):,}"


def fmt_pct(fraction: float, decimals: int = 1, signed: bool = False) -> str:
    sign = "+" if signed else ""
    return f"{fraction * 100:{sign}.{decimals}f}%"


def fmt_k(value: float) -> str:
    return f"{value / 1000:.1f}K"


def cache_value(value: float) -> float:
    """A chart-cache number with at most 12 significant digits."""
    return float(f"{value:.12g}")


def read_display(text: str) -> tuple[float, float]:
    """(value, half unit) of a printed number as the interpreter should read it (percent -> fraction)."""
    t = text.strip().replace(",", "").replace("−", "-")
    scale = Decimal(1)
    if t.endswith("%"):
        t, scale = t[:-1], Decimal("0.01")
    elif t[-1:] in ("K", "k"):
        t, scale = t[:-1], Decimal(1000)
    elif t[-1:] == "M":
        t, scale = t[:-1], Decimal(1_000_000)
    number = Decimal(t)
    exponent = number.as_tuple().exponent
    step = Decimal(1).scaleb(exponent if isinstance(exponent, int) else 0)
    return float(number * scale), float(step * scale / 2)


def scenario_values(data: truth.PlatformData | None = None) -> dict[str, Any]:
    """Every number the reports print, derived from ``eval.truth`` (plus labeled perturbations)."""
    d = data or truth.load_platform()
    t = truth
    v: dict[str, Any] = {
        "trend": [(cat, pe, cache_value(t.trx(pe, data=d))) for cat, pe in TREND],
        "trx": t.trx(MAY, data=d), "nrx": t.nrx(MAY, data=d), "nbrx": t.nbrx(MAY, data=d),
        "trx_py": t.trx(PY_MAY, data=d), "nrx_py": t.nrx(PY_MAY, data=d), "nbrx_py": t.nbrx(PY_MAY, data=d),
        "share": t.trx_share(MAY, data=d), "share_il17": t.trx_share(MAY, basket="il17_class", data=d),
        "samples": t.samples(MAY, data=d),
        "attain": t.call_attainment(MAY, data=d),
        "attain_writers": t.call_attainment(MAY, eligibility="writers_only", data=d),
        "nrx_region": {r: cache_value(t.nrx(MAY, scope={"region": r}, data=d)) for r in REGIONS},
        "attain_region": {r: cache_value(t.call_attainment(MAY, scope={"region": r}, data=d)) for r in REGIONS},
        "calls": t.calls(MAY, data=d), "new_writers": t.new_writers(MAY, data=d),
        "attendances": t.speaker_attendance(MAY, data=d),
        "distinct_hcps": t.speaker_attendance(MAY, measure="distinct_hcps", data=d),
        "referrals": t.referrals_in(MAY, data=d),
        "referrals_projected": t.referrals_in(MAY, projection="projected", data=d),
        "exfactory": t.exfactory_units(MAY, data=d),
        "exfactory_trend": [(cat, t.exfactory_units(pe, data=d)) for cat, pe in TREND],
        "q2_trx": t.trx(JUNE, "quarter", data=d), "q2_trx_panel": t.trx(JUNE, "quarter", source="panel", data=d),
        "q2_share": t.trx_share(JUNE, "quarter", data=d), "q2_calls": t.calls(JUNE, "quarter", data=d),
        "jun_share": t.trx_share(JUNE, data=d), "jun_share_il17": t.trx_share(JUNE, basket="il17_class", data=d),
        "jun_calls": t.calls(JUNE, data=d),
    }
    v["samples_reported"] = v["samples"] * SAMPLES_INFLATION
    v["nrx_conflict"] = v["nrx"] * CONFLICT_FACTOR
    return v


def _pin_core_dates(xml: bytes) -> bytes:
    return _CORE_DATE.sub(lambda m: m.group(1) + _CORE_STAMP.encode() + m.group(2), xml)


def normalize_ooxml(blob: bytes) -> bytes:
    """Re-zip an OOXML package with fixed member timestamps and pinned core dates."""
    src = zipfile.ZipFile(io.BytesIO(blob))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename.lower().endswith((".xlsx", ".xlsm")):
                data = normalize_ooxml(data)
            elif info.filename == "docProps/core.xml":
                data = _pin_core_dates(data)
            member = zipfile.ZipInfo(info.filename, date_time=_ZIP_TIME)
            member.compress_type = zipfile.ZIP_DEFLATED
            member.create_system = 0
            member.external_attr = 0
            dst.writestr(member, data, compresslevel=6)
    return out.getvalue()


def _new_presentation(title: str, subject: str, *, revision: int = 1, created: datetime | None = None):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    stamp = created or datetime(2026, 6, 10, 9, 0, 0)
    cp = prs.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.title = title
    cp.subject = subject
    cp.keywords = ""
    cp.comments = "Synthetic report generated by kpi_evidence.eval.synth_reports (not real data)."
    cp.category = "synthetic"
    cp.revision = revision
    cp.created = stamp
    cp.modified = stamp
    return prs


def _title_slide(prs, title: str, subtitle: str):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    return slide


def _content_slide(prs, title: str):
    from pptx.util import Pt

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    return slide


def _text_box(slide, text: str, left: float, top: float, width: float, height: float, size: int = 14,
              name: str | None = None):
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    if name:
        box.name = name
    box.text_frame.text = text
    for p in box.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
    return box


def _footer(slide, text: str = "Data through May 2026"):
    return _text_box(slide, text, 0.5, 6.85, 6.0, 0.4, size=11, name="Data footer")


def _table(slide, rows: list[list[str]], left: float, top: float, width: float, col_widths: list[float],
           name: str):
    from pptx.util import Inches, Pt

    height = 0.45 * len(rows)
    frame = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(height))
    frame.name = name
    table = frame.table
    table.first_row = True
    for j, w in enumerate(col_widths):
        table.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = text
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(16)
    return frame


def _category_chart(slide, kind: str, title: str, categories: list[str], series: list[tuple[str, list[float]]],
                    number_format: str, left: float, top: float, width: float, height: float, name: str):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt

    data = CategoryChartData(number_format=number_format)
    data.categories = categories
    for series_name, values in series:
        data.add_series(series_name, values, number_format=number_format)
    chart_type = {"line": XL_CHART_TYPE.LINE_MARKERS, "column": XL_CHART_TYPE.COLUMN_CLUSTERED}[kind]
    frame = slide.shapes.add_chart(chart_type, Inches(left), Inches(top), Inches(width), Inches(height), data)
    frame.name = name
    chart = frame.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    chart.has_legend = False
    axis = chart.value_axis
    axis.tick_labels.number_format = number_format
    axis.tick_labels.number_format_is_linked = False
    return frame


def _picture(slide, png: bytes, left: float, top: float, width: float, height: float, alt_text: str):
    from pptx.util import Inches

    pic = slide.shapes.add_picture(io.BytesIO(png), Inches(left), Inches(top), Inches(width), Inches(height))
    pic._element.nvPicPr.cNvPr.set("descr", alt_text)
    return pic


def _save(prs) -> bytes:
    buf = io.BytesIO()
    prs.save(buf)
    return normalize_ooxml(buf.getvalue())


def _font(size: int):
    from PIL import ImageFont

    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()


def exfactory_chart_png(trend: list[tuple[str, float]], width: int = 1200, height: int = 640) -> bytes:
    """A bar chart rendered to pixels: the numbers exist only as an image (no native chart data)."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    draw.text((40, 20), "Ex-factory units", fill="black", font=_font(36))
    left, right, top, bottom = 110, width - 40, 100, height - 80
    draw.line([(left, top), (left, bottom), (right, bottom)], fill="black", width=2)
    peak = max(v for _, v in trend) * 1.15
    slot = (right - left) / len(trend)
    for i, (cat, value) in enumerate(trend):
        x0 = left + i * slot + slot * 0.2
        x1 = left + (i + 1) * slot - slot * 0.2
        y0 = bottom - (bottom - top) * value / peak
        draw.rectangle([x0, y0, x1, bottom], fill=(47, 95, 152))
        draw.text((x0, y0 - 34), fmt_int(value), fill="black", font=_font(24))
        draw.text((x0 + 6, bottom + 14), cat, fill="black", font=_font(24))
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=False)
    return buf.getvalue()


def scan_page_png(lines: list[str], *, seed: int = SEED, width: int = 1275, height: int = 1650) -> bytes:
    """A 'scanned' page: text rendered to grayscale pixels with seeded speckle noise and a slight skew."""
    from PIL import Image, ImageDraw

    rng = random.Random(seed)
    img = Image.new("L", (width, height), 250)
    draw = ImageDraw.Draw(img)
    y = 140
    for i, line in enumerate(lines):
        draw.text((120, y), line, fill=25, font=_font(48 if i == 0 else 36))
        y += 110 if i == 0 else 80
    for _ in range(2500):
        x, yy = rng.randrange(width), rng.randrange(height)
        img.putpixel((x, yy), rng.randrange(120, 200))
    img = img.rotate(0.6, resample=0, expand=False, fillcolor=250)
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=False)
    return buf.getvalue()


def build_monthly_review(values: dict[str, Any], *, version: int = 1) -> tuple[bytes, list[Label]]:
    """Deck 1 (version=1) or its re-issue with the corrected samples figure (version=2)."""
    file = DECK1 if version == 1 else DECK1_V2
    title = "Kestrel Monthly Brand Review - June 2026" + (" (re-issue)" if version > 1 else "")
    prs = _new_presentation(title, "Monthly brand review, data through May 2026", revision=version,
                            created=datetime(2026, 6, 10 if version == 1 else 12, 9, 0, 0))
    labels: list[Label] = []
    may = MAY.isoformat()

    def add(page, hint, kpi, value_text, status, reason, note, *, period=may, grain="month", scope=None,
            exact=None, unit=None, also=None):
        if exact is not None:
            value, half = exact, 0.0
        elif value_text is None:
            value, half = None, 0.0
        else:
            value, half = read_display(value_text)
        labels.append(Label(file=file, page=page, locator_hint=hint, kpi_id=kpi, period_end=period, period_grain=grain,
                            scope=dict(scope or {}), report_value=value, display_half_unit=half,
                            expected_status=status, expected_reason=reason, note=note, display_text=value_text,
                            unit=unit, also_acceptable=list(also or [])))

    subtitle = "June 2026 edition · Data through May 2026"
    if version > 1:
        subtitle += "\nRe-issue: the samples figure is corrected in this version."
    _title_slide(prs, "Kestrel Monthly Brand Review", subtitle)

    s = _content_slide(prs, "TRx trend")
    cats = [c for c, _, _ in values["trend"]]
    _category_chart(s, "line", "TRx by month", cats, [("TRx", [v for _, _, v in values["trend"]])], "#,##0",
                    0.8, 1.5, 11.7, 5.1, "TRx trend chart")
    _footer(s)
    for cat, pe, v in values["trend"]:
        add(2, {"object": "TRx by month", "series": "TRx", "category": cat}, "trx", None, "VERIFIED",
            "match_under_governed_default", f"Native chart cache (12 significant digits); equals account-feed TRx for {cat}.",
            period=pe.isoformat(), exact=v, unit="count")
        labels[-1].display_text = repr(v)

    s = _content_slide(prs, "Key performance indicators – May 2026")
    samples_text = fmt_int(values["samples_reported"] if version == 1 else values["samples"])
    rows = [
        ["KPI", "May 2026", "vs PY %"],
        ["TRx", fmt_k(values["trx"]), fmt_pct(values["trx"] / values["trx_py"] - 1, signed=True)],
        ["NRx", fmt_int(values["nrx"]), fmt_pct(values["nrx"] / values["nrx_py"] - 1, signed=True)],
        ["NBRx", fmt_int(values["nbrx"]), fmt_pct(values["nbrx"] / values["nbrx_py"] - 1, signed=True)],
        ["TRx share", fmt_pct(values["share"]), "–"],
        ["IL-17 class share", fmt_pct(values["share_il17"]), "–"],
        ["Samples", samples_text, "–"],
        ["Call attainment (writers only)", fmt_pct(values["attain_writers"]), "–"],
    ]
    _table(s, rows, 0.8, 1.5, 11.0, [5.0, 3.0, 3.0], "KPI table")
    _footer(s)
    obj = "Key performance indicators"
    add(3, {"object": obj, "row": "TRx", "col": "May 2026"}, "trx", rows[1][1], "VERIFIED",
        "match_under_governed_default", "Rounded display (5.7K = 5,700, half unit 50); within display rounding.",
        unit="count")
    add(3, {"object": obj, "row": "NRx", "col": "May 2026"}, "nrx", rows[2][1], "REPORT_CONFLICT",
        "conflicting_reports", "Agrees with the platform, but the executive summary prints a different May NRx "
        f"({fmt_int(values['nrx_conflict'])}).", unit="count")
    add(3, {"object": obj, "row": "NBRx", "col": "May 2026"}, "nbrx", rows[3][1], "VERIFIED",
        "match_under_governed_default", "Rounded to the unit; within display rounding.", unit="count")
    add(3, {"object": obj, "row": "TRx share", "col": "May 2026"}, "trx_share", rows[4][1], "VERIFIED",
        "match_under_governed_default", "Governed advanced-therapy basket, rounded to 0.1 point.", unit="fraction")
    add(3, {"object": obj, "row": "IL-17 class share", "col": "May 2026"}, "trx_share", rows[5][1], "VERIFIED",
        "match_under_stated_permitted_variant", "The row states the IL-17 class basket, a permitted variant.",
        unit="fraction")
    if version == 1:
        add(3, {"object": obj, "row": "Samples", "col": "May 2026"}, "samples", samples_text, "VARIANCE_REVIEW", "",
            "Overstated by 6% versus the CRM feed; no declared definition explains it.", unit="count")
    else:
        add(3, {"object": obj, "row": "Samples", "col": "May 2026"}, "samples", samples_text, "VERIFIED",
            "match_under_governed_default", "The re-issue corrects the samples figure; it now matches the platform.",
            unit="count")
    add(3, {"object": obj, "row": "Call attainment (writers only)", "col": "May 2026"}, "call_attainment",
        rows[7][1], "INCOMPATIBLE_DEFINITION", "matches_alternative_definition",
        "Writers-only eligibility is a declared but non-permitted definition (governed: all HCPs with a plan).",
        unit="fraction")
    for i, name in ((1, "TRx"), (2, "NRx"), (3, "NBRx")):
        add(3, {"object": obj, "row": name, "col": "vs PY %"}, None, rows[i][2], "UNMAPPED_CAPTURED",
            "no_registry_match", "A change versus prior year (DERIVED_CHANGE): captured, not a registry KPI.",
            unit="fraction")

    s = _content_slide(prs, "Regional performance – May 2026")
    _category_chart(s, "column", "NRx by region", list(REGIONS), [("May 2026", [values["nrx_region"][r] for r in REGIONS])],
                    "#,##0", 0.5, 1.5, 6.0, 5.0, "NRx by region chart")
    _category_chart(s, "column", "Call attainment by region", list(REGIONS),
                    [("May 2026", [values["attain_region"][r] for r in REGIONS])], "0.0%", 6.8, 1.5, 6.0, 5.0,
                    "Call attainment by region chart")
    _footer(s)
    for r in REGIONS:
        add(4, {"object": "NRx by region", "series": "May 2026", "category": r}, "nrx", None, "VERIFIED",
            "match_under_governed_default", f"Region scope from the category ({r}); exact chart cache.",
            scope={"region": r}, exact=values["nrx_region"][r], unit="count")
        labels[-1].display_text = repr(values["nrx_region"][r])
    for r in REGIONS:
        add(4, {"object": "Call attainment by region", "series": "May 2026", "category": r}, "call_attainment", None,
            "VERIFIED", "match_under_governed_default",
            f"Percent-formatted cache keeps the fraction; governed ratio of sums for {r}.",
            scope={"region": r}, exact=values["attain_region"][r], unit="fraction")
        labels[-1].display_text = repr(values["attain_region"][r])

    s = _content_slide(prs, "Field activity – May 2026")
    rows = [
        ["Field activity", "May 2026"],
        ["Calls (details delivered)", fmt_int(values["calls"])],
        ["HCPs attended speaker programs", fmt_int(values["distinct_hcps"])],
        ["Referrals (projected)", fmt_int(values["referrals_projected"])],
        ["New writers", fmt_int(values["new_writers"])],
    ]
    _table(s, rows, 0.8, 1.5, 9.0, [6.0, 3.0], "Field activity table")
    _footer(s)
    obj = "Field activity"
    add(5, {"object": obj, "row": rows[1][0], "col": "May 2026"}, "calls", rows[1][1], "VERIFIED",
        "match_under_governed_default", "Exact count of details delivered.", unit="count")
    add(5, {"object": obj, "row": rows[2][0], "col": "May 2026"}, "speaker_attendance", rows[2][1],
        "INCOMPATIBLE_DEFINITION", "matches_alternative_definition",
        f"Distinct HCPs ({fmt_int(values['distinct_hcps'])}), not attendances ({fmt_int(values['attendances'])}).",
        unit="count")
    add(5, {"object": obj, "row": rows[3][0], "col": "May 2026"}, "referrals_in", rows[3][1],
        "INCOMPATIBLE_DEFINITION", "matches_alternative_definition",
        f"Projected (observed {fmt_int(values['referrals'])} / 0.80); the governed KPI is observed only.", unit="count")
    add(5, {"object": obj, "row": rows[4][0], "col": "May 2026"}, "new_writers", rows[4][1], "VERIFIED",
        "match_under_governed_default", "Exact count of first-time writers.", unit="count")

    s = _content_slide(prs, "Market research")
    tracker = [["Brand tracker", "Q1 2026"], ["Brand awareness (aided)", BRAND_AWARENESS],
               ["12-month persistency", PERSISTENCY]]
    audit = [["Promotion audit", "May 2026"], ["Share of voice", SHARE_OF_VOICE]]
    _table(s, tracker, 0.8, 1.5, 9.0, [6.0, 3.0], "Brand tracker table")
    _table(s, audit, 0.8, 3.7, 9.0, [6.0, 3.0], "Promotion audit table")
    _footer(s)
    note = "Survey / vendor KPI with no platform table: captured, not verified."
    for i, kpi in ((1, "brand_awareness"), (2, "persistency_12m")):
        add(6, {"object": "Brand tracker", "row": tracker[i][0], "col": "Q1 2026"}, kpi, tracker[i][1], "REPORT_ONLY",
            "no_platform_definition", note + " Quarterly tracker wave.", period=Q1_END.isoformat(), grain="quarter",
            unit="fraction")
    add(6, {"object": "Promotion audit", "row": audit[1][0], "col": "May 2026"}, "share_of_voice", audit[1][1],
        "REPORT_ONLY", "no_platform_definition", note, unit="fraction")

    s = _content_slide(prs, "Supply – ex-factory units")
    rows = [["Supply", "May 2026"], ["Ex-factory units", fmt_int(values["exfactory"])]]
    _table(s, rows, 0.8, 1.5, 9.0, [6.0, 3.0], "Ex-factory table")
    _footer(s)
    add(7, {"object": "Supply", "row": "Ex-factory units", "col": "May 2026"}, "exfactory_units", rows[1][1],
        "MISSING_ACCESS", "permission_denied",
        "The governed source exists, but the pipeline identity cannot read the restricted schema.", unit="count")

    s = _content_slide(prs, "Ex-factory units trend")
    _picture(s, exfactory_chart_png(values["exfactory_trend"]), 1.2, 1.4, 10.9, 5.6,
             "Bar chart image of ex-factory units by month")
    add(8, {"object": "Picture", "kind": "image"}, None, None, "UNCERTAIN_EXTRACTION", "image_not_parsed",
        "Image-only slide: nothing is read from the pixels locally (ai_parse_document on Databricks), so no KPI is "
        "assigned; the slide title only suggests ex-factory units.")

    s = _content_slide(prs, "Speaker program lift – May 2026")
    _text_box(s, SPEAKER_LIFT_TEXT, 0.8, 1.8, 11.5, 1.0, size=24, name="Lift callout")
    _footer(s)
    add(9, {"object": "Lift callout", "text": SPEAKER_LIFT_TEXT}, "speaker_program_lift", "+8%", "REPORT_ONLY",
        "causal_claim", "A causal claim is captured as stated and never verified.", scope={"region": "East"},
        unit="fraction")

    return _save(prs), labels


def build_update_deck(values: dict[str, Any]) -> tuple[bytes, list[Label]]:
    prs = _new_presentation("Field performance update", "Latest month", created=datetime(2026, 7, 8, 9, 0, 0))
    _title_slide(prs, "Field performance update", "Latest month")
    s = _content_slide(prs, "Headline numbers")
    rows = [["Measure", "Latest month"], ["Share", fmt_pct(values["jun_share_il17"])], ["Calls", fmt_int(values["jun_calls"])]]
    _table(s, rows, 0.8, 1.5, 9.0, [6.0, 3.0], "Headline table")
    june = JUNE.isoformat()
    share_value, share_half = read_display(rows[1][1])
    calls_value, calls_half = read_display(rows[2][1])
    labels = [
        Label(UPDATE_DECK, 2, {"object": "Headline numbers", "row": "Share", "col": "Latest month"}, "trx_share", june,
              "month", {}, share_value, share_half, "INCOMPATIBLE_DEFINITION", "matches_alternative_definition",
              "After confirmation (June 2026): an unqualified 'share' equal to the IL-17 class share; the governed "
              f"basket gives {fmt_pct(values['jun_share'], 2)}.", rows[1][1], "fraction"),
        Label(UPDATE_DECK, 2, {"object": "Headline numbers", "row": "Calls", "col": "Latest month"}, "calls", june,
              "month", {}, calls_value, calls_half, "VERIFIED", "match_under_governed_default",
              "After confirmation (June 2026): exact count of details delivered.", rows[2][1], "count"),
    ]
    return _save(prs), labels


def _pdf_doc(buf: io.BytesIO, title: str, subject: str):
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate

    return SimpleDocTemplate(buf, pagesize=letter, invariant=1, title=title, author=AUTHOR, subject=subject,
                             creator="kpi_evidence.eval.synth_reports", leftMargin=54, rightMargin=54,
                             topMargin=54, bottomMargin=54)


def _pdf_table(rows: list[list[str]]):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    table = Table(rows, colWidths=[260, 140])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.75, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 11),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return table


def build_exec_summary(values: dict[str, Any]) -> tuple[bytes, list[Label]]:
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, Spacer

    styles = getSampleStyleSheet()
    q2 = [["KPI", "Q2 2026"], ["TRx (retail panel)", fmt_int(values["q2_trx_panel"])],
          ["TRx share", fmt_pct(values["q2_share"])], ["Calls", fmt_int(values["q2_calls"])]]
    snap = [["KPI", "May 2026"], ["NRx", fmt_int(values["nrx_conflict"])]]
    buf = io.BytesIO()
    doc = _pdf_doc(buf, "Kestrel Executive Summary Q2 2026", "Quarterly business review")
    doc.build([
        Paragraph("Kestrel Executive Summary – Q2 2026", styles["Title"]),
        Paragraph("Quarterly business review for the leadership team.", styles["Normal"]),
        Spacer(1, 16),
        Paragraph("Q2 2026 headline KPIs", styles["Heading2"]),
        _pdf_table(q2),
        Spacer(1, 20),
        Paragraph("May 2026 snapshot", styles["Heading2"]),
        _pdf_table(snap),
    ])
    q2_end = JUNE.isoformat()

    def lab(hint, kpi, text, period, grain, status, reason, note, unit, also=()):
        value, half = read_display(text)
        return Label(EXEC_PDF, 1, hint, kpi, period, grain, {}, value, half, status, reason, note, text, unit,
                     list(also))

    labels = [
        lab({"object": "Q2 2026 headline KPIs", "row": "TRx (retail panel)", "col": "Q2 2026"}, "trx", q2[1][1], q2_end,
            "quarter", "VARIANCE_REVIEW", "single_period_gap_unexplained",
            f"Panel source with a one-month lag: Apr+May only ({fmt_int(values['q2_trx_panel'])} vs governed "
            f"{fmt_int(values['q2_trx'])}, {values['q2_trx_panel'] / values['q2_trx'] - 1:+.1%}). The panel has no "
            f"June, so it cannot reproduce the quarter.", "count"),
        lab({"object": "Q2 2026 headline KPIs", "row": "TRx share", "col": "Q2 2026"}, "trx_share", q2[2][1], q2_end,
            "quarter", "VERIFIED", "match_under_governed_default", "Quarter share from the account feed.", "fraction"),
        lab({"object": "Q2 2026 headline KPIs", "row": "Calls", "col": "Q2 2026"}, "calls", q2[3][1], q2_end,
            "quarter", "VERIFIED", "match_under_governed_default", "Exact quarter sum of details delivered.", "count"),
        lab({"object": "May 2026 snapshot", "row": "NRx", "col": "May 2026"}, "nrx", snap[1][1], MAY.isoformat(),
            "month", "REPORT_CONFLICT", "conflicting_reports",
            f"Printed 4% below the monthly brand review ({fmt_int(values['nrx'])}) for the same KPI and period.", "count"),
    ]
    return buf.getvalue(), labels


def build_scan(values: dict[str, Any]) -> tuple[bytes, list[Label]]:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    png = scan_page_png(["Kestrel field update - May 2026", f"Calls: {fmt_int(values['calls'])}",
                         f"Samples: {fmt_int(values['samples'])}", f"Call attainment: {fmt_pct(values['attain'])}"])
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter, invariant=1)
    c.setAuthor(AUTHOR)
    c.setTitle("Scanned field update")
    c.setCreator("kpi_evidence.eval.synth_reports")
    w, h = letter
    c.drawImage(ImageReader(io.BytesIO(png)), 0, 0, width=w, height=h)
    c.showPage()
    c.save()
    labels = [Label(SCAN_PDF, 1, {"object": "Page image", "kind": "scan"}, None, MAY.isoformat(), "month", {}, None,
                    0.0, "UNCERTAIN_EXTRACTION", "image_not_parsed",
                    "Scanned page without a text layer: an image placeholder locally (ai_parse_document on Databricks "
                    "would read calls, samples and attainment).")]
    return buf.getvalue(), labels


def document_labels(hashes: dict[str, str]) -> list[DocumentLabel]:
    may_start, may_end = "2026-05-01", MAY.isoformat()
    docs = [
        DocumentLabel(DECK1, "PPTX", BRAND, "monthly_brand_review", "May 2026", may_start, may_end, "month", "2026-06-10",
                      1, None, False, ["READY", "PROCESSED"], "SUPERSEDED",
                      "Brand and report type from the file name and title; period from the 'Data through May 2026' footer."),
        DocumentLabel(EXEC_PDF, "PDF", BRAND, "executive_summary", "Q2 2026", "2026-04-01", JUNE.isoformat(), "quarter",
                      "2026-07-06", 1, None, False, ["READY", "PROCESSED"], "PROCESSED",
                      "Digital PDF (text layer); period from the title 'Q2 2026'."),
        DocumentLabel(UPDATE_DECK, "PPTX", BRAND, "field_update", "June 2026", "2026-06-01", JUNE.isoformat(), "month",
                      "2026-07-08", 1, None, True, ["AWAITING_CONFIRMATION"], "PROCESSED",
                      "No brand in the title and no period anywhere: a person confirms brand/type/period (June 2026)."),
        DocumentLabel(SCAN_PDF, "PDF", BRAND, "field_update", "May 2026", may_start, may_end, "month", "2026-06-05",
                      1, None, False, ["READY", "AWAITING_CONFIRMATION", "PROCESSED"], "PROCESSED",
                      "Image-only page; the period is not machine-readable locally, so it may need confirmation."),
        DocumentLabel(DECK1_V2, "PPTX", BRAND, "monthly_brand_review", "May 2026", may_start, may_end, "month",
                      "2026-06-12", 2, DECK1, False, ["READY", "PROCESSED"], "PROCESSED",
                      "Re-issue of the monthly brand review with the samples figure corrected; supersedes version 1."),
    ]
    for d in docs:
        d.sha256 = hashes.get(d.file, "")
    return docs


def coverage_expectations() -> dict[str, Any]:
    """Goal-list states expected when coverage is assessed at the end of June 2026."""
    expected = {k: "FOUND" for k in ("trx", "nrx", "nbrx", "trx_share", "calls", "call_attainment", "samples",
                                     "speaker_attendance", "new_writers", "referrals_in", "brand_awareness",
                                     "share_of_voice", "persistency_12m")}
    expected["exfactory_units"] = "NO_ACCESS"
    expected["active_referrers"] = "NOT_FOUND"
    return {"brand": BRAND, "as_of": JUNE.isoformat(),
            "expected_period_end": {"monthly": MAY.isoformat(), "quarterly": Q1_END.isoformat()},
            "expected": dict(sorted(expected.items())),
            "note": "active_referrers is on the goal list but in no report; ex-factory is reported but unreadable."}


def _assign_ids(labels: list[Label]) -> None:
    from ..ids import stable_id

    for lab in labels:
        lab.label_id = stable_id(lab.file, lab.page, lab.locator_hint, lab.kpi_id, lab.period_end, lab.scope,
                                 length=16)


def build_labels_document(labels: list[Label], docs: list[DocumentLabel], values: dict[str, Any]) -> dict[str, Any]:
    counts: dict[str, int] = {}
    for lab in labels:
        counts[lab.expected_status] = counts.get(lab.expected_status, 0) + 1
    return {
        "dataset": DATASET,
        "generator_version": GENERATOR_VERSION,
        "brand": BRAND,
        "platform_reference": {
            "june_2026": {"trx": 5783.041, "q2_2026_trx": 17220.318, "trx_share_advanced_therapy": 0.0917,
                          "trx_share_il17_class": 0.2621, "call_attainment_all_hcps": 0.8442,
                          "call_attainment_writers_only": 0.9192, "call_attainment_mean_of_ratios": 0.8576,
                          "referrals_observed": 147, "referrals_projected": 183.75, "speaker_attendances": 64,
                          "speaker_distinct_hcps": 47, "west_trx": 1265.7},
            "may_2026": {"trx": round(values["trx"], 6), "trx_share": round(values["share"], 6),
                         "trx_share_il17_class": round(values["share_il17"], 6),
                         "call_attainment": round(values["attain"], 6),
                         "call_attainment_writers_only": round(values["attain_writers"], 6),
                         "speaker_attendances": values["attendances"], "speaker_distinct_hcps": values["distinct_hcps"],
                         "referrals_observed": values["referrals"], "referrals_projected": values["referrals_projected"],
                         "samples": values["samples"]},
        },
        "documents": [asdict(d) for d in docs],
        "coverage": coverage_expectations(),
        "status_counts": dict(sorted(counts.items())),
        "labels": [lab.to_dict() for lab in labels],
    }


def _readme(doc: dict[str, Any]) -> str:
    plain = {
        "VERIFIED": "matches the platform under the governed definition",
        "REPORT_ONLY": "captured as reported; no platform table can check it",
        "UNMAPPED_CAPTURED": "captured, but it is not a registry KPI",
        "VARIANCE_REVIEW": "differs from the platform and nothing declared explains it; a person reviews",
        "INCOMPATIBLE_DEFINITION": "the report uses a different (declared) definition; a person reviews",
        "UNCERTAIN_EXTRACTION": "the number could not be read with confidence; a person checks it",
        "REPORT_CONFLICT": "two current reports disagree on the same KPI and period",
        "MISSING_ACCESS": "the platform source exists but the pipeline may not read it",
    }
    lines = [
        "# Synthetic report set (labeled)",
        "",
        "Generated by `kpi_evidence.eval.synth_reports.generate_all()` from the synthetic platform data. The",
        "output is byte-stable; **do not edit these files by hand**.",
        "",
        "Every number is computed by the independent pandas truth path (`kpi_evidence.eval.truth`) and then",
        "changed only where a scenario says so. `labels.json` lists, for every labeled number, where it is,",
        "which KPI it is, the value the pipeline should read (percent as a fraction, `5.7K` as 5,700), the",
        "display half-unit and the expected status and reason. `kpi_evidence.eval.evaluate.run_evaluation()`",
        "runs the pipeline over these files and scores it.",
        "",
        "Deck 1 is published in June with **data through May 2026**, so it prints May values. The",
        "reference values (TRx 5,783.041, share 9.17% / 26.21%, attainment 84.42% / 91.92%, 47 distinct HCPs,",
        "183.75 projected referrals) are the June 2026 platform values; the June update deck uses them.",
        "In May the same traps give 46 distinct HCPs (83 attendances), 210 projected",
        "referrals (168 observed) and 95.8% writers-only attainment (88.1% governed).",
        "",
        "## Documents",
        "",
    ]
    for d in doc["documents"]:
        lines += [f"### `{d['file']}`", "",
                  f"* {d['note']}",
                  f"* Metadata: brand {d['brand']}, report type `{d['report_type']}`, period {d['period_label']} "
                  f"({d['period_grain']}, ending {d['period_end']}), version {d['version']}"
                  + (f", supersedes `{d['supersedes']}`" if d["supersedes"] else "") + ".",
                  f"* Expected document status after the evaluation run: `{d['final_status']}`"
                  + (" (after a person confirms the metadata)" if d["needs_confirmation"] else "") + ".",
                  "",
                  "| Page | Where | Printed | KPI | Expected | Why |",
                  "|---:|---|---|---|---|---|"]
        for lab in (x for x in doc["labels"] if x["file"] == d["file"]):
            hint = lab["locator_hint"]
            where = " / ".join(str(hint[k]) for k in ("object", "row", "col", "series", "category") if k in hint)
            if lab["report_value"] is None:
                shown = "(image)"
            elif "series" in hint:
                value = lab["report_value"]
                shown = "cache " + (fmt_pct(value) if lab["unit"] == "fraction" else f"{value:,.1f}")
            else:
                shown = lab["display_text"]
            scope = ", ".join(f"{k}={v}" for k, v in lab["scope"].items())
            kpi = (lab["kpi_id"] or "(none)") + (f" [{scope}]" if scope else "")
            lines.append(f"| {lab['page']} | {where} | {shown} | {kpi} | `{lab['expected_status']}` | {lab['note']} |")
        lines.append("")
    lines += ["## Expected statuses in plain language", ""]
    for status, n in doc["status_counts"].items():
        lines.append(f"* `{status}` ({n} labels): {plain.get(status, status)}.")
    cov = doc["coverage"]
    lines += ["", "## Goal-list coverage", "",
              f"Assessed for {cov['brand']} as of {cov['as_of']} (one period of lag: May 2026 for monthly KPIs, "
              "Q1 2026 for the quarterly brand tracker): every goal-list KPI is FOUND except `active_referrers` "
              "(NOT_FOUND: no report contains it) and `exfactory_units` (NO_ACCESS: reported, but the platform "
              "source is not readable).", ""]
    return "\n".join(lines)


def generate_all(out_dir: str | Path | None = None, *, data_dir: str | Path | None = None) -> list[Path]:
    """Write the five synthetic reports, ``labels.json`` and ``README.md``; return the report paths."""
    from ..settings import demo_dir

    out = Path(out_dir) if out_dir is not None else demo_dir("reports")
    out.mkdir(parents=True, exist_ok=True)
    values = scenario_values(truth.load_platform(data_dir))
    built: dict[str, tuple[bytes, list[Label]]] = {
        DECK1: build_monthly_review(values, version=1),
        EXEC_PDF: build_exec_summary(values),
        UPDATE_DECK: build_update_deck(values),
        SCAN_PDF: build_scan(values),
        DECK1_V2: build_monthly_review(values, version=2),
    }
    paths: list[Path] = []
    labels: list[Label] = []
    for name in REPORT_FILES:
        blob, file_labels = built[name]
        path = out / name
        if not path.exists() or path.read_bytes() != blob:
            path.write_bytes(blob)
        paths.append(path)
        labels += file_labels
    _assign_ids(labels)
    doc = build_labels_document(labels, document_labels({n: sha256_bytes(built[n][0]) for n in REPORT_FILES}), values)
    _write_text(out / LABELS_FILE, json.dumps(doc, indent=2, ensure_ascii=False) + "\n")
    _write_text(out / README_FILE, _readme(doc))
    return paths


def _write_text(path: Path, text: str) -> None:
    data = text.encode("utf-8")
    if not path.exists() or path.read_bytes() != data:
        path.write_bytes(data)


if __name__ == "__main__":
    for p in generate_all():
        print(p)
