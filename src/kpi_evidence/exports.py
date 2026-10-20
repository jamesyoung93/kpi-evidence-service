from __future__ import annotations

import io
import json
from datetime import date
from typing import Any

import pandas as pd

from .ids import utcnow
from .serving import format_value, in_list, is_missing, kpi_labels, published_current, scope_text, to_date

SCORECARD_COLUMNS = (
    "kpi_id", "kpi_label", "scope", "period", "period_end", "value", "unit", "status", "validation_status",
    "platform_value", "variance_rel", "approved_by", "approved_at", "evidence", "caveat", "publication_id",
    "evidence_uri",
)

DISPLAY_HEADERS = {
    "kpi_label": "KPI", "scope": "Scope", "period": "Period", "value": "Value", "unit": "Unit",
    "status": "Status", "platform_value": "Platform value", "variance_rel": "Difference vs platform",
    "approved_by": "Approved by", "approved_at": "Approved at", "evidence": "Source", "caveat": "Caveat",
    "kpi_id": "KPI id", "publication_id": "Publication id",
}
_SHEET_COLUMNS = ("kpi_label", "scope", "period", "value", "unit", "status", "platform_value", "variance_rel",
                  "approved_by", "approved_at", "evidence", "caveat", "kpi_id", "publication_id")
_ROWS_PER_SLIDE = 12
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _as_dict(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    if isinstance(value, str) and value:
        try:
            out = json.loads(value)
        except ValueError:
            return {}
        return out if isinstance(out, dict) else {}
    return {}


def locator_text(file_name: str | None, locator: Any, context: Any = None) -> str:
    """"deck.pptx, slide 2, chart 'TRx by month', series 'TRx', point 'May-26'" from a locator + raw context."""
    loc, ctx = _as_dict(locator), _as_dict(context)
    parts = [file_name] if file_name else []
    page = loc.get("page") if loc.get("page") is not None else loc.get("slide")
    if page is not None:
        word = "slide" if str(file_name or "").lower().endswith((".pptx", ".pptm")) else "page"
        parts.append(f"{word} {page}")
    for key, word in (("chart_title", "chart"), ("table_title", "table")):
        if ctx.get(key):
            parts.append(f"{word} '{ctx[key]}'")
    if loc.get("sheet"):
        parts.append(f"sheet '{loc['sheet']}'")
    series = ctx.get("series_name") if ctx.get("series_name") is not None else loc.get("series")
    if series is not None:
        parts.append(f"series '{series}'")
    if ctx.get("category") is not None:
        parts.append(f"point '{ctx['category']}'")
    elif loc.get("point_index") is not None:
        parts.append(f"point {int(loc['point_index']) + 1}")
    if ctx.get("row_label") and ctx.get("col_header"):
        parts.append(f"row '{ctx['row_label']}', column '{ctx['col_header']}'")
    elif loc.get("row") is not None and loc.get("col") is not None:
        parts.append(f"row {loc['row']}, column {loc['col']}")
    if loc.get("cell"):
        parts.append(f"cell {loc['cell']}")
    return ", ".join(str(p) for p in parts)


def _value_contexts(store, result_ids: list[str]) -> dict[str, dict[str, Any]]:
    if not result_ids:
        return {}
    marks, params = in_list("r", result_ids)
    df = store.query(
        f"SELECT r.result_id, v.context FROM {store.fq('kpi_results')} r "
        f"JOIN {store.fq('observations')} o ON o.observation_id = r.observation_id "
        f"JOIN {store.fq('extracted_values')} v ON v.value_id = o.value_id WHERE r.result_id IN ({marks})", params)
    return {str(row.result_id): _as_dict(row.context) for row in df.itertuples()}


def _caveat(row: dict[str, Any]) -> str:
    label = str(row.get("display_label") or "")
    unit = row.get("unit")
    tol = _as_dict(row.get("tolerance"))
    rationale = str(row.get("rationale") or "").strip().rstrip(".")
    platform = row.get("platform_value")
    variance = row.get("variance_rel")
    if label == "Matches platform":
        rel = tol.get("rel")
        within = f" (tolerance {float(rel) * 100:.1f}%)" if not is_missing(rel) and float(rel) > 0 else ""
        return f"Matches the platform value computed with the governed definition{within}."
    if label == "Accepted variance":
        diff = f", {float(variance) * 100:+.1f}%" if not is_missing(variance) else ""
        base = (f"Differs from the platform value ({format_value(platform, unit)}{diff}); accepted by a reviewer"
                if not is_missing(platform) else "Differs from the platform definition; accepted by a reviewer")
        return f"{base}: {rationale}." if rationale else f"{base}."
    if label == "Report-only":
        return "Captured from the report; no platform source holds this KPI, so it is not independently verified."
    if label == "Corrected by reviewer":
        return f"Value corrected by a reviewer: {rationale}." if rationale else "Value corrected by a reviewer."
    return label


def scorecard_frame(store, registry, brand: str) -> pd.DataFrame:
    """Published current values for ``brand`` with plain labels, evidence text and a caveat sentence."""
    pubs = published_current(store, brand)
    if pubs.empty:
        return pd.DataFrame(columns=list(SCORECARD_COLUMNS))
    labels = kpi_labels(store, registry)
    doc_ids = sorted({str(d) for d in pubs["evidence_hash"] if not is_missing(d)})
    files: dict[str, str] = {}
    if doc_ids:
        marks, params = in_list("d", doc_ids)
        df = store.query(f"SELECT doc_id, file_name FROM {store.fq('documents')} WHERE doc_id IN ({marks})", params)
        files = {str(r.doc_id): str(r.file_name) for r in df.itertuples()}
    contexts = _value_contexts(store, sorted({str(r) for r in pubs["result_id"]}))
    order = {g["kpi_id"]: i for i, g in enumerate(registry.goal_kpis(brand))} if registry is not None else {}
    rows = []
    for r in pubs.to_dict(orient="records"):
        end = to_date(r.get("period_end"))
        rows.append({
            "kpi_id": r["kpi_id"], "kpi_label": labels.get(r["kpi_id"], r["kpi_id"]),
            "scope": scope_text(r.get("scope_key")),
            "period": r.get("period_label") or (f"{end:%Y-%m}" if end else ""), "period_end": end,
            "value": r.get("value"), "unit": r.get("unit"), "status": r.get("display_label"),
            "validation_status": r.get("validation_status"),
            "platform_value": None if is_missing(r.get("platform_value")) else r.get("platform_value"),
            "variance_rel": None if is_missing(r.get("variance_rel")) else r.get("variance_rel"),
            "approved_by": r.get("approved_by"), "approved_at": r.get("approved_at"),
            "evidence": locator_text(files.get(str(r.get("evidence_hash"))), r.get("evidence_locator"),
                                     contexts.get(str(r["result_id"]))),
            "caveat": _caveat(r), "publication_id": r.get("publication_id"), "evidence_uri": r.get("evidence_uri"),
            "_order": order.get(r["kpi_id"], len(order)),
        })
    frame = pd.DataFrame(rows)
    frame["_end"] = [e or date.min for e in frame["period_end"]]
    frame = frame.sort_values(["_order", "kpi_label", "scope", "_end"], ascending=[True, True, True, False],
                              kind="stable")
    return frame.reset_index(drop=True)[list(SCORECARD_COLUMNS)]


def csv_text(value: Any) -> Any:
    """A text value a spreadsheet would read as a formula gets a leading apostrophe; other values are unchanged."""
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value
    return value


def scorecard_csv(store, registry, brand: str) -> bytes:
    """UTF-8 CSV (with BOM so Excel opens it cleanly) of the scorecard frame."""
    frame = scorecard_frame(store, registry, brand).copy()
    for column in frame.columns:
        dtype = frame[column].dtype
        if not (pd.api.types.is_numeric_dtype(dtype) or pd.api.types.is_datetime64_any_dtype(dtype)):
            frame[column] = frame[column].map(csv_text)
    return frame.to_csv(index=False).encode("utf-8-sig")


def _append_row(ws, values: list[Any]) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _number_format(unit: Any, value: Any) -> str:
    if unit == "fraction":
        return "0.0%"
    if unit == "usd":
        return '"$"#,##0'
    if not is_missing(value) and abs(float(value) - round(float(value))) > 1e-9:
        return "#,##0.0"
    return "#,##0"


def scorecard_xlsx(store, registry, brand: str) -> bytes:
    """Excel workbook: a formatted 'Scorecard' sheet (percent formats for fraction units) and a 'Notes' sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    frame = scorecard_frame(store, registry, brand)
    wb = Workbook()
    ws = wb.active
    ws.title = "Scorecard"
    headers = [DISPLAY_HEADERS[c] for c in _SHEET_COLUMNS]
    _append_row(ws, headers)
    fill = PatternFill("solid", fgColor="1F3864")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    col = {c: i + 1 for i, c in enumerate(_SHEET_COLUMNS)}
    for r in frame.to_dict(orient="records"):
        _append_row(ws, [_cell(r.get(c)) for c in _SHEET_COLUMNS])
        row = ws.max_row
        for name in ("value", "platform_value"):
            ws.cell(row=row, column=col[name]).number_format = _number_format(r.get("unit"), r.get(name))
        ws.cell(row=row, column=col["variance_rel"]).number_format = "+0.0%;-0.0%;0.0%"
        ws.cell(row=row, column=col["approved_at"]).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row=row, column=col["caveat"]).alignment = Alignment(wrap_text=True, vertical="top")
    widths = {"kpi_label": 30, "scope": 18, "period": 14, "value": 14, "unit": 9, "status": 20, "platform_value": 15,
              "variance_rel": 14, "approved_by": 28, "approved_at": 18, "evidence": 50, "caveat": 60, "kpi_id": 18,
              "publication_id": 34}
    for name, width in widths.items():
        ws.column_dimensions[get_column_letter(col[name])].width = width
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_SHEET_COLUMNS))}{max(ws.max_row, 1)}"
    notes = wb.create_sheet("Notes")
    _append_row(notes, ["Brand", brand])
    _append_row(notes, ["Generated at (UTC)", utcnow()])
    _append_row(notes, ["Published values", len(frame)])
    notes.append([])
    _append_row(notes, ["Status", "Meaning"])
    for label, meaning in (
            ("Matches platform", "The platform computes the same number with the governed definition."),
            ("Accepted variance", "The platform differs; a reviewer accepted the report value with a rationale."),
            ("Report-only", "No platform source holds this KPI; the value was captured from the report."),
            ("Corrected by reviewer", "A reviewer corrected the value read from the report.")):
        _append_row(notes, [label, meaning])
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(value: Any) -> Any:
    if is_missing(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def scorecard_pptx(store, registry, brand: str) -> bytes:
    """PowerPoint scorecard: one title slide plus one table slide per 12 published values."""
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("PowerPoint export needs python-pptx (pip install python-pptx).") from exc

    frame = scorecard_frame(store, registry, brand)
    display = registry.brands.get(brand, {}).get("display_name", brand) if registry is not None else brand
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = f"{display} KPI scorecard"
    subtitle = (f"{len(frame)} published value(s) - generated {utcnow():%Y-%m-%d %H:%M} UTC" if len(frame)
                else f"No approved values have been published for {display} yet.")
    title.placeholders[1].text = subtitle
    headers = ["KPI", "Scope", "Period", "Value", "Status", "Platform", "Difference", "Source"]
    widths = [2.6, 1.3, 1.2, 1.1, 1.6, 1.1, 1.0, 3.433]
    for start in range(0, len(frame), _ROWS_PER_SLIDE):
        chunk = frame.iloc[start:start + _ROWS_PER_SLIDE]
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        page = start // _ROWS_PER_SLIDE + 1
        pages = (len(frame) + _ROWS_PER_SLIDE - 1) // _ROWS_PER_SLIDE
        slide.shapes.title.text = f"{display} KPIs ({page} of {pages})"
        shape = slide.shapes.add_table(len(chunk) + 1, len(headers), Inches(0.3), Inches(1.4),
                                       Inches(sum(widths)), Inches(0.4 * (len(chunk) + 1)))
        table = shape.table
        for i, w in enumerate(widths):
            table.columns[i].width = Inches(w)
        for j, h in enumerate(headers):
            _set_text(table.cell(0, j), h, Pt(11), bold=True)
        for i, r in enumerate(chunk.to_dict(orient="records"), start=1):
            variance = r.get("variance_rel")
            cells = [r["kpi_label"], r["scope"], r["period"], format_value(r["value"], r["unit"]), r["status"] or "",
                     format_value(r.get("platform_value"), r["unit"]),
                     f"{float(variance) * 100:+.1f}%" if not is_missing(variance) else "", r["evidence"]]
            for j, text in enumerate(cells):
                _set_text(table.cell(i, j), str(text), Pt(10))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _set_text(cell, text: str, size, bold: bool = False) -> None:
    cell.text = text
    for paragraph in cell.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = size
            run.font.bold = bold
