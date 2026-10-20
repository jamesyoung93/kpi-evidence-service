from __future__ import annotations

import copy
import dataclasses
import math
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd

from .contracts import RegistryStatus
from .registry import DIMENSIONS, GRAINS, UNITS, KpiDefinition, Registry, lint, normalize_label, parse_kpi

COLUMNS = ("kpi", "aliases", "table", "value_sql", "numerator", "denominator", "date_column", "date_format",
           "filters", "brand_column", "scope_columns", "unit", "grain", "tolerance", "alternatives", "brand",
           "priority", "owner", "notes")

COLUMN_SYNONYMS: dict[str, str] = {
    "name": "kpi", "kpiname": "kpi", "kpiid": "kpi", "metric": "kpi", "metricname": "kpi",
    "alias": "aliases", "labels": "aliases", "label": "aliases",
    "tables": "table", "sourcetable": "table", "sourcetables": "table", "candidatetables": "table",
    "sql": "value_sql", "expression": "value_sql", "measure": "value_sql", "value": "value_sql",
    "valueexpression": "value_sql", "valuesql": "value_sql",
    "date": "date_column", "periodcolumn": "date_column", "datecol": "date_column", "monthcolumn": "date_column",
    "dateformat": "date_format", "periodformat": "date_format",
    "where": "filters", "filter": "filters", "predicate": "filters",
    "brandcolumn": "brand_column", "scopecolumns": "scope_columns", "scope": "scope_columns",
    "grains": "grain", "periods": "grain", "alternative": "alternatives", "tol": "tolerance",
    "note": "notes", "comment": "notes", "comments": "notes",
    "prio": "priority", "rank": "priority", "order": "priority",
}

FORBIDDEN_KEYWORDS = frozenset({
    "INSERT", "UPDATE", "DELETE", "MERGE", "UPSERT", "DROP", "CREATE", "ALTER", "TRUNCATE", "GRANT", "REVOKE",
    "DENY", "COPY", "CALL", "EXEC", "EXECUTE", "ATTACH", "DETACH", "INSTALL", "LOAD", "PRAGMA", "VACUUM",
    "OPTIMIZE", "REFRESH", "MSCK", "UNDROP", "RESTORE", "COMMIT", "ROLLBACK", "SET", "USE", "CACHE", "UNCACHE",
    "EXPORT", "IMPORT",
})

DATE_FORMATS = ("date", "yyyy-MM", "yyyyMM")
_DATE_FORMAT_SPELLINGS = {
    "": "date", "date": "date", "timestamp": "date", "datetime": "date", "yyyy-mm-dd": "date", "%y-%m-%d": "date",
    "yyyy-mm": "yyyy-MM", "%y-%m": "yyyy-MM", "month": "yyyy-MM",
    "yyyymm": "yyyyMM", "%y%m": "yyyyMM",
}
_UNIT_SPELLINGS = {
    "count": "count", "counts": "count", "number": "count", "units": "count", "unit": "count", "#": "count",
    "usd": "usd", "$": "usd", "dollars": "usd", "dollar": "usd", "currency": "usd",
    "fraction": "fraction", "percent": "fraction", "percentage": "fraction", "pct": "fraction", "%": "fraction",
    "share": "fraction", "ratio": "fraction", "rate": "fraction",
}
_GRAIN_SPELLINGS = {
    "month": "month", "months": "month", "monthly": "month", "quarter": "quarter", "quarters": "quarter",
    "quarterly": "quarter", "qtr": "quarter", "r3m": "r3m", "r6m": "r6m", "r12m": "r12m", "ytd": "ytd",
}
DEFAULT_TOLERANCE = {"count": {"rel": 0.005, "abs": 0.0}, "usd": {"rel": 0.005, "abs": 0.0},
                     "fraction": {"rel": 0.0, "abs": 0.0005}}
GRID_BUDGET = 16
PRIMARY = "primary"
HINT_VERSION = 0

_TABLE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*(?:\.[A-Za-z_][A-Za-z0-9_]*){0,2}$")
_IDENT = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_KPI_ID = re.compile(r"^[a-z][a-z0-9_]*$")
_NUMBER = r"([0-9]*\.?[0-9]+)"
MESSAGE_COLUMNS = ("row", "kpi", "level", "column", "message")


@dataclass
class HintRow:
    """One KPI table row after parsing (``accepted`` is False when the row is skipped; see the messages)."""

    row: int
    kpi: str
    kpi_id: str | None
    aliases: list[str] = field(default_factory=list)
    tables: list[str] = field(default_factory=list)
    value_sql: str | None = None
    numerator: str | None = None
    denominator: str | None = None
    date_column: str | None = None
    date_format: str = "date"
    filters: str | None = None
    brand_column: str | None = None
    scope_columns: dict[str, str] = field(default_factory=dict)
    unit: str | None = None
    grains: list[str] = field(default_factory=list)
    tolerance: dict[str, float] | None = None
    alternatives: dict[str, str] = field(default_factory=dict)
    brand: str | None = None
    priority: int | None = None
    owner: str | None = None
    notes: str | None = None
    overrides: str | None = None
    accepted: bool = True

    @property
    def report_only(self) -> bool:
        return not self.tables

    def summary(self) -> str:
        """One plain line: what the hint computes and where (for reports and messages)."""
        if self.report_only:
            return "report only (no table)"
        if self.value_sql is not None:
            what = self.value_sql
        elif self.numerator is not None and self.denominator is not None:
            what = f"({self.numerator}) / ({self.denominator})"
        else:
            what = "(no valid expression)"
        text = f"{what} from {self.tables[0]}"
        if len(self.tables) > 1:
            text += f" (candidates: {', '.join(self.tables[1:])})"
        details = [f"date {self.date_column} ({self.date_format})"] if self.date_column else []
        if self.filters:
            details.append(f"where {self.filters}")
        if self.brand_column:
            details.append(f"brand column {self.brand_column}")
        if self.alternatives:
            details.append("alternatives " + ", ".join(self.alternatives))
        return text + (f"; {'; '.join(details)}" if details else "")


@dataclass
class HintSet:
    """Parsed KPI table: every row, the KPI definitions built from accepted rows, and row-level messages."""

    rows: list[HintRow]
    kpis: dict[str, KpiDefinition]
    messages: pd.DataFrame
    source: str = "kpi_table"

    @property
    def errors(self) -> pd.DataFrame:
        return self.messages[self.messages["level"] == "error"].reset_index(drop=True)

    @property
    def warnings(self) -> pd.DataFrame:
        return self.messages[self.messages["level"] == "warning"].reset_index(drop=True)

    @property
    def brands(self) -> list[str]:
        """Distinct brand ids named by accepted rows, in table order."""
        return list(dict.fromkeys(r.brand for r in self.rows if r.accepted and r.brand))

    @property
    def default_brand(self) -> str | None:
        """The table's brand when the accepted rows name exactly one, else None."""
        brands = self.brands
        return brands[0] if len(brands) == 1 else None

    def overview(self) -> pd.DataFrame:
        """One row per KPI table row: what it defines, whether it is used, and its messages."""
        out = []
        for r in self.rows:
            msgs = self.messages[self.messages["row"] == r.row]
            out.append({
                "row": r.row, "kpi": r.kpi, "kpi_id": r.kpi_id, "used": r.accepted,
                "check": "report only" if r.report_only else "platform",
                "tables": "; ".join(r.tables), "definition": r.summary(), "overrides_registry": r.overrides,
                "brand": r.brand, "priority": r.priority, "errors": int((msgs["level"] == "error").sum()),
                "warnings": int((msgs["level"] == "warning").sum()),
                "messages": " | ".join(f"{m.level}: {m.message}" for m in msgs.itertuples()),
            })
        cols = ["row", "kpi", "kpi_id", "used", "check", "tables", "definition", "overrides_registry", "brand",
                "priority", "errors", "warnings", "messages"]
        frame = pd.DataFrame(out, columns=cols)
        frame["priority"] = frame["priority"].astype("Int64")
        return frame


class _Messages:
    def __init__(self) -> None:
        self.items: list[dict[str, Any]] = []

    def add(self, row: int | None, kpi: str | None, level: str, column: str | None, message: str) -> None:
        self.items.append({"row": row, "kpi": kpi, "level": level, "column": column, "message": message})

    def frame(self) -> pd.DataFrame:
        frame = pd.DataFrame(self.items, columns=list(MESSAGE_COLUMNS))
        frame["row"] = frame["row"].astype("Int64")
        return frame


def _column_key(name: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(name).strip().lower())


def canonical_column(name: Any) -> str | None:
    """The canonical KPI table column for a header (case/space/underscore insensitive), or None."""
    key = _column_key(name)
    for col in COLUMNS:
        if key == col.replace("_", ""):
            return col
    return COLUMN_SYNONYMS.get(key)


def normalize_columns(frame: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """Rename headers to canonical columns; coalesce duplicates (first non-blank wins)."""
    notes: list[str] = []
    groups: dict[str, list[Any]] = {}
    unknown: list[str] = []
    for col in frame.columns:
        canon = canonical_column(col)
        if canon is None:
            unknown.append(str(col))
            continue
        groups.setdefault(canon, []).append(col)
    out = pd.DataFrame(index=frame.index)
    for canon in [c for c in COLUMNS if c in groups]:
        sources = groups[canon]
        series = frame[sources[0]].astype(object)
        if len(sources) > 1:
            notes.append(f"columns {', '.join(map(str, sources))} all mean '{canon}'; the first non-blank value is used")
            for extra in sources[1:]:
                blank = series.map(lambda v: cell_text(v) is None)
                series = series.where(~blank, frame[extra].astype(object))
        out[canon] = series
    for col in unknown:
        notes.append(f"column '{col}' is not a KPI table column and is ignored")
    out.attrs.update(frame.attrs)
    out.attrs["unknown_columns"] = unknown
    out.attrs["column_notes"] = notes
    return out, notes


def _is_spark_frame(obj: Any) -> bool:
    return hasattr(obj, "toPandas") and hasattr(obj, "schema")


def _read_file(path: Path, sheet: str | int | None) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix in (".csv", ".txt"):
        return pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    if suffix in (".xlsx", ".xlsm"):
        return pd.read_excel(path, sheet_name=0 if sheet is None else sheet, dtype=str, keep_default_na=False)
    raise ValueError(f"unsupported KPI table file type {suffix!r} ({path.name}); use .csv or .xlsx")


def load_kpi_table(table: Any, *, sheet: str | int | None = None, spark: Any = None) -> pd.DataFrame:
    """Read a KPI table into a pandas DataFrame with canonical column names."""
    source = "kpi_table"
    if isinstance(table, pd.DataFrame):
        frame = table.copy()
    elif _is_spark_frame(table):
        frame = table.toPandas()
        source = "spark_dataframe"
    elif isinstance(table, (str, Path)):
        text = str(table).strip()
        if text.startswith("dbfs:/Volumes/"):
            text = text[len("dbfs:"):]
        path = Path(text)
        if path.suffix.lower() in (".csv", ".txt", ".xlsx", ".xlsm", ".xls"):
            if not path.exists():
                raise FileNotFoundError(f"KPI table not found: {text}")
            frame = _read_file(path, sheet)
            source = path.name
        elif _TABLE.match(text) and not path.exists():
            if spark is None:
                raise ValueError(f"{text!r} looks like a table name; pass spark= to read it, or give a .csv/.xlsx path")
            frame = spark.table(text).toPandas()
            source = text
        else:
            raise ValueError(f"unsupported KPI table {text!r}; give a .csv or .xlsx path, a DataFrame or a list of dicts")
    elif isinstance(table, Mapping):
        columnar = any(isinstance(v, (list, tuple, pd.Series)) for v in table.values())
        frame = pd.DataFrame(dict(table)) if columnar else pd.DataFrame([dict(table)])
    elif isinstance(table, Sequence) and all(isinstance(r, Mapping) for r in table):
        frame = pd.DataFrame.from_records([dict(r) for r in table])
    else:
        raise TypeError(f"unsupported KPI table input of type {type(table).__name__}")
    frame, _notes = normalize_columns(frame)
    if "kpi" not in frame.columns:
        found = ", ".join(map(str, frame.attrs.get("unknown_columns", []))) or "none"
        raise ValueError(f"the KPI table has no 'kpi' column (unrecognized columns: {found})")
    frame.attrs["kpi_table_source"] = source
    return frame.reset_index(drop=True)


def example_kpi_table_path() -> Path:
    """The bundled example KPI table (``samples/kpi_table_example.csv`` in the repo)."""
    from .settings import demo_dir

    return demo_dir("reports").parent / "kpi_table_example.csv"


def save_kpi_table(frame: pd.DataFrame, path: str | Path) -> str:
    """Write a KPI table to .csv or .xlsx (bold frozen header, filters, wrapped SQL columns)."""
    path = Path(path)
    frame, _ = normalize_columns(frame)
    frame = frame.reindex(columns=[c for c in COLUMNS if c in frame.columns] or list(COLUMNS))
    frame = pd.DataFrame({col: frame[col].map(lambda v: cell_text(v) or "") for col in frame.columns},
                         index=frame.index, columns=frame.columns)
    if path.suffix.lower() == ".csv":
        frame.to_csv(path, index=False, encoding="utf-8", lineterminator="\n")
        return str(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("save_kpi_table writes .csv or .xlsx")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "KPIs"
    ws.append(list(frame.columns))
    for values in frame.itertuples(index=False):
        ws.append(list(values))
    widths = {"kpi": 18, "aliases": 30, "table": 44, "value_sql": 46, "numerator": 24, "denominator": 26,
              "scope_columns": 40, "alternatives": 60, "notes": 60, "owner": 28, "filters": 20}
    for idx, col in enumerate(frame.columns, start=1):
        letter = ws.cell(row=1, column=idx).column_letter
        ws.column_dimensions[letter].width = widths.get(col, 14)
        ws.cell(row=1, column=idx).font = Font(bold=True)
        if col in ("value_sql", "numerator", "denominator", "alternatives", "notes", "scope_columns", "table"):
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=idx).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.properties.creator = "kpi_evidence"
    wb.save(path)
    return str(path)


def cell_text(value: Any) -> str | None:
    """Cell value as stripped text; None for blanks (None, NaN, NA, empty string)."""
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if not isinstance(value, (str, bytes, list, tuple, dict, set)):
        try:
            if pd.isna(value):
                return None
        except (TypeError, ValueError):
            pass
    text = str(value).strip()
    return text or None


def split_list(text: str | None, *, newlines: bool = True) -> list[str]:
    """Split a ``;``-separated (or, with ``newlines``, line-separated) cell into stripped, non-empty parts."""
    if not text:
        return []
    return [p.strip() for p in re.split(r"[;\n]" if newlines else ";", text) if p.strip()]


def kpi_slug(name: str) -> str:
    """lower_snake_case id for a KPI name ('TRx share' -> 'trx_share')."""
    s = normalize_label(name).replace("%", " pct ").replace("$", " usd ")
    s = re.sub(r"[^a-z0-9]+", "_", s).strip("_")
    if s and s[0].isdigit():
        s = "kpi_" + s
    return s


def brand_id_for(text: str, registry: Registry | None = None) -> str:
    """Brand id for a brand name: the registry brand it is an alias of, else an UPPER_SNAKE id."""
    if registry is not None:
        hit = registry.brand_aliases().get(normalize_label(text))
        if hit:
            return hit
    return re.sub(r"[^A-Z0-9]+", "_", text.upper()).strip("_")


def parse_unit(text: str | None) -> tuple[str | None, str | None]:
    """(canonical unit | None, error | None); percent units become 'fraction'."""
    if text is None:
        return None, None
    unit = _UNIT_SPELLINGS.get(text.strip().lower())
    if unit is None or unit not in UNITS:
        return None, f"unit {text!r} is not one of count, usd, percent / fraction"
    return unit, None


def default_tolerance(unit: str) -> dict[str, float]:
    return dict(DEFAULT_TOLERANCE.get(unit, DEFAULT_TOLERANCE["count"]))


def parse_tolerance(text: str | None, unit: str) -> tuple[dict[str, float] | None, list[tuple[str, str]]]:
    """Parse a tolerance cell."""
    if text is None:
        return None, []
    rel = abs_ = 0.0
    notes: list[tuple[str, str]] = []
    for part in [p for p in re.split(r"[;,]", text) if p.strip()]:
        p = part.strip().lower().replace(" ", "")
        if m := re.fullmatch(rf"(?:rel[:=])?{_NUMBER}%", p):
            rel = max(rel, float(m.group(1)) / 100)
        elif m := re.fullmatch(rf"{_NUMBER}pp", p):
            if unit != "fraction":
                notes.append(("warning", f"tolerance {part.strip()!r} is in percentage points but the unit is {unit}; "
                                         "default tolerance used"))
                return None, notes
            abs_ = max(abs_, float(m.group(1)) / 100)
        elif m := re.fullmatch(rf"abs[:=]{_NUMBER}", p):
            abs_ = max(abs_, float(m.group(1)))
        elif (m := re.fullmatch(rf"(?:rel[:=])?{_NUMBER}", p)) and float(m.group(1)) < 1:
            value = float(m.group(1))
            rel = max(rel, value)
            if value >= 0.05:
                notes.append(("warning", f"tolerance {part.strip()!r} is read as a relative {value * 100:g}%; "
                                         "write e.g. 0.5% to be explicit"))
        else:
            notes.append(("warning", f"tolerance {part.strip()!r} is not understood (use 0.5%, 0.05pp or abs:12); "
                                     "default tolerance used"))
            return None, notes
    return {"rel": rel, "abs": abs_}, notes


def parse_priority(text: str | None) -> tuple[int | None, str | None]:
    """(priority as a whole number >= 1 | None, problem | None); blank means lowest."""
    if text is None:
        return None, None
    try:
        number = float(text)
    except ValueError:
        return None, f"priority {text!r} is not a whole number (1 = highest); treated as blank"
    if not math.isfinite(number) or number < 1 or number != int(number):
        return None, f"priority {text!r} is not a whole number of 1 or more; treated as blank"
    return int(number), None


def parse_date_format(text: str | None) -> tuple[str | None, str | None]:
    """(one of DATE_FORMATS | None, error | None); blank means 'date'."""
    key = (text or "").strip().lower()
    fmt = _DATE_FORMAT_SPELLINGS.get(key)
    if fmt is None:
        return None, f"date_format {text!r} is not one of date, yyyy-MM, yyyyMM"
    return fmt, None


def date_expression(column: str, date_format: str) -> str:
    """SQL giving each row's date from ``column`` (the KPI's ``date_expr``)."""
    if date_format == "yyyy-MM":
        return f"CAST({column} || '-01' AS DATE)"
    if date_format == "yyyyMM":
        text = f"CAST({column} AS STRING)"
        return f"CAST(CONCAT(SUBSTR({text}, 1, 4), '-', SUBSTR({text}, 5, 2), '-01') AS DATE)"
    return column


def parse_grains(text: str | None) -> tuple[list[str], list[str]]:
    """(grains, unknown spellings) from a ``;``-separated cell."""
    grains, unknown = [], []
    for part in split_list(text.replace(",", ";") if text else None):
        g = _GRAIN_SPELLINGS.get(part.lower())
        if g is None:
            unknown.append(part)
        elif g not in grains:
            grains.append(g)
    return grains, unknown


def parse_pairs(text: str | None, *, newlines: bool = True) -> tuple[list[tuple[str, str]], list[str]]:
    """``a=x; b=y`` -> ([(a, x), (b, y)], malformed parts)."""
    pairs, bad = [], []
    for part in split_list(text, newlines=newlines):
        m = re.match(r"^\s*([^=:]+?)\s*[=:]\s*(.+?)\s*$", part, re.S)
        if m:
            pairs.append((m.group(1), m.group(2)))
        else:
            bad.append(part)
    return pairs, bad


def _strip_quoted(text: str) -> tuple[str, bool, bool]:
    out, i, n, escaped = [], 0, len(text), False
    while i < n:
        ch = text[i]
        if ch in "'\"`":
            j = i + 1
            while j < n:
                if text[j] == ch:
                    if j + 1 < n and text[j + 1] == ch:
                        escaped = escaped or ch != "`"
                        j += 2
                        continue
                    break
                j += 1
            if j >= n:
                return "".join(out), False, escaped
            out.append(ch + ch)
            i = j + 1
            continue
        out.append(ch)
        i += 1
    return "".join(out), True, escaped


def sql_fragment_problems(text: str) -> list[str]:
    """Why a user SQL fragment is rejected ([] = acceptable): ';', comments, braces, keywords, escaped quotes."""
    problems = []
    if ";" in text:
        problems.append("';' is not allowed (one expression, not statements)")
    if "--" in text or "/*" in text or "*/" in text:
        problems.append("SQL comments are not allowed")
    if "{" in text or "}" in text:
        problems.append("braces are not allowed (they are template placeholders)")
    if "\\" in text:
        problems.append("backslashes are not allowed (Spark and DuckDB read them differently inside strings)")
    stripped, closed, escaped = _strip_quoted(text)
    if not closed:
        problems.append("a quote is not closed")
    if escaped:
        problems.append("a quote inside a string is not allowed (Spark and DuckDB read doubled quotes differently)")
    words = {w.upper() for w in re.findall(r"[A-Za-z_][A-Za-z0-9_]*", stripped)}
    bad = sorted(words & FORBIDDEN_KEYWORDS)
    if bad:
        problems.append(f"statement keywords are not allowed: {', '.join(bad)}")
    return problems


class _RowParser:
    def __init__(self, number: int, record: dict[str, Any], messages: _Messages, base: Registry | None):
        self.number, self.record, self.messages, self.base = number, record, messages, base
        self.failed = False
        self.kpi_name = cell_text(record.get("kpi")) or ""

    def text(self, column: str) -> str | None:
        return cell_text(self.record.get(column))

    def error(self, column: str | None, message: str) -> None:
        self.failed = True
        self.messages.add(self.number, self.kpi_name or None, "error", column, message)

    def warn(self, column: str | None, message: str) -> None:
        self.messages.add(self.number, self.kpi_name or None, "warning", column, message)

    def sql(self, column: str, text: str | None) -> str | None:
        if text is None:
            return None
        problems = sql_fragment_problems(text)
        for p in problems:
            self.error(column, f"{column}: {p}")
        return None if problems else text

    def ident(self, column: str, text: str | None) -> str | None:
        if text is None:
            return None
        if not _IDENT.match(text):
            self.error(column, f"{column} {text!r} must be a plain column name (letters, digits, underscore)")
            return None
        return text

    def resolve_id(self) -> tuple[str | None, str | None]:
        """(kpi_id, overridden base KPI id)."""
        slug = kpi_slug(self.kpi_name)
        if not slug or not _KPI_ID.match(slug):
            self.error("kpi", f"kpi {self.kpi_name!r} does not give a usable id (lower_snake_case)")
            return None, None
        if self.base is None:
            return slug, None
        if slug in self.base.kpis:
            return slug, slug
        name = normalize_label(self.kpi_name)
        hits = sorted({k.kpi_id for k in self.base.kpis.values()
                       if name in {normalize_label(s) for s in (k.label, *k.synonyms)}})
        if len(hits) == 1:
            self.warn("kpi", f"'{self.kpi_name}' is registry KPI {hits[0]} (matched by name); this row redefines it. "
                             f"Write kpi = {hits[0]} to be explicit")
            return hits[0], hits[0]
        return slug, None

    def tables(self) -> list[str]:
        out = []
        for t in split_list(self.text("table")):
            if not _TABLE.match(t):
                self.error("table", f"table {t!r} must be a dotted identifier such as catalog.schema.table")
                continue
            if "." not in t:
                self.warn("table", f"table {t!r} has no schema; it resolves against the session's current schema")
            if t in out:
                self.warn("table", f"table {t!r} is listed twice; the repeat is ignored")
                continue
            out.append(t)
        return out

    def scope_columns(self) -> dict[str, str]:
        pairs, bad = parse_pairs(self.text("scope_columns"))
        out: dict[str, str] = {}
        for part in bad:
            pairs.append((part, part))
        for dim, col in pairs:
            key = re.sub(r"[\s-]+", "_", dim.strip().lower())
            if key not in DIMENSIONS:
                self.warn("scope_columns", f"scope dimension {dim!r} is not one of {', '.join(DIMENSIONS)}; ignored")
                continue
            if not _IDENT.match(col):
                self.error("scope_columns", f"scope column {col!r} for {key} must be a plain column name")
                continue
            out[key] = col
        return out

    def alternatives(self) -> dict[str, str]:
        pairs, bad = parse_pairs(self.text("alternatives"), newlines=False)
        for part in bad:
            self.error("alternatives", f"alternative {part!r} must be written name=EXPRESSION")
        out: dict[str, str] = {}
        for name, expr in pairs:
            key = kpi_slug(name)
            if not key or key == PRIMARY:
                self.error("alternatives", f"alternative name {name!r} is not usable ({PRIMARY!r} is reserved)")
                continue
            if key in out:
                self.error("alternatives", f"alternative {key!r} is named twice")
                continue
            sql = self.sql("alternatives", expr)
            if sql is not None:
                out[key] = sql
        return out

    def parse(self) -> HintRow:
        kpi_id, overrides = self.resolve_id()
        row = HintRow(row=self.number, kpi=self.kpi_name, kpi_id=kpi_id, overrides=overrides)
        row.aliases = split_list(self.text("aliases"))
        row.tables = self.tables()
        row.value_sql = self.sql("value_sql", self.text("value_sql"))
        row.numerator = self.sql("numerator", self.text("numerator"))
        row.denominator = self.sql("denominator", self.text("denominator"))
        row.filters = self.sql("filters", self.text("filters"))
        row.date_column = self.ident("date_column", self.text("date_column"))
        row.brand_column = self.ident("brand_column", self.text("brand_column"))
        row.scope_columns = self.scope_columns()
        fmt, err = parse_date_format(self.text("date_format"))
        if err:
            self.error("date_format", err)
        row.date_format = fmt or "date"
        row.unit, err = parse_unit(self.text("unit"))
        if err:
            self.error("unit", err)
        row.grains, unknown = parse_grains(self.text("grain"))
        for g in unknown:
            self.warn("grain", f"grain {g!r} is not one of {', '.join(GRAINS)}; ignored")
        row.alternatives = self.alternatives()
        brand = self.text("brand")
        row.brand = brand_id_for(brand, self.base) if brand else None
        row.owner, row.notes = self.text("owner"), self.text("notes")
        row.priority, problem = parse_priority(self.text("priority"))
        if problem:
            self.warn("priority", problem)
        self.check_structure(row)
        unit = row.unit or self.base_unit(row) or "count"
        row.tolerance, notes = parse_tolerance(self.text("tolerance"), unit)
        for level, message in notes:
            (self.error if level == "error" else self.warn)("tolerance", message)
        row.accepted = not self.failed
        return row

    def base_unit(self, row: HintRow) -> str | None:
        if self.base is not None and row.overrides:
            return self.base.kpis[row.overrides].unit
        return None

    def check_structure(self, row: HintRow) -> None:
        has_value = self.text("value_sql") is not None
        has_num, has_den = self.text("numerator") is not None, self.text("denominator") is not None
        if not row.tables:
            if self.text("table") is None and (has_value or has_num or has_den):
                self.error("table", "an expression is given but no table; add the table, or clear the expressions "
                                    "for a report-only KPI")
            return
        if has_value and (has_num or has_den):
            self.error("value_sql", "give value_sql or numerator + denominator, not both")
        elif has_num != has_den:
            self.error("numerator" if not has_num else "denominator",
                       "a ratio needs both numerator and denominator")
        elif not (has_value or has_num):
            self.error("value_sql", "a table is given but no value_sql (or numerator + denominator)")
        if self.text("date_column") is None:
            self.error("date_column", "date_column is required for a platform check (the column with each row's date)")
        if row.alternatives and not has_value:
            self.warn("alternatives", "alternatives work with value_sql only; write the ratio as value_sql "
                                      "(e.g. SUM(a) / NULLIF(SUM(b), 0)) to test them. Ignored")
            row.alternatives = {}
        combos = max(len(row.tables), 1) * (len(row.alternatives) + 1)
        if combos > GRID_BUDGET:
            self.warn("alternatives", f"{combos} candidate definitions exceed the grid budget of {GRID_BUDGET}; "
                                      "only the judge search runs (no full grid)")


def _option_names(n: int) -> list[str]:
    return [PRIMARY] + [f"alt_{i}" for i in range(1, n)]


def _describe(expr: str, limit: int = 120) -> str:
    text = " ".join(expr.split())
    return text if len(text) <= limit else text[:limit - 3] + "..."


def platform_block(row: HintRow) -> dict[str, Any] | str:
    """The registry shorthand platform block for a hint row ('none' for a report-only row)."""
    if row.report_only:
        return "none"
    params: dict[str, Any] = {}
    if len(row.tables) == 1:
        block: dict[str, Any] = {"from": row.tables[0]}
    else:
        names = _option_names(len(row.tables))
        block = {"view": "{source}"}
        params["source"] = {"hypothesis": "source", "default": PRIMARY, "options": dict(zip(names, row.tables)),
                            "describe": dict(zip(names, row.tables)), "permitted": [PRIMARY]}
    if row.value_sql is not None:
        if row.alternatives:
            options = {PRIMARY: row.value_sql, **row.alternatives}
            params["measure"] = {
                "hypothesis": "measure", "default": PRIMARY, "options": options,
                "describe": {name: _describe(expr) for name, expr in options.items()},
                "keywords": {name: [name.replace("_", " ")] for name in row.alternatives},
                "permitted": [PRIMARY]}
            block["value"] = "{measure}"
        else:
            block["value"] = row.value_sql
    else:
        block["numerator"], block["denominator"] = row.numerator, row.denominator
    if row.filters:
        block["where"] = f"({row.filters})"
    block["brand_column"] = row.brand_column
    block["date_expr"] = date_expression(row.date_column, row.date_format)
    block["scope_columns"] = dict(row.scope_columns)
    block["params"] = params
    block["source_views"] = list(row.tables)
    return block


def _aggregation(row: HintRow, unit: str) -> str:
    if row.numerator is not None:
        return "ratio_of_sums"
    expr = (row.value_sql or "").strip().upper()
    if unit == "fraction":
        return "ratio_of_sums" if "/" in expr else "point"
    return "count" if expr.startswith("COUNT(") else "sum"


def _synonyms(*groups: Sequence[str]) -> list[str]:
    seen: dict[str, str] = {}
    for group in groups:
        for s in group:
            key = normalize_label(s)
            if key and key not in seen:
                seen[key] = s
    return list(seen.values())


def raw_definition(row: HintRow, base: KpiDefinition | None, source: str) -> dict[str, Any]:
    """The registry-style raw dict for a hint row (DRAFT, ``definition_source = 'hint'``)."""
    unit = row.unit or (base.unit if base else "count")
    if row.tolerance is not None:
        tolerance = {**row.tolerance, "use_display_rounding": True}
    elif base is not None and row.unit in (None, base.unit):
        tolerance = {"rel": base.tolerance.rel, "abs": base.tolerance.abs,
                     "use_display_rounding": base.tolerance.use_display_rounding}
    else:
        tolerance = {**default_tolerance(unit), "use_display_rounding": True}
    summary = row.summary()
    hint_fields = {k: v for k, v in dataclasses.asdict(row).items()
                   if k not in ("accepted", "row") and v not in (None, [], {})}
    if base is not None:
        dims = list(dict.fromkeys([*base.dimensions, *row.scope_columns]))
        raw = {**copy.deepcopy(base.raw)}
        raw.update({"label": base.label, "definition": base.definition, "owner": row.owner or base.owner,
                    "synonyms": _synonyms(base.synonyms, row.aliases, [row.kpi]), "dimensions": dims,
                    "grains": row.grains or list(base.grains), "harness_link": base.harness_link})
    else:
        definition = f"KPI table hint: {summary}."
        if row.notes:
            definition += f" {row.notes}"
        raw = {"kpi_id": row.kpi_id, "label": row.kpi, "definition": definition,
               "aggregation": _aggregation(row, unit), "claim_class": "descriptive", "owner": row.owner or "",
               "synonyms": _synonyms(row.aliases, [row.kpi]), "dimensions": list(row.scope_columns),
               "grains": row.grains or list(GRAINS), "harness_link": None,
               "mapping_description": f"{row.kpi} (from the KPI table)"}
    raw.update({"kpi_id": row.kpi_id, "version": HINT_VERSION, "status": RegistryStatus.DRAFT.value,
                "approved_by": None, "approved_at": None,
                "unit": unit, "tolerance": tolerance, "platform": platform_block(row),
                "definition_source": "hint", "hint_row": row.row, "hint_source": source, "hint_summary": summary,
                "hint_brand": row.brand, "hint_notes": row.notes, "hint": hint_fields, "priority": row.priority})
    return raw


def _lint_one(kpi: KpiDefinition) -> list[str]:
    solo = Registry(kpis={kpi.kpi_id: kpi}, brands={}, report_types={}, goal_list=[], policy={}, directory=Path("."))
    return lint(solo)


def is_hint(kpi: KpiDefinition | None) -> bool:
    """True for a KPI defined by a KPI table row."""
    return kpi is not None and (kpi.raw or {}).get("definition_source") == "hint"


def definition_source(kpi: KpiDefinition | None) -> str | None:
    """'hint' or 'registry' (None when there is no KPI)."""
    if kpi is None:
        return None
    return "hint" if is_hint(kpi) else "registry"


def _choose_rows(rows: list[HintRow], messages: _Messages, default_brand: str | None) -> None:
    by_id: dict[str, list[HintRow]] = {}
    for r in rows:
        if r.accepted and r.kpi_id:
            by_id.setdefault(r.kpi_id, []).append(r)
    for kpi_id, group in by_id.items():
        if default_brand:
            for r in group:
                if r.brand and r.brand != default_brand:
                    r.accepted = False
                    messages.add(r.row, r.kpi, "warning", "brand",
                                 f"skipped: the row is for brand {r.brand}; this check runs for {default_brand}")
            group = [r for r in group if r.accepted]
            specific = [r for r in group if r.brand == default_brand]
            if specific:
                for r in group:
                    if not r.brand:
                        r.accepted = False
                        messages.add(r.row, r.kpi, "warning", "brand",
                                     f"skipped: row {specific[0].row} defines {kpi_id} for brand {default_brand}")
                group = [r for r in group if r.accepted]
        if len(group) <= 1:
            continue
        keep = group[0]
        brands = sorted({r.brand for r in group if r.brand})
        for r in group[1:]:
            r.accepted = False
            if len(brands) > 1:
                messages.add(r.row, r.kpi, "warning", "brand",
                             f"skipped: {kpi_id} has rows for several brands ({', '.join(brands)}); row {keep.row} "
                             "is used. Pass brand= to choose")
            else:
                messages.add(r.row, r.kpi, "error", "kpi", f"skipped: duplicate row for {kpi_id} (row {keep.row} is used)")


def _claim_labels(rows: list[HintRow], kpis: dict[str, KpiDefinition], messages: _Messages,
                  base: Registry | None) -> dict[str, KpiDefinition]:
    owner: dict[str, str] = {}
    out: dict[str, KpiDefinition] = {}
    by_id = {r.kpi_id: r for r in rows if r.accepted}
    hint_ids = set(kpis)
    base_labels: dict[str, str] = {}
    if base is not None:
        for k in base.kpis.values():
            if k.kpi_id not in hint_ids:
                for s in k.synonyms:
                    base_labels.setdefault(normalize_label(s), k.kpi_id)
    for kpi_id, kpi in kpis.items():
        row = by_id[kpi_id]
        keep = []
        for s in kpi.synonyms:
            key = normalize_label(s)
            if key in owner and owner[key] != kpi_id:
                messages.add(row.row, row.kpi, "warning", "aliases",
                             f"label '{s}' already belongs to {owner[key]}; it is not used for {kpi_id}")
                continue
            owner[key] = kpi_id
            keep.append(s)
            if key in base_labels:
                messages.add(row.row, row.kpi, "warning", "aliases",
                             f"label '{s}' is also a synonym of registry KPI {base_labels[key]}; the KPI table takes it over")
        if tuple(keep) != kpi.synonyms:
            kpi = dataclasses.replace(kpi, synonyms=tuple(keep), raw={**kpi.raw, "synonyms": keep})
        out[kpi_id] = kpi
    return out


def parse_hints(frame: Any, *, base_registry: Registry | None = None, default_brand: str | None = None) -> HintSet:
    """Validate a KPI table and build one DRAFT ``KpiDefinition`` per accepted row."""
    if not isinstance(frame, pd.DataFrame):
        frame = load_kpi_table(frame)
    source = str(frame.attrs.get("kpi_table_source", "kpi_table"))
    earlier = list(frame.attrs.get("column_notes", []))
    frame, notes = normalize_columns(frame)
    messages = _Messages()
    for note in dict.fromkeys(earlier + notes):
        messages.add(None, None, "warning", None, note)
    if "kpi" not in frame.columns:
        messages.add(None, None, "error", "kpi", "the KPI table has no 'kpi' column")
        return HintSet(rows=[], kpis={}, messages=messages.frame(), source=source)
    if default_brand:
        default_brand = brand_id_for(default_brand, base_registry)
    rows: list[HintRow] = []
    for number, record in enumerate(frame.to_dict(orient="records"), start=1):
        if all(cell_text(v) is None for v in record.values()):
            continue
        parser = _RowParser(number, record, messages, base_registry)
        if not parser.kpi_name:
            parser.error("kpi", "kpi is blank; the row is skipped")
            rows.append(HintRow(row=number, kpi="", kpi_id=None, accepted=False))
            continue
        rows.append(parser.parse())
    if not rows:
        messages.add(None, None, "warning", None, "the KPI table has no rows")
    _choose_rows(rows, messages, default_brand)
    kpis: dict[str, KpiDefinition] = {}
    for row in rows:
        if not row.accepted:
            continue
        base = base_registry.kpis.get(row.overrides) if base_registry is not None and row.overrides else None
        try:
            kpi = parse_kpi(raw_definition(row, base, source), source_file=f"{source}#row{row.row}")
            problems = _lint_one(kpi)
        except Exception as exc:
            problems = [f"{type(exc).__name__}: {exc}"]
        if problems:
            row.accepted = False
            for p in problems:
                messages.add(row.row, row.kpi, "error", None, f"skipped: {p}")
            continue
        kpis[row.kpi_id] = kpi
    kpis = _claim_labels(rows, kpis, messages, base_registry)
    return HintSet(rows=rows, kpis=kpis, messages=messages.frame(), source=source)


@dataclass
class CheckRegistry(Registry):
    """A registry for one check: the base KPIs plus the KPI table's DRAFT hint definitions."""

    def approved(self) -> list[KpiDefinition]:
        return [k for k in self.kpis.values() if k.status == RegistryStatus.APPROVED or is_hint(k)]


def merge_registry(base: Registry | None, hints: HintSet) -> Registry:
    """Base registry (or an empty one) plus the hint KPIs; a hint wins on id and takes over its labels."""
    kpis: dict[str, KpiDefinition] = dict(base.kpis) if base is not None else {}
    claimed = {normalize_label(s) for k in hints.kpis.values() for s in k.synonyms}
    for kpi_id, kpi in list(kpis.items()):
        if kpi_id in hints.kpis:
            continue
        keep = tuple(s for s in kpi.synonyms if normalize_label(s) not in claimed)
        if keep != kpi.synonyms:
            kpis[kpi_id] = dataclasses.replace(kpi, synonyms=keep, raw={**kpi.raw, "synonyms": list(keep)})
    kpis.update(hints.kpis)
    brands = copy.deepcopy(base.brands) if base is not None else {}
    report_types = copy.deepcopy(base.report_types) if base is not None else {}
    for spec in report_types.values():
        assumed = spec.get("assumed_definitions")
        if isinstance(assumed, dict):
            spec["assumed_definitions"] = {k: v for k, v in assumed.items() if k not in hints.kpis}
    probe = Registry(kpis={}, brands=brands, report_types={}, goal_list=[], policy={}, directory=Path("."))
    for brand in hints.brands:
        aliases = probe.brand_aliases()
        if brand not in brands and normalize_label(brand) not in aliases:
            brands[brand] = {"brand_id": brand, "display_name": brand.replace("_", " ").title(),
                             "aliases": [brand.lower().replace("_", " ")], "fiscal_year_start_month": 1}
    return CheckRegistry(
        kpis=kpis, brands=brands, report_types=report_types,
        goal_list=list(copy.deepcopy(base.goal_list)) if base is not None else [],
        policy=copy.deepcopy(base.policy) if base is not None else {},
        directory=base.directory if base is not None else Path(hints.source),
        param_groups=copy.deepcopy(base.param_groups) if base is not None else {})


def lint_for_hints(registry: Registry) -> list[str]:
    """``registry.lint`` for a merged registry; DRAFT hint KPIs need no approver ([] = consistent)."""
    return [p for p in lint(registry) if not p.endswith("APPROVED needs approved_by")]


def scope_values_sql(hints: HintSet) -> list[tuple[str, str]]:
    """``[(dimension, sql)]``: ``SELECT DISTINCT <column> AS v FROM <primary table> LIMIT 1000`` per scope column."""
    out: set[tuple[str, str]] = set()
    for row in hints.rows:
        if not row.accepted or row.report_only:
            continue
        for dim, col in row.scope_columns.items():
            out.add((dim, f"SELECT DISTINCT {col} AS v FROM {row.tables[0]} LIMIT 1000"))
    return sorted(out)
