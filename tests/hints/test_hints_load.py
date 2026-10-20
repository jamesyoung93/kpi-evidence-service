from __future__ import annotations

import pandas as pd
import pytest
from hints_testkit import EXAMPLE_CSV, EXAMPLE_XLSX, row

from kpi_evidence.hints import (
    COLUMNS, canonical_column, cell_text, load_kpi_table, normalize_columns, parse_hints, save_kpi_table,
)


class FakeSparkFrame:
    def __init__(self, frame: pd.DataFrame):
        self._frame = frame
        self.schema = list(frame.columns)

    def toPandas(self) -> pd.DataFrame:
        return self._frame.copy()


class FakeSpark:
    def __init__(self, tables: dict[str, pd.DataFrame]):
        self.tables = tables
        self.calls: list[str] = []

    def table(self, name: str) -> FakeSparkFrame:
        self.calls.append(name)
        return FakeSparkFrame(self.tables[name])


def _clean(frame: pd.DataFrame) -> pd.DataFrame:
    return frame.apply(lambda col: col.map(lambda v: cell_text(v) or ""))


def test_example_csv_and_xlsx_load_to_the_same_table():
    csv, xlsx = load_kpi_table(EXAMPLE_CSV), load_kpi_table(EXAMPLE_XLSX)
    assert list(csv.columns) == list(COLUMNS)
    pd.testing.assert_frame_equal(_clean(csv), _clean(xlsx))
    assert csv.attrs["kpi_table_source"] == "kpi_table_example.csv"


@pytest.mark.parametrize("header, canonical", [
    ("KPI", "kpi"), ("Name", "kpi"), ("kpi name", "kpi"), ("Value SQL", "value_sql"), ("SQL", "value_sql"),
    ("expression", "value_sql"), ("Measure", "value_sql"), ("Date", "date_column"), ("period_column", "date_column"),
    ("Date-Column", "date_column"), ("WHERE", "filters"), ("filter", "filters"), ("Scope Columns", "scope_columns"),
    ("Tables", "table"), ("brand column", "brand_column"), ("Grains", "grain"), ("unknown thing", None),
])
def test_canonical_column_names(header, canonical):
    assert canonical_column(header) == canonical


def test_synonym_headers_load_and_parse():
    frame = pd.DataFrame([{"Name": "TRx", "Tables": "platform_demo.rx_account_monthly", "SQL": "SUM(trx_units)",
                           "Period Column": "month", "Date Format": "yyyy-MM", "Where": "trx_units > 0",
                           "Shade": "blue"}])
    loaded = load_kpi_table(frame)
    assert set(loaded.columns) == {"kpi", "table", "value_sql", "date_column", "date_format", "filters"}
    hints = parse_hints(loaded)
    assert hints.kpis["trx"].platform.template.endswith("AND (trx_units > 0)")
    assert any("'Shade' is not a KPI table column" in m for m in hints.messages["message"])


def test_duplicate_columns_coalesce_first_non_blank():
    frame = pd.DataFrame([{"kpi": "a", "sql": "", "value_sql": "SUM(x)"}, {"kpi": "b", "sql": "SUM(y)", "value_sql": "SUM(z)"}])
    out, notes = normalize_columns(frame)
    assert list(out["value_sql"]) == ["SUM(x)", "SUM(y)"]
    assert any("all mean 'value_sql'" in n for n in notes)


def test_list_of_dicts_dict_of_lists_and_dataframe_copy():
    records = [row(), row(kpi="nrx", value_sql="SUM(nrx)")]
    assert list(load_kpi_table(records)["kpi"]) == ["trx", "nrx"]
    assert list(load_kpi_table({"kpi": ["a"], "unit": ["count"]})["kpi"]) == ["a"]
    assert list(load_kpi_table(row())["kpi"]) == ["trx"]
    original = pd.DataFrame(records)
    loaded = load_kpi_table(original)
    loaded.loc[:, "kpi"] = "changed"
    assert list(original["kpi"]) == ["trx", "nrx"]


def test_spark_dataframe_and_table_name():
    frame = pd.DataFrame([row()])
    assert list(load_kpi_table(FakeSparkFrame(frame))["kpi"]) == ["trx"]
    spark = FakeSpark({"main.brand.kpi_hints": frame})
    loaded = load_kpi_table("main.brand.kpi_hints", spark=spark)
    assert spark.calls == ["main.brand.kpi_hints"] and loaded.attrs["kpi_table_source"] == "main.brand.kpi_hints"
    with pytest.raises(ValueError, match="pass spark="):
        load_kpi_table("main.brand.kpi_hints")


def test_file_errors(tmp_path):
    with pytest.raises(FileNotFoundError, match="/Volumes/main/x/kpi.csv"):
        load_kpi_table("dbfs:/Volumes/main/x/kpi.csv")
    with pytest.raises(FileNotFoundError):
        load_kpi_table(tmp_path / "missing.xlsx")
    (tmp_path / "old.xls").write_bytes(b"not really")
    with pytest.raises(ValueError, match="use .csv or .xlsx"):
        load_kpi_table(tmp_path / "old.xls")
    (tmp_path / "nokpi.csv").write_text("table,value_sql\nx.y,SUM(a)\n", encoding="utf-8")
    with pytest.raises(ValueError, match="no 'kpi' column"):
        load_kpi_table(tmp_path / "nokpi.csv")
    with pytest.raises(TypeError):
        load_kpi_table(42)


def test_csv_keeps_text_like_na_and_leading_zeros(tmp_path):
    path = tmp_path / "t.csv"
    path.write_text("kpi,aliases,table,value_sql,date_column,filters\n"
                    "trx,NA,platform_demo.rx_account_monthly,SUM(trx_units),month,territory = '007'\n", encoding="utf-8")
    frame = load_kpi_table(path)
    assert frame.loc[0, "aliases"] == "NA" and frame.loc[0, "filters"] == "territory = '007'"


def test_excel_sheet_selection_and_percent_cells(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Readme"
    wb.active.append(["This sheet explains the table"])
    ws = wb.create_sheet("KPIs")
    ws.append(["KPI", "Table", "Value SQL", "Date column", "Date format", "Tolerance"])
    ws.append(["trx", "platform_demo.rx_account_monthly", "SUM(trx_units)", "month", "yyyy-MM", 0.01])
    ws.cell(row=2, column=6).number_format = "0.0%"
    path = tmp_path / "kpis.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="no 'kpi' column"):
        load_kpi_table(path)
    hints = parse_hints(load_kpi_table(path, sheet="KPIs"))
    assert hints.kpis["trx"].tolerance.rel == pytest.approx(0.01)


def test_save_kpi_table_round_trip(tmp_path):
    frame = load_kpi_table(EXAMPLE_CSV)
    for name in ("copy.csv", "copy.xlsx"):
        out = save_kpi_table(frame, tmp_path / name)
        pd.testing.assert_frame_equal(_clean(load_kpi_table(out)), _clean(frame))
    with pytest.raises(ValueError):
        save_kpi_table(frame, tmp_path / "copy.json")
