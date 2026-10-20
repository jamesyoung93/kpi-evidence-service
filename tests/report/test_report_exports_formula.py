from __future__ import annotations

import io
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from kpi_evidence import exports
from kpi_evidence.exports import SCORECARD_COLUMNS, csv_text, locator_text


@pytest.fixture()
def formula_frame(monkeypatch):
    row = {c: None for c in SCORECARD_COLUMNS}
    row.update(kpi_id="trx", kpi_label="+TRx", scope="Brand total", period="May 2026", period_end=date(2026, 5, 31),
               value=5910.0, unit="count", status="Matches platform", platform_value=-12.5,
               evidence=locator_text("=1+1.pptx", {"slide": 2}, {"chart_title": "TRx by month"}),
               caveat="@SUM(A1)", approved_by="-reviewer")
    monkeypatch.setattr(exports, "scorecard_frame",
                        lambda store, registry, brand: pd.DataFrame([row])[list(SCORECARD_COLUMNS)])
    return row


def test_xlsx_writes_formula_looking_text_as_text(formula_frame):
    ws = load_workbook(io.BytesIO(exports.scorecard_xlsx(None, None, "=BRAND")))["Scorecard"]
    head = [c.value for c in ws[1]]
    for header in ("Source", "Caveat", "KPI", "Approved by"):
        cell = ws.cell(row=2, column=head.index(header) + 1)
        assert cell.data_type == "s", header
    assert ws.cell(row=2, column=head.index("Source") + 1).value.startswith("=1+1.pptx, slide 2")
    assert ws.cell(row=2, column=head.index("Platform value") + 1).value == -12.5
    notes = load_workbook(io.BytesIO(exports.scorecard_xlsx(None, None, "=BRAND")))["Notes"]
    assert notes["B1"].value == "=BRAND" and notes["B1"].data_type == "s"


def test_csv_prefixes_formula_looking_text(formula_frame):
    back = pd.read_csv(io.BytesIO(exports.scorecard_csv(None, None, "KESTREL")), encoding="utf-8-sig")
    row = back.iloc[0]
    assert row["evidence"].startswith("'=1+1.pptx, slide 2")
    assert row["caveat"] == "'@SUM(A1)" and row["kpi_label"] == "'+TRx" and row["approved_by"] == "'-reviewer"
    assert row["platform_value"] == -12.5 and row["value"] == 5910.0
    assert row["scope"] == "Brand total" and row["period_end"] == "2026-05-31"


@pytest.mark.parametrize("value, expected", [
    ("=HYPERLINK(1)", "'=HYPERLINK(1)"), ("+1", "'+1"), ("-1", "'-1"), ("@x", "'@x"), ("\tx", "'\tx"),
    ("\rx", "'\rx"), ("plain", "plain"), ("", ""), (-1.5, -1.5), (None, None)])
def test_csv_text(value, expected):
    assert csv_text(value) == expected
