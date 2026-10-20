from __future__ import annotations

import dataclasses
from pathlib import Path

import openpyxl
import pytest

from kpi_evidence.report import FRAME_NAMES, PLATFORM_COLUMNS, sql_file_text


def _dbfs(path: Path) -> str:
    return "dbfs:/" + path.relative_to(path.anchor).as_posix()


@pytest.fixture()
def in_tmp(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    return tmp_path


def test_platform_values_cover_every_sql_artifact(report_result):
    frame = report_result.platform_values
    assert list(frame.columns) == list(PLATFORM_COLUMNS)
    assert len(frame) == len(report_result._index.series) > 0
    sql = report_result.sql
    assert set(frame["artifact_id"]) == set(sql.loc[~sql["alternative"], "artifact_id"])
    assert frame["value"].dtype == "float64"
    trx = frame[(frame["kpi_id"] == "trx") & (frame["scope"] == "total") & (frame["period_grain"] == "month")]
    assert len(trx) > 1
    for _artifact, series in trx.groupby("artifact_id"):
        assert series["period_end"].is_monotonic_increasing and series["period_end"].is_unique
    assert "platform_values" in FRAME_NAMES and "platform_values" in report_result.frames()


def test_excel_has_a_platform_series_sheet(report_result, tmp_path):
    wb = openpyxl.load_workbook(report_result.save_excel(tmp_path / "book.xlsx"))
    assert "Platform series" in wb.sheetnames
    ws = wb["Platform series"]
    head = [c.value for c in ws[1]]
    assert head == list(PLATFORM_COLUMNS) and ws.max_row == len(report_result.platform_values) + 1
    col = head.index("value") + 1
    units = head.index("unit") + 1
    fractions = [r for r in range(2, ws.max_row + 1) if ws.cell(r, units).value == "fraction"
                 and ws.cell(r, col).value is not None]
    assert fractions and all("%" in ws.cell(r, col).number_format for r in fractions)


def test_save_sql_writes_one_file_per_artifact_from_the_stored_text(report_result, tmp_path):
    sql = report_result.sql.copy()
    sql["export_path"] = str(tmp_path / "gone" / "missing.sql")
    result = dataclasses.replace(report_result, sql=sql)
    paths = result.save_sql(tmp_path / "sql")
    assert len(paths) == len(sql) == len(set(paths))
    for path, row in zip(paths, sql.to_dict(orient="records"), strict=True):
        text = Path(path).read_text(encoding="utf-8")
        assert Path(path).parent.name == row["kpi_id"] and path.endswith(".sql")
        assert text == sql_file_text(row)
        assert row["value_sql"].strip().rstrip(";") in text and row["artifact_id"] in text
        assert "CREATE OR REPLACE FUNCTION" not in text.upper()
    assert not (tmp_path / "gone").exists()


def test_sql_file_text_keeps_header_values_on_comment_lines():
    row = {"kpi_id": "k", "kpi_label": "Label", "definition": "two\nlines", "definition_source": "hint",
           "choices": "(no options)", "scope": "total", "period_grain": "month", "artifact_id": "a1", "sql_sha": "s",
           "value_sql": "SELECT 1 AS value", "components_sql": "", "checks_sql": "", "function_sql": ""}
    text = sql_file_text(row)
    assert "-- Definition: two lines\n" in text
    assert text.endswith("-- Value\nSELECT 1 AS value;\n")


def test_outputs_accept_dbfs_volume_paths(report_result, in_tmp):
    out = in_tmp / "Volumes" / "main" / "brand" / "out"
    html = report_result.save_html(_dbfs(out / "kpi.html"))
    book = report_result.save_excel(_dbfs(out / "kpi.xlsx"))
    files = report_result.save_sql(_dbfs(out / "sql"))
    assert Path(html).resolve() == (out / "kpi.html").resolve() and Path(html).is_file()
    assert Path(book).resolve() == (out / "kpi.xlsx").resolve() and Path(book).is_file()
    assert files and all(Path(f).resolve().is_relative_to(out.resolve()) for f in files)
    assert not any(str(p).startswith("dbfs:") for p in (html, book, *files))
    assert not (in_tmp / "dbfs:").exists()
