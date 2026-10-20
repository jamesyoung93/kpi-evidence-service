from __future__ import annotations

import builtins
import dataclasses
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

from kpi_evidence.report import FRAME_NAMES, build_result
from kpi_evidence.report_html import render

SCRATCH = Path(__file__).resolve().parents[2] / ".kpi_evidence_scratch" / "report"

TILES = ("Matches platform", "Captured", "Needs review", "Blocked", "Other numbers")
EXTERNAL = re.compile(r"""(?:src|href)\s*=\s*["']?\s*(?:https?:)?//|url\(\s*["']?(?:https?:)?//|@import|<link\b|"""
                      r"""https?://""", re.I)


@pytest.fixture(scope="module")
def report_html(report_result):
    return report_result.to_html(title="KPI check: Kestrel sample reports")


def test_html_is_self_contained(report_html):
    assert report_html.startswith("<!doctype html>")
    assert not EXTERNAL.search(report_html), EXTERNAL.search(report_html).group(0)
    assert "<style>" in report_html and "<script>" in report_html
    assert "<title>KPI check: Kestrel sample reports</title>" in report_html


def test_html_tiles_and_verdicts(report_result, report_html):
    for name in TILES:
        assert f'<div class="l">{name}</div>' in report_html
    assert 'class="tile good"' in report_html and 'title="The platform number' in report_html
    assert "KPI verdicts" in report_html
    for verdict in set(report_result.kpis["verdict"]):
        assert f"</span>{verdict}</span>" in report_html
    assert report_html.count('class="sortable"') >= 4


def test_one_chart_per_kpi_with_platform_values(report_result, report_html):
    df = report_result.comparisons
    charted = df["platform_value"].notna() | df["alternative_value"].notna()
    with_platform = df[charted]["kpi_id"].nunique()
    assert with_platform >= 8
    assert report_html.count("<svg") == with_platform
    assert len(re.findall(r'<svg [^>]*role="img" aria-labelledby="[^"]+"><title id=', report_html)) == with_platform
    assert "No platform values to chart for this KPI" in report_html


def test_rows_have_evidence_drawers_with_sql(report_result, report_html):
    assert report_html.count('<details class="drawer">') == len(report_result.comparisons)
    assert report_html.count('<tr class="ev">') == len(report_result.comparisons)
    assert "<pre><code>SELECT SUM(trx_units) AS value" in report_html
    assert "Alternatives tried" in report_html and "matches the report" in report_html
    assert "Value as printed" in report_html and "5.7K" in report_html
    for label in ("Matches platform", "Definition differs", "No access to platform data", "Reports disagree"):
        assert label in report_html


def test_status_chip_tooltips_come_from_status_info(report_html):
    from kpi_evidence.contracts import STATUS_INFO, ValidationStatus
    import html

    info = STATUS_INFO[ValidationStatus.MISSING_ACCESS]
    assert f'title="{html.escape(info.tooltip, quote=True)}"' in report_html


def test_html_escapes_malicious_text(report_result):
    payload = '<script>alert("x")</script><img src=x onerror=alert(1)>'
    comps = report_result.comparisons.copy()
    comps.loc[comps.index[0], "label_as_printed"] = payload
    comps.loc[comps.index[1], "object"] = payload
    comps.loc[comps.index[2], "report_file"] = payload + ".pptx"
    kpis = report_result.kpis.copy()
    kpis.loc[kpis.index[0], "kpi_label"] = payload
    evil = dataclasses.replace(report_result, comparisons=comps, kpis=kpis)
    out = render(evil)
    assert payload not in out and "<script>alert" not in out and "<img src=x" not in out
    assert "&lt;script&gt;alert(&quot;x&quot;)&lt;/script&gt;" in out
    assert out.count("<script>") == 1


def test_demo_html_written_for_inspection(report_result):
    path = report_result.save_html(SCRATCH / "demo_report.html", title="KPI check: Kestrel sample reports")
    text = open(path, encoding="utf-8").read()
    assert text.startswith("<!doctype html>") and len(text) > 10_000


def test_excel_workbook(report_result, tmp_path):
    path = report_result.save_excel(tmp_path / "out" / "kpi_check.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Summary", "Comparisons", "KPI verdicts", "SQL", "Platform series", "Documents",
                             "Extracted values", "Issues"]
    frames = {"Comparisons": report_result.comparisons, "KPI verdicts": report_result.kpis, "SQL": report_result.sql,
              "Platform series": report_result.platform_values, "Documents": report_result.documents,
              "Extracted values": report_result.extraction, "Issues": report_result.issues}
    for sheet, frame in frames.items():
        ws = wb[sheet]
        assert [c.value for c in ws[1]] == list(frame.columns), sheet
        assert ws.max_row == len(frame) + 1, sheet
        assert ws.freeze_panes == "A2" and ws.auto_filter.ref, sheet
        assert ws["A1"].font.bold, sheet
    ws = wb["Comparisons"]
    head = [c.value for c in ws[1]]
    col = {name: head.index(name) + 1 for name in head}
    fractions = [r for r in range(2, ws.max_row + 1) if ws.cell(r, col["unit"]).value == "fraction"
                 and ws.cell(r, col["report_value"]).value is not None]
    assert fractions and all("%" in ws.cell(r, col["report_value"]).number_format for r in fractions)
    assert isinstance(ws.cell(2, col["period_end"]).value, datetime)
    sql = wb["SQL"]
    sql_head = [c.value for c in sql[1]]
    assert sql.cell(2, sql_head.index("value_sql") + 1).alignment.wrap_text
    summary = wb["Summary"]
    assert summary["A1"].value == "KPI check report"
    labels = {summary.cell(r, 1).value for r in range(1, summary.max_row + 1)}
    assert {"Matches platform", "Needs review", "Comparisons by flag"} <= labels


def test_excel_never_writes_formulas(report_result, tmp_path):
    comps = report_result.comparisons.copy()
    comps.loc[comps.index[0], "label_as_printed"] = '=HYPERLINK("http://x","click")'
    comps.loc[comps.index[1], "object"] = "Line oneline two"
    path = dataclasses.replace(report_result, comparisons=comps).save_excel(tmp_path / "f.xlsx")
    ws = openpyxl.load_workbook(path)["Comparisons"]
    head = [c.value for c in ws[1]]
    cell = ws.cell(2, head.index("label_as_printed") + 1)
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")
    assert ws.cell(3, head.index("object") + 1).value == "Line oneline two"
    assert ws.cell(2, head.index("metadata_confirmed") + 1).data_type == "b"


def test_check_level_issues_are_counted_and_rendered(report_run, tmp_path):
    r = report_run
    extra = [{"severity": "warning", "category": "KPI table", "message": "column 'grain' unknown", "count": 1},
             {"severity": "info", "category": "engine note", "message": "periods listed from the compared reports"}]
    result = build_result(r.store, r.registry, r.settings, r.doc_ids, run_id=r.run_id, extra_issues=extra)
    issues = result.issues
    assert {"KPI table", "engine note"} <= set(issues["category"])
    assert result.summary["issues"] == len(issues) and result.summary["hint_messages"] == 1
    html_text = result.to_html()
    assert f"Issues ({len(issues)})" in html_text and "column &#x27;grain&#x27; unknown" in html_text
    ws = openpyxl.load_workbook(result.save_excel(tmp_path / "issues.xlsx"))["Issues"]
    assert ws.max_row == len(issues) + 1


class FakeSpark:
    def __init__(self):
        self.calls = []

    def createDataFrame(self, rows, schema=None):
        self.calls.append((rows, schema))
        return {"rows": rows, "schema": schema}


def test_to_spark_with_fake_session(report_result):
    spark = FakeSpark()
    out = report_result.to_spark(spark)
    assert set(out) == set(FRAME_NAMES) and len(spark.calls) == len(FRAME_NAMES)
    comp = out["comparisons"]
    assert len(comp["rows"]) == len(report_result.comparisons)
    schema = comp["schema"]
    for part in ("`report_value` DOUBLE", "`period_end` DATE", "`metadata_confirmed` BOOLEAN", "`page` BIGINT",
                 "`sql` STRING", "`kpi_id` STRING"):
        assert part in schema
    assert schema.count("`") == 2 * len(report_result.comparisons.columns)
    allowed = (str, float, int, bool, date, type(None))
    for rows, _schema in spark.calls:
        for row in rows:
            assert isinstance(row, tuple)
            assert all(isinstance(v, allowed) for v in row), [type(v) for v in row]
    kinds = dict(part.rsplit(" ", 1) for part in schema.split(", "))
    for row in comp["rows"]:
        for (name, kind), value in zip(kinds.items(), row, strict=True):
            if value is not None and kind == "STRING":
                assert isinstance(value, str), name


def test_to_spark_empty_frames_keep_a_schema(report_result):
    empty = dataclasses.replace(report_result, issues=report_result.issues.iloc[0:0])
    spark = FakeSpark()
    out = empty.to_spark(spark)
    assert out["issues"]["rows"] == [] and "`severity` STRING" in out["issues"]["schema"]


def test_to_spark_without_session_is_an_explicit_error(report_result):
    with pytest.raises(RuntimeError, match="SparkSession|pyspark"):
        report_result.to_spark()


def test_evidence_by_index_result_and_observation(report_result):
    df = report_result.comparisons
    i = int(df.index[(df["kpi_id"] == "call_attainment") & (df["status"] == "INCOMPATIBLE_DEFINITION")][0])
    ev = report_result.evidence(i)
    row = df.loc[i]
    assert ev["comparison"]["result_id"] == row["result_id"]
    assert report_result.evidence(row["result_id"])["comparison"]["result_id"] == row["result_id"]
    assert report_result.evidence(row["observation_id"])["comparison"]["result_id"] == row["result_id"]
    assert ev["trials"] and {"phase", "choices", "value", "sql", "as_of"} <= set(ev["trials"][0])
    assert any(t["phase"] == "phase_a" and t["for_this_period"] for t in ev["trials"])
    assert ev["extracted_value"]["raw_text"] == "95.8%" and ev["extracted_value"]["location"].startswith("slide 3")
    assert ev["page"]["page_number"] == 3 and ev["page"]["title"]
    assert ev["recon_run"]["start_choices"] == {"aggregation": "ratio_of_sums", "eligibility": "all_hcps"}
    assert ev["sql_artifact"]["value_sql_template"]
    assert ev["explanation"]["alternative"] == "eligibility=writers_only"
    matching = [a for a in ev["alternatives"] if a["matches_report"]]
    assert matching and matching[0]["differs"] == {"eligibility": "writers_only"}
    assert ev["status_tooltip"].startswith("The report seems to use a different definition")


def test_evidence_unknown_key(report_result):
    with pytest.raises(KeyError):
        report_result.evidence("no-such-id")
    with pytest.raises(KeyError):
        report_result.evidence(10_000)


def test_filter(report_result):
    review = report_result.filter(flag="review")
    assert len(review) and set(review["flag"]) == {"REVIEW"}
    trx = report_result.filter(kpi="Total prescriptions (TRx)")
    assert set(trx["kpi_id"]) == {"trx"}
    both = report_result.filter(flag=["GOOD", "INFO"], report="brand_review_2026-06_v2")
    assert len(both) and set(both["report_file"]) == {"Kestrel_Monthly_Brand_Review_2026-06_v2.pptx"}
    assert set(both["flag"]) <= {"GOOD", "INFO"}


def test_show_uses_display_html_when_available(report_result, monkeypatch):
    shown = []
    monkeypatch.setattr(builtins, "displayHTML", shown.append, raising=False)
    report_result.show()
    assert len(shown) == 1 and shown[0].startswith("<!doctype html>")


def test_show_falls_back_to_text(report_result, monkeypatch, capsys):
    monkeypatch.delattr(builtins, "displayHTML", raising=False)
    report_result.show()
    out = capsys.readouterr().out
    assert "Matches platform" in out and "result.comparisons" in out
    assert "KPI check" in repr(report_result)
