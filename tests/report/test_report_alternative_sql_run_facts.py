from __future__ import annotations

import dataclasses
import html
import inspect
import re
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from kpi_evidence import report, report_excel, report_html
from kpi_evidence.contracts import DocumentStatus, ValidationStatus
from kpi_evidence.reconcile import OUTSIDE_FOCUS
from kpi_evidence.report import (
    ISSUE_COLUMNS, alternative_sql, build_result, durable_path, extraction_frame, kpi_verdict, matches_alternative,
    sql_file_text, summary_dict,
)
from kpi_evidence.report_html import ai_text
from kpi_evidence.store.duckdb_store import DuckDBStore

READ_STATEMENT = re.compile(r"^\s*(SELECT|WITH)\b", re.IGNORECASE)


def _alternative_results(result):
    return [r for r in result._index.results.values() if matches_alternative(r) and alternative_sql(result._index, r)]


def _statements(text: str) -> list[str]:
    body = "\n".join(line for line in text.splitlines() if not line.lstrip().startswith("--"))
    return [s.strip() for s in body.split(";") if s.strip()]


def test_alternative_sql_reaches_frame_drawer_excel_and_files(report_result, tmp_path):
    matched = _alternative_results(report_result)
    assert matched, "the sample reports hold values that only a declared alternative reproduces"
    sql = report_result.sql
    rows = sql[sql["alternative"]]
    texts = {alternative_sql(report_result._index, r) for r in matched}
    assert set(rows["value_sql"]) == texts
    assert rows["comparisons"].sum() == len(matched) and not rows["governed_default"].any()
    assert (rows["export_path"] == "").all() and (rows["function_sql"] == "").all()
    page = report_result.to_html()
    for text in texts:
        assert html.escape(text, quote=True) in page
    assert page.count("Alternative SQL") == len(matched)
    book = openpyxl.load_workbook(report_result.save_excel(tmp_path / "check.xlsx"))["SQL"]
    head = [c.value for c in book[1]]
    cells = {book.cell(r, head.index("value_sql") + 1).value for r in range(2, book.max_row + 1)}
    assert texts <= cells
    files = report_result.save_sql(tmp_path / "sql")
    saved = {Path(f).read_text(encoding="utf-8") for f in files if "_alternative_" in Path(f).name}
    assert len(saved) == len(texts)
    for text in texts:
        assert any(text.rstrip(";") + ";" in body for body in saved)
    for kpi_id in {r.kpi_id for r in matched}:
        assert (tmp_path / "sql" / kpi_id).is_dir()


def test_saved_sql_files_hold_read_queries_only(report_result, tmp_path):
    files = report_result.save_sql(tmp_path / "sql")
    assert files and report_result.sql["function_sql"].str.contains("CREATE OR REPLACE FUNCTION").any()
    for path in files:
        statements = _statements(Path(path).read_text(encoding="utf-8"))
        assert statements and all(READ_STATEMENT.match(s) for s in statements), path
    row = report_result.sql[report_result.sql["alternative"]].iloc[0].to_dict()
    assert "{as_of}" not in sql_file_text(row)


def test_extraction_reason_comes_from_the_document_state_first(report_run, report_result):
    r = report_run
    focused = build_result(r.store, r.registry, r.settings, r.doc_ids, run_id=r.run_id, focus_kpis={"trx"})
    index = focused._index
    doc = next(d for d in index.docs.values() if d.status == DocumentStatus.PROCESSED
               and any(v.doc_id == d.doc_id for v in index.values.values()))
    waiting = dataclasses.replace(doc, status=DocumentStatus.AWAITING_CONFIRMATION, needs_confirmation=["period_end"])
    changed = dataclasses.replace(index, docs={**index.docs, doc.doc_id: waiting},
                                  results={k: v for k, v in index.results.items() if v.doc_id != doc.doc_id})
    frame = extraction_frame(changed)
    rows = frame[frame["doc_id"] == doc.doc_id]
    assert set(rows["reason"]) == {"Not compared: the report is awaiting confirmation (period_end)."}
    assert rows["status"].isna().all()
    for result in (report_result, focused):
        docs = result.documents
        superseded = set(docs.loc[docs["status"] == "SUPERSEDED", "doc_id"])
        old = result.extraction[result.extraction["doc_id"].isin(superseded)]
        assert len(old) and old["reason"].str.startswith("Not compared: superseded by ").all()
    processed = set(focused.documents.loc[focused.documents["status"] == "PROCESSED", "doc_id"])
    outside = focused.extraction[focused.extraction["reason"] == OUTSIDE_FOCUS]
    assert len(outside) and outside["doc_id"].isin(processed).all() and not outside["in_focus"].any()


def _copies(result, status, *keys):
    out = []
    for i, (end, scope) in enumerate(keys):
        out.append(dataclasses.replace(result, result_id=f"{result.result_id}-{i}", period_end=end, scope_key=scope,
                                       validation_status=status[i], platform_value=1.0, explanation={},
                                       reason_code="match"))
    return out


def test_verdict_detail_counts_period_and_scope_combinations(report_result):
    index = report_result._index
    base = next(r for r in index.results.values() if r.kpi_id and r.period_end is not None)
    kpi = index.kpis.get(base.kpi_id)
    ok, bad = ValidationStatus.VERIFIED, ValidationStatus.VARIANCE_REVIEW
    end = base.period_end
    scoped = kpi_verdict(kpi, _copies(base, [bad, ok, ok], (end, "{}"), (end, '{"region": "East"}'),
                                      (end, '{"region": "West"}')))
    detail = scoped["verdict_detail"]
    assert "2 of 3 period and scope combination(s)" in detail and f"not in {base.period_grain} {end} total" in detail
    assert "period(s)" not in detail
    later = end.replace(year=end.year + 1)
    plain = kpi_verdict(kpi, _copies(base, [ok, bad], (end, "{}"), (later, "{}")))
    assert "Validated in 1 of 2 period(s); not in" in plain["verdict_detail"]
    assert f"{later} total" not in plain["verdict_detail"]


def test_saved_results_line_needs_a_state_file(report_run, report_result, mem_store, registry, tmp_path):
    assert report_result.summary["state_saved"] is False
    assert "Saved results" not in report_result.to_html()
    kept = dataclasses.replace(report_result, summary={**report_result.summary, "state_saved": True})
    assert "Saved results" in kept.to_html() and "0 of " in kept.to_html()
    store = DuckDBStore(mem_store.settings, tmp_path / "state.duckdb")
    store.ensure_schema()
    assert build_result(store, registry, mem_store.settings, []).summary["state_saved"] is True
    assert build_result(mem_store, registry, mem_store.settings, []).summary["state_saved"] is False


def test_run_facts_show_the_ai_mode_and_failed_model_calls(report_run, report_result):
    summary = report_result.summary
    assert summary["ai_mode"] == report_run.settings.ai_mode and summary["ai_failures"] == 0
    assert "<dt>AI</dt><dd>off (no model calls)</dd>" in report_result.to_html()
    index = report_result._index
    key, mapping = next(iter(index.mappings.items()))
    failed = dataclasses.replace(mapping, candidates=[*(mapping.candidates or []), {
        "kpi_id": None, "method": "MODEL", "confidence": 0.0, "source": "model", "note": "error"}])
    changed = dataclasses.replace(index, mappings={**index.mappings, key: failed})
    issues = pd.DataFrame([{**dict.fromkeys(ISSUE_COLUMNS), "severity": "warning", "category": "AI calls",
                            "message": "ai_query failed"}], columns=list(ISSUE_COLUMNS))
    frames = report_result.frames()
    facts = summary_dict(changed, frames, issues, hints=None, run_id="r", platform_description="",
                         timings=None, settings=dataclasses.replace(report_run.settings, ai_mode="ai_query"))
    assert facts["ai_mode"] == "ai_query" and facts["ai_failures"] == 1 and facts["ai_issues"] == 1
    text = ai_text(facts)
    assert text.startswith("ai_query") and "failed for 1 value(s)" in text and "1 AI message(s) under Issues" in text
    page = dataclasses.replace(report_result, summary={**summary, **facts}).to_html()
    assert html.escape(text) in page


def test_html_never_prints_temporary_paths(report_result):
    temp = Path(tempfile.gettempdir())
    assert durable_path(temp / "kpi_evidence_check_x" / "a.sql") == ""
    assert durable_path("/Volumes/main/brand/out/a.sql") == "/Volumes/main/brand/out/a.sql"
    project = Path(__file__).resolve().parent / "a.sql"
    if not str(project).lower().startswith(str(temp.resolve()).lower()):
        assert durable_path(project) == str(project)
    assert (report_result.sql["export_path"] == "").all()
    page = report_result.to_html()
    assert str(temp) not in page and temp.as_posix() not in page and "Saved to" not in page
    sql = report_result.sql.copy()
    sql["export_path"] = [str(temp / "x.sql")] + ["/Volumes/main/brand/sql/kpi.sql"] * (len(sql) - 1)
    page = dataclasses.replace(report_result, sql=sql).to_html()
    assert str(temp) not in page and "Saved to /Volumes/main/brand/sql/kpi.sql" in page


def test_report_docstrings_are_single_lines_without_document_references():
    seen = 0
    for module in (report, report_html, report_excel):
        members = [obj for name, obj in vars(module).items() if not name.startswith("_")
                   and (inspect.isfunction(obj) or inspect.isclass(obj)) and obj.__module__ == module.__name__]
        members += [obj for cls in list(members) if inspect.isclass(cls)
                    for name, obj in vars(cls).items() if not name.startswith("_") and inspect.isfunction(obj)]
        for obj in members:
            doc = obj.__doc__
            if doc:
                seen += 1
                assert "\n" not in doc.strip() and ".md" not in doc and not re.search(r"\bsections?\s+\d", doc), obj
    assert seen > 20
