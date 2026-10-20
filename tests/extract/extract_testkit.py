from __future__ import annotations

import io
import zipfile
from datetime import date, datetime

import pytest

from kpi_evidence.contracts import Document, DocumentFormat, DocumentStatus
from kpi_evidence.ids import sha256_bytes

TRX = [5909.796, 5800.1, 5750.25, 5810.0, 5799.5, 5783.041]
MONTHS = [date(2025, 12, 31), date(2026, 1, 31), date(2026, 2, 28), date(2026, 3, 31), date(2026, 4, 30), date(2026, 5, 31)]


def png_bytes(size=(320, 200), color=(40, 90, 160)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    return buf.getvalue()


def rewrite_zip(blob: bytes, changes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    parts.update(changes(parts))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return out.getvalue()


def _textbox(slide, x, y, w, h, text):
    from pptx.util import Inches

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    lines = text.split("\n")
    frame.text = lines[0]
    for line in lines[1:]:
        frame.add_paragraph().text = line
    return box


def build_deck() -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    deck = Presentation()
    deck.core_properties.title = "Kestrel Monthly Brand Review"
    deck.core_properties.created = datetime(2026, 6, 3, 9, 0, 0)
    deck.core_properties.modified = datetime(2026, 6, 4, 10, 30, 0)
    layout = deck.slide_layouts[5]

    s1 = deck.slides.add_slide(layout)
    s1.shapes.title.text = "TRx trend"
    data = CategoryChartData(number_format="#,##0")
    data.categories = MONTHS
    data.add_series("TRx", TRX)
    chart = s1.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(1), Inches(1.5), Inches(8), Inches(4), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Monthly TRx"
    _textbox(s1, 1, 6.0, 8, 0.4, "Source: account feed. Data through May 2026")
    s1.notes_slide.notes_text_frame.text = "Illustrative data only."

    s2 = deck.slides.add_slide(layout)
    s2.shapes.title.text = "KPI summary"
    rows = [["Metric", "May 2026", "vs PY"], ["TRx", "5.9K", "+6%"], ["TRx share", "84.4%", "(0.3%)"],
            ["Net change", "(1,234)", "n/a"]]
    table = s2.shapes.add_table(4, 3, Inches(1), Inches(1.5), Inches(8), Inches(2)).table
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    s2.shapes.add_picture(io.BytesIO(png_bytes((40, 40))), Inches(9.2), Inches(0.2), Inches(0.5), Inches(0.5))

    s3 = deck.slides.add_slide(layout)
    s3.shapes.title.text = "Regional performance"
    for i, (heading, series, values, fmt) in enumerate([
            ("NRx by region", "NRx", [1200, 1265.7], "#,##0"),
            ("Call attainment by region", "Attainment", [0.844, 0.919], "0.0%")]):
        _textbox(s3, 0.5 + i * 4.75, 1.3, 4.25, 0.4, heading)
        cd = CategoryChartData(number_format=fmt)
        cd.categories = ["East", "West"]
        cd.add_series(series, values)
        s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5 + i * 4.75), Inches(1.8), Inches(4.25),
                            Inches(3.5), cd)

    s4 = deck.slides.add_slide(layout)
    s4.shapes.title.text = "Highlights"
    _textbox(s4, 0.5, 1.5, 5, 2.5, "Brand awareness: 62%\nSpeaker programs drove +8% TRx\n47 HCPs attended\n"
                                    "IL-17 class share 26.2%\nPage 4 of 6")
    _textbox(s4, 6, 1.5, 3, 1.2, "184\nReferrals (projected)")

    s5 = deck.slides.add_slide(layout)
    s5.shapes.title.text = "Backup: samples"
    _textbox(s5, 1, 1.5, 6, 0.5, "Samples 1,200")
    s5._element.set("show", "0")

    s6 = deck.slides.add_slide(layout)
    s6.shapes.title.text = "Field dashboard"
    s6.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(1.5), Inches(8), Inches(5))

    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def build_workbook() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "KPIs"
    ws.append(["Metric", "May 2026", "Apr 2026"])
    ws.append(["TRx", 5909.796, 5799.5])
    ws.append(["TRx share", 0.0917, 0.091])
    ws.append(["Calls", "=SUM(B2:B3)", 2031.5])
    ws.append([])
    ws.append(["Region", "TRx", "vs PY"])
    ws.append(["West", 1265.7, 0.06])
    for cell in ("B2", "C2"):
        ws[cell].number_format = "#,##0"
    for cell in ("B3", "C3", "C7"):
        ws[cell].number_format = "0.0%"
    ws.row_dimensions[7].hidden = True
    notes = wb.create_sheet("Notes")
    notes.sheet_state = "hidden"
    notes.append(["Persistency", 0.62])
    wb.properties.title = "Kestrel KPI workbook"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf() -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Table, TableStyle

    buf = io.BytesIO()
    pdf = canvas.Canvas(buf, pagesize=letter)
    pdf.setTitle("Kestrel Executive Summary")
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(72, 720, "Kestrel Executive Summary Q2 2026")
    pdf.setFont("Helvetica", 10)
    pdf.drawString(72, 690, "Call attainment: 84.4% of plan")
    table = Table([["KPI", "Q2 2026", "vs PY"], ["TRx", "17,220", "-31.6%"], ["TRx share", "9.2%", "+0.4pp"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    table.wrapOn(pdf, 400, 200)
    table.drawOn(pdf, 72, 590)
    pdf.drawString(72, 560, "Data through June 2026")
    pdf.showPage()
    pdf.drawImage(ImageReader(io.BytesIO(png_bytes((600, 800), (200, 200, 200)))), 36, 36, 540, 720)
    pdf.showPage()
    pdf.save()
    return buf.getvalue()


def make_doc(blob: bytes, file_name: str, fmt: DocumentFormat) -> Document:
    now = datetime(2026, 6, 5, 12, 0, 0)
    return Document(doc_id=sha256_bytes(blob), file_name=file_name, file_path=f"/tmp/{file_name}",
                    size_bytes=len(blob), format=fmt, status=DocumentStatus.REGISTERED, registered_at=now, updated_at=now)


@pytest.fixture(scope="session")
def deck_bytes() -> bytes:
    return build_deck()


@pytest.fixture(scope="session")
def workbook_bytes() -> bytes:
    return build_workbook()


@pytest.fixture(scope="session")
def pdf_bytes() -> bytes:
    return build_pdf()


@pytest.fixture(scope="session")
def deck_output(deck_bytes):
    from kpi_evidence.extract.pptx import extract_pptx_values

    return extract_pptx_values(sha256_bytes(deck_bytes), deck_bytes, "Kestrel_Monthly_Brand_Review_2026-06.pptx",
                               run_id="run-1")
