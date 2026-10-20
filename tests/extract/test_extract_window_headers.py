from __future__ import annotations

import io

from kpi_evidence.contracts import ExtractionMethod
from kpi_evidence.extract._common import analyze_grid, text_cell_kind
from kpi_evidence.extract.pptx import extract_pptx_values
from kpi_evidence.extract.xlsx import extract_xlsx_values
from kpi_evidence.ids import sha256_bytes

ROWS = [["Metric", "1M", "3M", "12M"], ["TRx", "5,736", "17,694", "72,105"], ["TRx share", "9.1%", "9.4%", "9.8%"]]


def _layout(rows):
    kinds = [[text_cell_kind(t)[0] for t in row] for row in rows]
    return analyze_grid(rows, kinds)


def test_window_tokens_make_a_header_row():
    layout = _layout(ROWS)
    assert layout.header_rows == 1
    assert layout.col_headers == ["Metric", "1M", "3M", "12M"]
    assert layout.header_cells == frozenset({(0, 1), (0, 2), (0, 3)})


def test_suffixed_values_below_keep_the_row_as_data():
    rows = [["Impressions", "3M", "6M"], ["Reach", "1.2M", "2.5M"]]
    layout = _layout(rows)
    assert layout.header_rows == 0 and layout.header_cells == frozenset()


def test_a_value_row_after_a_header_is_not_taken_as_a_second_header():
    for rows in ([["Metric", "Apr 2026", "May 2026"], ["Impressions", "3M", "6M"], ["Reach", "1.2M", "2.5M"]],
                 [["Metric", "Apr 2026", "May 2026"], ["Impressions", "3M", "6M"], ["TRx", "5,736", "5,800"]]):
        layout = _layout(rows)
        assert layout.header_rows == 1 and layout.header_cells == frozenset()


def test_windows_under_a_merged_group_header():
    rows = [["", "TRx", "", "NRx", ""], ["Region", "3M", "12M", "3M", "12M"], ["West", "1,200", "4,800", "300", "1,250"]]
    layout = _layout(rows)
    assert layout.header_rows == 2
    assert layout.header_cells == frozenset({(1, 1), (1, 2), (1, 3), (1, 4)})
    spanned = analyze_grid(rows, [[text_cell_kind(t)[0] for t in row] for row in rows], spans=[(0, 1, 2), (0, 3, 2)])
    assert spanned.col_headers[1:] == ["TRx / 3M", "TRx / 12M", "NRx / 3M", "NRx / 12M"]


def _deck() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "TRx windows"
    table = slide.shapes.add_table(3, 4, Inches(1), Inches(1.5), Inches(8), Inches(2)).table
    for r, row in enumerate(ROWS):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def test_pptx_table_with_window_headers():
    blob = _deck()
    out = extract_pptx_values(sha256_bytes(blob), blob, "deck.pptx")
    cells = [v for v in out.values if v.method == ExtractionMethod.NATIVE_TABLE]
    assert [v.raw_text for v in cells] == ["5,736", "17,694", "72,105", "9.1%", "9.4%", "9.8%"]
    assert [v.context["col_header"] for v in cells] == ["1M", "3M", "12M"] * 2
    assert {v.context["row_label"] for v in cells} == {"TRx", "TRx share"}


def test_worksheet_with_window_headers():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "KPIs"
    ws.append(["Metric", "1M", "3M", "12M"])
    ws.append(["TRx", 5736, 17694, 72105])
    buf = io.BytesIO()
    wb.save(buf)
    blob = buf.getvalue()
    out = extract_xlsx_values(sha256_bytes(blob), blob, "kpis.xlsx")
    assert [(v.locator["cell"], v.raw_number, v.context["col_header"]) for v in out.values] == [
        ("B2", 5736.0, "1M"), ("C2", 17694.0, "3M"), ("D2", 72105.0, "12M")]
