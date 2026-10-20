from __future__ import annotations

import io
from xml.etree import ElementTree as X

import pytest
from extract_testkit import MONTHS, TRX, build_deck, rewrite_zip

from kpi_evidence.contracts import DocumentPage, ExtractedValue, ExtractionMethod
from kpi_evidence.extract._vendor import pptx_extract as px
from kpi_evidence.extract.pptx import extract_pptx_values, formula_cells
from kpi_evidence.ids import sha256_bytes, stable_id

LOCATOR_KEYS = {"page", "record_id", "part", "shape_id", "shape_name", "chart_index", "series_index", "point_index",
                "row", "col", "sheet", "cell", "bbox", "element_id",
                "line", "span", "image_index"}
CONTEXT_KEYS = {"object_kind", "page_title", "object_title", "series_name", "category", "row_label", "col_header",
                "header_row", "axis_title", "nearby_text", "footnotes", "text_line", "hidden", "position",
                "number_source", "chart_type", "series_count", "category_count",
                "label_before", "label_after", "row_labels", "category_levels", "number_format_source",
                "workbook_ref", "x", "alt_text", "area_fraction", "media_part", "media_sha256", "external",
                "placeholder_type", "source_kind", "formula"}


def _extract(blob, **kw):
    return extract_pptx_values(sha256_bytes(blob), blob, "deck.pptx", **kw)


def _by(out, method, page=None):
    return [v for v in out.values if v.method == method and (page is None or v.slide_or_page == page)]


def _codes(value):
    return [f["code"] for f in value.flags]


def test_line_chart_gives_six_monthly_points(deck_output):
    points = _by(deck_output, ExtractionMethod.NATIVE_CHART, page=1)
    assert [v.raw_number for v in points] == TRX
    assert [v.context["category"] for v in points] == [m.isoformat() for m in MONTHS]
    first = points[0]
    assert first.raw_text == "5909.796" and first.number_format == "#,##0" and first.confidence == 1.0
    assert first.locator["part"] == "ppt/charts/chart1.xml" and first.locator["point_index"] == 0
    assert first.locator["chart_index"] == 1 and first.locator["series_index"] == 0
    assert len(first.locator["bbox"]) == 4
    ctx = first.context
    assert ctx["object_kind"] == "chart" and ctx["object_title"] == "Monthly TRx" and ctx["series_name"] == "TRx"
    assert ctx["page_title"] == "TRx trend" and ctx["number_source"] == "stored" and ctx["chart_type"] == "lineChart"
    assert ctx["series_count"] == 1 and ctx["category_count"] == 6
    assert "Source: account feed. Data through May 2026" in ctx["footnotes"]
    assert first.value_id == stable_id(first.doc_id, "NATIVE_CHART", first.locator)
    assert first.source_record_id and first.run_id == "run-1"


def test_chart_linked_workbook_is_a_freshness_check_not_values(deck_output):
    assert _by(deck_output, ExtractionMethod.EMBEDDED_WORKSHEET) == []
    for v in _by(deck_output, ExtractionMethod.NATIVE_CHART):
        assert "CACHE_MATCHES_WORKBOOK" in _codes(v)
        assert v.context["workbook_ref"]["sheet"] == "Sheet1"


def test_stale_chart_cache_is_flagged(deck_bytes):
    def patch(parts):
        root = X.fromstring(parts["ppt/charts/chart1.xml"])
        root.find(".//c:ser/c:val/c:numRef/c:numCache/c:pt/c:v", px.NS).text = "6000"
        return {"ppt/charts/chart1.xml": X.tostring(root)}

    points = _by(_extract(rewrite_zip(deck_bytes, patch)), ExtractionMethod.NATIVE_CHART, page=1)
    assert points[0].raw_number == 6000.0 and "CACHE_DIFFERS_FROM_WORKBOOK" in _codes(points[0])
    assert "5909.796" in next(f["detail"] for f in points[0].flags if f["code"] == "CACHE_DIFFERS_FROM_WORKBOOK")
    assert all("CACHE_MATCHES_WORKBOOK" in _codes(v) for v in points[1:])


def test_percent_chart_keeps_stored_fraction(deck_output):
    attainment = [v for v in _by(deck_output, ExtractionMethod.NATIVE_CHART, page=3) if v.context["series_name"] == "Attainment"]
    assert [v.raw_number for v in attainment] == [0.844, 0.919]
    assert {v.number_format for v in attainment} == {"0.0%"}


def test_side_by_side_charts_stay_separate(deck_output):
    points = _by(deck_output, ExtractionMethod.NATIVE_CHART, page=3)
    left = [v for v in points if v.locator["chart_index"] == 1]
    right = [v for v in points if v.locator["chart_index"] == 2]
    assert {v.context["series_name"] for v in left} == {"NRx"} and {v.context["series_name"] for v in right} == {"Attainment"}
    assert left[0].locator["shape_id"] != right[0].locator["shape_id"]
    assert left[0].context["position"] == "middle-left" and right[0].context["position"] == "middle-right"
    assert left[0].context["nearby_text"] == ["NRx by region"]
    assert right[0].context["nearby_text"] == ["Call attainment by region"]
    for v in points:
        assert {"MULTIPLE_CHARTS_WITHOUT_LOCAL_LABEL", "NEARBY_LABEL_UNVERIFIED"} <= set(_codes(v))
        assert next(f for f in v.flags if f["code"] == "NEARBY_LABEL_UNVERIFIED")["dimension"] == "meaning"


def test_table_cells_with_header_and_row_labels(deck_output):
    cells = {(v.locator["row"], v.locator["col"]): v for v in _by(deck_output, ExtractionMethod.NATIVE_TABLE, page=2)}
    assert set(cells) == {(1, 1), (1, 2), (2, 1), (2, 2), (3, 1)}
    k = cells[(1, 1)]
    assert (k.raw_text, k.raw_number, k.number_format, k.confidence) == ("5.9K", 5.9, None, 1.0)
    assert k.context["row_label"] == "TRx" and k.context["col_header"] == "May 2026"
    assert k.context["header_row"] == ["Metric", "May 2026", "vs PY"] and k.context["number_source"] == "text"
    assert cells[(2, 1)].raw_text == "84.4%" and cells[(2, 1)].raw_number == 84.4
    assert cells[(3, 1)].raw_text == "(1,234)" and cells[(3, 1)].raw_number == -1234.0
    assert "DERIVED_CHANGE" in _codes(cells[(1, 2)]) and "DERIVED_CHANGE" not in _codes(k)
    assert k.locator["shape_name"] == "Table 2" and k.locator["part"] == "ppt/slides/slide2.xml"


def test_slide_text_labeled_numbers(deck_output):
    text = {v.raw_text: v for v in _by(deck_output, ExtractionMethod.SLIDE_TEXT, page=4)}
    assert set(text) == {"62%", "+8%", "47", "26.2%", "184"}
    awareness = text["62%"]
    assert awareness.raw_number == 62.0 and awareness.confidence == 0.7
    assert awareness.context["row_label"] == "Brand awareness"
    assert awareness.context["text_line"] == "Brand awareness: 62%"
    assert awareness.context["object_kind"] == "slide_text" and awareness.context["number_source"] == "text"
    assert awareness.locator["shape_id"] and len(awareness.locator["bbox"]) == 4
    assert text["+8%"].context["label_after"] == "TRx"
    assert text["47"].context["row_label"] == "HCPs attended"
    callout = text["184"]
    assert callout.context["row_label"] == "Referrals (projected)" and "LABEL_FROM_ADJACENT_LINE" in _codes(callout)


def test_footnote_lines_are_context_not_values(deck_output):
    assert not [v for v in deck_output.values if "2026" == (v.raw_text or "")]
    assert not _by(deck_output, ExtractionMethod.SLIDE_TEXT, page=1)


def test_hidden_slide_values_are_flagged(deck_output):
    hidden = _by(deck_output, ExtractionMethod.SLIDE_TEXT, page=5)
    assert [v.raw_number for v in hidden] == [1200.0]
    assert hidden[0].context["hidden"] is True and "HIDDEN_SLIDE" in _codes(hidden[0])
    assert all(not v.context.get("hidden") for v in deck_output.values if v.slide_or_page != 5)


def test_large_picture_becomes_placeholder_and_logo_is_skipped(deck_output):
    placeholders = _by(deck_output, ExtractionMethod.IMAGE_PLACEHOLDER)
    assert [v.slide_or_page for v in placeholders] == [6]
    image = placeholders[0]
    assert image.raw_number is None and image.raw_text is None and image.confidence == 0.0
    assert "IMAGE_NOT_PARSED" in _codes(image) and image.context["object_kind"] == "image"
    assert image.context["area_fraction"] > 0.5 and len(image.context["media_sha256"]) == 64
    assert image.locator["part"].startswith("ppt/media/")


def test_one_page_row_per_slide(deck_output):
    pages = {p.page_number: p for p in deck_output.pages}
    assert sorted(pages) == [1, 2, 3, 4, 5, 6]
    assert pages[1].title == "TRx trend" and pages[1].notes == "Illustrative data only."
    assert "Brand awareness: 62%" in pages[4].text
    assert pages[5].hidden is True and not pages[4].hidden
    assert pages[6].image_count == 1 and pages[2].image_count == 1
    assert (pages[1].page_width, pages[1].page_height) == (9144000.0, 6858000.0)
    kinds = [s["kind"] for s in pages[3].shapes]
    assert kinds.count("chart") == 2 and "title" in kinds
    assert all(len(s["bbox"]) == 4 for s in pages[3].shapes if s["bbox"])


def test_document_properties_and_versions(deck_output):
    props = deck_output.document_properties
    assert props["title"] == "Kestrel Monthly Brand Review" and props["page_count"] == 6
    assert props["created"].startswith("2026-06-03") and props["modified"].startswith("2026-06-04")
    assert deck_output.extractor_versions["pptx_extract"] == "1.1.0"
    assert deck_output.extractor_versions["pptx_extract_schema"] == "1.1"
    assert deck_output.issues == [] and not deck_output.failed


def test_field_contract(deck_output):
    ids = [v.value_id for v in deck_output.values]
    assert len(ids) == len(set(ids))
    for v in deck_output.values:
        assert set(v.locator) <= LOCATOR_KEYS, set(v.locator) - LOCATOR_KEYS
        assert set(v.context) <= CONTEXT_KEYS, set(v.context) - CONTEXT_KEYS
        assert v.slide_or_page == v.locator.get("page")
        for f in v.flags:
            assert set(f) == {"code", "dimension", "detail"} and f["dimension"] in {"extraction", "meaning", "association"}
        if v.method != ExtractionMethod.IMAGE_PLACEHOLDER:
            assert v.raw_text is not None and v.raw_number is not None


def test_store_roundtrip(deck_output, mem_store):
    doc_id = deck_output.values[0].doc_id
    mem_store.replace(ExtractedValue, {"doc_id": doc_id}, deck_output.values)
    mem_store.replace(DocumentPage, {"doc_id": doc_id}, deck_output.pages)
    back = {v.value_id: v for v in mem_store.read(ExtractedValue, "doc_id = :d", {"d": doc_id})}
    assert len(back) == len(deck_output.values)
    sample = deck_output.values[0]
    assert back[sample.value_id].locator == sample.locator and back[sample.value_id].flags == sample.flags
    assert mem_store.count(DocumentPage, "doc_id = :d", {"d": doc_id}) == 6


def test_rerun_gives_identical_ids(deck_bytes, deck_output):
    again = extract_pptx_values(sha256_bytes(deck_bytes), deck_bytes, "renamed.pptx")
    assert [v.value_id for v in again.values] == [v.value_id for v in deck_output.values]


def test_value_ids_survive_config_toggles(deck_bytes, deck_output):
    toggled = _extract(deck_bytes, config=px.ExtractionConfig(include_notes=False))
    assert {v.value_id for v in toggled.values} == {v.value_id for v in deck_output.values}
    assert {v.source_record_id for v in toggled.values} != {v.source_record_id for v in deck_output.values}


def test_hidden_slide_switch_is_respected(deck_bytes):
    out = _extract(deck_bytes, config=px.ExtractionConfig(include_hidden_slides=False))
    assert not [v for v in out.values if v.slide_or_page == 5]
    assert len(_by(out, ExtractionMethod.SLIDE_TEXT, page=4)) == 5


def test_sparse_cache_does_not_shift_values(deck_bytes):
    def patch(parts):
        root = X.fromstring(parts["ppt/charts/chart1.xml"])
        cache = root.find(".//c:ser/c:val/c:numRef/c:numCache", px.NS)
        cache.remove(cache.findall("c:pt", px.NS)[1])
        return {"ppt/charts/chart1.xml": X.tostring(root)}

    points = _by(_extract(rewrite_zip(deck_bytes, patch)), ExtractionMethod.NATIVE_CHART, page=1)
    assert [v.locator["point_index"] for v in points] == [0, 2, 3, 4, 5]
    assert points[1].raw_number == TRX[2] and points[1].context["category"] == MONTHS[2].isoformat()


def test_error_token_is_kept_without_a_number(deck_bytes):
    def patch(parts):
        root = X.fromstring(parts["ppt/charts/chart1.xml"])
        root.find(".//c:ser/c:val/c:numRef/c:numCache/c:pt/c:v", px.NS).text = "#N/A"
        return {"ppt/charts/chart1.xml": X.tostring(root)}

    points = _by(_extract(rewrite_zip(deck_bytes, patch)), ExtractionMethod.NATIVE_CHART, page=1)
    assert points[0].raw_text == "#N/A" and points[0].raw_number is None and points[0].confidence == 0.0
    assert "NONNUMERIC_CHART_VALUE" in _codes(points[0])
    assert all("NONNUMERIC_CHART_VALUE" not in _codes(v) for v in points[1:])


def test_malformed_relationship_loses_one_slide_not_the_deck(deck_bytes):
    def patch(parts):
        rels = parts["ppt/slides/_rels/slide4.xml.rels"].decode()
        rels = rels.replace("</Relationships>", '<Relationship Id="rIdBad" Type="http://schemas.openxmlformats.org/'
                            'officeDocument/2006/relationships/image" Target="..\\media\\bad.png"/></Relationships>')
        return {"ppt/slides/_rels/slide4.xml.rels": rels.encode()}

    out = _extract(rewrite_zip(deck_bytes, patch))
    assert any(i["code"] == "EXTRACTION_ISSUE" and i["page"] == 4 for i in out.issues)
    assert {v.raw_text for v in _by(out, ExtractionMethod.SLIDE_TEXT, page=4)} >= {"62%", "47"}
    assert len(_by(out, ExtractionMethod.NATIVE_CHART)) == 10 and len(out.pages) == 6


def test_limits_emit_issue_not_silent_truncation(deck_bytes):
    out = _extract(deck_bytes, config=px.ExtractionConfig(max_sheet_cells=2))
    assert any("max_sheet_cells" in i["detail"] for i in out.issues)
    points = _by(out, ExtractionMethod.NATIVE_CHART, page=1)
    assert len(points) == 6 and all("CACHE_NOT_COMPARED" in _codes(v) for v in points)


@pytest.mark.parametrize("blob, needle", [
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64, "OLE compound file"),
    (b"%PDF-1.7\n...", "PDF"),
    (b"not a deck", "not a ZIP"),
])
def test_sniff_before_parsing(blob, needle):
    out = _extract(blob)
    assert out.failed and out.issues[0]["code"] == "NOT_A_PPTX" and needle in out.issues[0]["detail"]


def test_unplotted_workbook_cells_are_values(deck_bytes):
    from openpyxl import load_workbook

    def patch(parts):
        part = "ppt/embeddings/Microsoft_Excel_Sheet1.xlsx"
        wb = load_workbook(io.BytesIO(parts[part]))
        ws = wb["Sheet1"]
        ws["D1"], ws["D2"] = "Plan", 6000
        out = io.BytesIO()
        wb.save(out)
        return {part: out.getvalue()}

    out = _extract(rewrite_zip(deck_bytes, patch))
    cells = _by(out, ExtractionMethod.EMBEDDED_WORKSHEET)
    assert [(v.locator["sheet"], v.locator["cell"], v.raw_number) for v in cells] == [("Sheet1", "D2", 6000.0)]
    assert cells[0].context["col_header"] == "Plan" and "WORKBOOK_UNPLOTTED" in _codes(cells[0])
    assert cells[0].context["object_kind"] == "worksheet" and cells[0].slide_or_page == 1
    assert all("CACHE_MATCHES_WORKBOOK" in _codes(v) for v in _by(out, ExtractionMethod.NATIVE_CHART, page=1))


def test_multilevel_categories_forward_fill():
    from datetime import datetime

    from kpi_evidence.extract.pptx import _Run, chart_values

    xml = f'''<c:chartSpace xmlns:c="{px.NS['c']}"><c:chart><c:plotArea><c:barChart><c:ser>
    <c:tx><c:v>Sales</c:v></c:tx><c:cat><c:multiLvlStrRef><c:f>Sheet1!A1:B2</c:f><c:multiLvlStrCache><c:ptCount val="2"/>
    <c:lvl><c:pt idx="0"><c:v>Q1</c:v></c:pt><c:pt idx="1"><c:v>Q2</c:v></c:pt></c:lvl>
    <c:lvl><c:pt idx="0"><c:v>2025</c:v></c:pt></c:lvl></c:multiLvlStrCache></c:multiLvlStrRef></c:cat>
    <c:val><c:numLit><c:ptCount val="2"/><c:pt idx="0"><c:v>1.20</c:v></c:pt><c:pt idx="1"><c:v>3</c:v></c:pt></c:numLit></c:val>
    </c:ser></c:barChart></c:plotArea></c:chart></c:chartSpace>'''
    data, _w = px._chart_data(X.fromstring(xml), px.ExtractionConfig())
    record = {"kind": "chart", "record_id": "r", "slide_number": 1, "slide_title": "Sales", "source_part": "ppt/charts/c.xml",
              "shape_id": "2", "shape_name": "Chart 1", "chart_index_on_slide": 1, "bounding_box": None,
              "position_on_slide": "unknown", "description_evidence": {}, "uncertainty_reasons": [], "data": data}
    run = _Run(doc_id="d", run_id=None, created_at=datetime(2026, 1, 1), version="t", scans={}, slide_area=None)
    values = chart_values(record, run, {}, {})
    assert [v.raw_text for v in values] == ["1.20", "3"]
    assert [v.context["category"] for v in values] == ["Q1", "Q2"]
    assert [v.context["category_levels"] for v in values] == [["2025", "Q1"], ["2025", "Q2"]]


def test_formula_cells():
    assert formula_cells("Sheet1!$B$2:$B$4") == [("Sheet1", "B2"), ("Sheet1", "B3"), ("Sheet1", "B4")]
    assert formula_cells("'My Sheet'!$A$1") == [("My Sheet", "A1")]
    assert formula_cells("(Sheet1!$B$2,Sheet1!$B$4)") == [("Sheet1", "B2"), ("Sheet1", "B4")]
    assert formula_cells("not a ref") == [] and formula_cells(None) == []


def test_deck_without_charts_or_text_still_has_pages():
    from pptx import Presentation

    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    buf = io.BytesIO()
    deck.save(buf)
    out = _extract(buf.getvalue())
    assert out.values == [] and len(out.pages) == 1 and not out.failed


def test_build_deck_is_deterministic_enough_for_ids():
    blob = build_deck()
    out = _extract(blob)
    assert len(out.values) == len({v.value_id for v in out.values})
