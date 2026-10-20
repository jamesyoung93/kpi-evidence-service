from __future__ import annotations

import io

import pytest

from kpi_evidence.contracts import ExtractionMethod
from kpi_evidence.extract import parse_number_parts, split_reference_mark
from kpi_evidence.extract._common import find_labeled_numbers
from kpi_evidence.extract.pdf import extract_pdf_values
from kpi_evidence.extract.pptx import extract_pptx_values
from kpi_evidence.extract.xlsx import extract_xlsx_values
from kpi_evidence.ids import sha256_bytes


def _codes(value):
    return [f["code"] for f in value.flags]


def _superscript_deck() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "KPI summary"
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(1.5), Inches(8), Inches(1)).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "May 2026"
    table.cell(1, 0).text = "NBRx"
    table.cell(2, 0).text = "TRx share"
    for r, number in ((1, "412"), (2, "9.1%")):
        paragraph = table.cell(r, 1).text_frame.paragraphs[0]
        paragraph.add_run().text = number
        mark = paragraph.add_run()
        mark.text = "1"
        mark.font._rPr.set("baseline", "30000")
    box = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(8), Inches(0.5))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.add_run().text = "New writers 87"
    mark = paragraph.add_run()
    mark.text = "2"
    mark.font._rPr.set("baseline", "30000")
    paragraph.add_run().text = " | Calls 1,204"
    box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(8), Inches(0.5))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.add_run().text = "Area 12 m"
    unit = paragraph.add_run()
    unit.text = "2"
    unit.font._rPr.set("baseline", "30000")
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def test_pptx_superscript_marks_are_not_glued_to_numbers():
    blob = _superscript_deck()
    out = extract_pptx_values(sha256_bytes(blob), blob, "deck.pptx")
    table = {v.context["row_label"]: v for v in out.values if v.method == ExtractionMethod.NATIVE_TABLE}
    assert (table["NBRx"].raw_text, table["NBRx"].raw_number) == ("412", 412.0)
    assert (table["TRx share"].raw_text, table["TRx share"].raw_number) == ("9.1%", 9.1)
    assert all("FOOTNOTE_MARK" in _codes(v) for v in table.values())
    text = {v.context["row_label"]: v for v in out.values if v.method == ExtractionMethod.SLIDE_TEXT}
    assert (text["New writers"].raw_text, text["New writers"].raw_number) == ("87", 87.0)
    assert "FOOTNOTE_MARK" in _codes(text["New writers"])
    assert text["Calls"].raw_number == 1204.0 and "FOOTNOTE_MARK" not in _codes(text["Calls"])
    assert text["New writers"].context["text_line"] == "New writers 87 | Calls 1,204"
    assert text["Area"].raw_number == 12.0 and text["Area"].context["text_line"] == "Area 12 m2"
    assert "FOOTNOTE_MARK" not in _codes(text["Area"])


def test_pdf_raised_small_mark_is_not_glued_to_the_number():
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    pdf = canvas.Canvas(buf, pagesize=letter)
    width, height = letter
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(72, height - 72, "Kestrel Executive Summary")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(72, height - 120, "NBRx 412")
    pdf.setFont("Helvetica", 7)
    pdf.drawString(72 + pdf.stringWidth("NBRx 412", "Helvetica", 11), height - 116, "1")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(72, height - 140, "TRx 5,910")
    pdf.showPage()
    pdf.save()
    blob = buf.getvalue()
    out = extract_pdf_values(sha256_bytes(blob), blob, "summary.pdf")
    values = {v.context["row_label"]: v for v in out.values if v.method == ExtractionMethod.PDF_TEXT}
    assert (values["NBRx"].raw_text, values["NBRx"].raw_number) == ("412", 412.0)
    assert "FOOTNOTE_MARK" in _codes(values["NBRx"])
    assert values["TRx"].raw_number == 5910.0 and "FOOTNOTE_MARK" not in _codes(values["TRx"])


@pytest.mark.parametrize("line", ["TRx 5,910 (1)", "TRx 5,910 [1]", "TRx 5,910(1)", "TRx 5,910 (12)", "TRx 5,910¹",
                                  "TRx 5,910⁴"])
def test_reference_marks_after_a_number_are_not_values(line):
    hits = find_labeled_numbers(line)
    assert [(h.parts.number, h.label) for h in hits] == [(5910.0, "TRx")]


def test_parenthesized_values_without_a_number_before_them_are_kept():
    assert [(h.parts.number, h.label) for h in find_labeled_numbers("Net change (1,234)")] == [(-1234.0, "Net change")]
    assert [h.parts.number for h in find_labeled_numbers("Open issues (3)")] == [-3.0]
    assert [h.parts.number for h in find_labeled_numbers("TRx 5,910 (0.3%)")] == [5910.0, -0.3]


@pytest.mark.parametrize("text, core, mark", [
    ("5,910 (1)", "5,910", "(1)"), ("9.1% [2]", "9.1%", "[2]"), ("(1,234) (1)", "(1,234)", "(1)"),
    ("(1)", "(1)", None), ("5,910", "5,910", None), ("TRx (1)", "TRx (1)", None),
])
def test_split_reference_mark(text, core, mark):
    assert split_reference_mark(text) == (core, mark)


def test_cell_numbers_with_reference_marks_parse():
    assert parse_number_parts("5,910 (1)").number == 5910.0
    assert parse_number_parts("9.1% (1)").kind == "percent"
    assert parse_number_parts("(1)").number == -1.0


def test_worksheet_text_cell_with_reference_mark():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "KPIs"
    ws.append(["Metric", "May 2026"])
    ws.append(["TRx", "5,910 (1)"])
    ws.append(["TRx share", "9.1% [2]"])
    buf = io.BytesIO()
    wb.save(buf)
    blob = buf.getvalue()
    out = extract_xlsx_values(sha256_bytes(blob), blob, "kpis.xlsx")
    cells = {v.locator["cell"]: v for v in out.values}
    assert (cells["B2"].raw_text, cells["B2"].raw_number) == ("5,910", 5910.0)
    assert (cells["B3"].raw_text, cells["B3"].raw_number) == ("9.1%", 9.1)
    assert all("FOOTNOTE_MARK" in _codes(v) for v in cells.values())
