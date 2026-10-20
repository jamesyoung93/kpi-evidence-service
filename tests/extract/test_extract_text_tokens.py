from __future__ import annotations

import io

import pytest

from kpi_evidence.contracts import ExtractionMethod
from kpi_evidence.extract import parse_number_parts
from kpi_evidence.extract._common import find_labeled_numbers, find_numbers, is_footnote, is_note_line
from kpi_evidence.extract.pdf import extract_pdf_values
from kpi_evidence.extract.pptx import extract_pptx_values
from kpi_evidence.ids import sha256_bytes

LINES = ["Net sales $12.3MM", "Gross sales $1.2BN", "Revenue reached $12.3 million", "Referrals YTD 147",
         "Q2 TRx of 17,220 (data through May 2026)", "TRx 5,910 (source: IQVIA)"]


@pytest.mark.parametrize("line, raw, number, suffix, label", [
    ("Net sales $12.3MM", "$12.3MM", 12.3, "MM", "Net sales"),
    ("Net sales $1.2BN", "$1.2BN", 1.2, "BN", "Net sales"),
    ("Net sales $12.3 MM", "$12.3 MM", 12.3, "MM", "Net sales"),
    ("Net sales $5Mn", "$5Mn", 5.0, "MN", "Net sales"),
    ("Net sales reached $12.3 million", "$12.3 million", 12.3, "M", "Net sales reached"),
    ("Net sales of $1.2 billion", "$1.2 billion", 1.2, "B", "Net sales of"),
    ("Patients 45 thousand", "45 thousand", 45.0, "K", "Patients"),
    ("Referrals YTD 147", "147", 147.0, "", "Referrals YTD"),
    ("Sales MAT 1240", "1240", 1240.0, "", "Sales MAT"),
    ("Spread 50 BPS", "50 BPS", 50.0, "", "Spread"),
])
def test_text_number_forms(line, raw, number, suffix, label):
    hits = find_labeled_numbers(line)
    assert [(h.raw_text, h.parts.number, h.parts.suffix, h.label) for h in hits] == [(raw, number, suffix, label)]


@pytest.mark.parametrize("line", ["Referrals YTD 25", "YTD '26 plan", "FY 26 plan", "MAT 24"])
def test_two_digit_years_after_period_words_are_not_values(line):
    assert find_numbers(line) == []


def test_scale_words_and_case_in_cell_parser():
    parts = parse_number_parts("$12.3 million")
    assert (parts.number, parts.suffix, parts.multiplier, parts.currency) == (12.3, "M", 1e6, "$")
    assert parse_number_parts("1.2 BILLION").multiplier == 1e9 and parse_number_parts("45 thousand").suffix == "K"
    assert parse_number_parts("12.3MM").multiplier == 1e6


def test_inline_notes_do_not_hide_the_line():
    assert is_footnote("Q2 TRx of 17,220 (data through May 2026)")
    assert not is_note_line("Q2 TRx of 17,220 (data through May 2026)")
    assert not is_note_line("TRx 5,910 (source: IQVIA)")
    for note in ("Source: IQVIA NPA, data through Mar 2026", "Data through May 2026", "* excludes samples",
                 "1 New-to-brand patients. Source: account feed.", "Notes: (1) excludes samples"):
        assert is_note_line(note)


def _expected():
    return {"Net sales $12.3MM": ("$12.3MM", 12.3), "Gross sales $1.2BN": ("$1.2BN", 1.2),
            "Revenue reached $12.3 million": ("$12.3 million", 12.3), "Referrals YTD 147": ("147", 147.0),
            "Q2 TRx of 17,220 (data through May 2026)": ("17,220", 17220.0),
            "TRx 5,910 (source: IQVIA)": ("5,910", 5910.0)}


def test_slide_text_lines_produce_values():
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Executive summary"
    box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(4))
    box.text_frame.text = LINES[0]
    for line in LINES[1:]:
        box.text_frame.add_paragraph().text = line
    buf = io.BytesIO()
    deck.save(buf)
    blob = buf.getvalue()
    out = extract_pptx_values(sha256_bytes(blob), blob, "deck.pptx")
    found = {v.context["text_line"]: (v.raw_text, v.raw_number) for v in out.values
             if v.method == ExtractionMethod.SLIDE_TEXT}
    assert found == _expected()


def test_pdf_text_lines_produce_values():
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    pdf = canvas.Canvas(buf, pagesize=letter)
    width, height = letter
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(72, height - 72, "Executive summary")
    pdf.setFont("Helvetica", 11)
    for i, line in enumerate(LINES):
        pdf.drawString(72, height - 120 - 20 * i, line)
    pdf.drawString(72, 60, "Source: account feed. Data through May 2026")
    pdf.showPage()
    pdf.save()
    blob = buf.getvalue()
    out = extract_pdf_values(sha256_bytes(blob), blob, "summary.pdf")
    found = {v.context["text_line"]: (v.raw_text, v.raw_number) for v in out.values if v.method == ExtractionMethod.PDF_TEXT}
    assert found == _expected()
    assert "Source: account feed. Data through May 2026" in out.values[0].context["footnotes"]
