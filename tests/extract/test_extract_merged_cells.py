from __future__ import annotations

import io

import pytest

from kpi_evidence.contracts import ExtractionMethod
from kpi_evidence.extract import pptx as pptx_route
from kpi_evidence.extract.pptx import extract_pptx_values
from kpi_evidence.ids import sha256_bytes


def _deck(merge: str | None) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    deck = Presentation()
    layout = deck.slide_layouts[5]
    slide = deck.slides.add_slide(layout)
    slide.shapes.title.text = "TRx trend"
    data = CategoryChartData(number_format="#,##0.0")
    data.categories = ["Mar 2026", "Apr 2026", "May 2026"]
    data.add_series("TRx", [6259.98, 5700.91, 5736.36])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(4), data)
    slide = deck.slides.add_slide(layout)
    slide.shapes.title.text = "KPI summary"
    table = slide.shapes.add_table(3, 3, Inches(1), Inches(1.5), Inches(8), Inches(2)).table
    if merge == "columns":
        table.cell(0, 1).merge(table.cell(0, 2))
    elif merge == "rows":
        table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "May 2026"
    table.cell(1, 1).text = "Actual"
    table.cell(1, 2).text = "Plan"
    table.cell(2, 0).text = "TRx"
    table.cell(2, 1).text = "5,736"
    table.cell(2, 2).text = "5,900"
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def _extract(blob):
    return extract_pptx_values(sha256_bytes(blob), blob, "deck.pptx")


@pytest.mark.parametrize("merge", [None, "columns", "rows"])
def test_merged_table_cells_keep_every_value(merge):
    out = _extract(_deck(merge))
    charts = [v for v in out.values if v.method == ExtractionMethod.NATIVE_CHART]
    cells = {(v.locator["row"], v.locator["col"]): v for v in out.values if v.method == ExtractionMethod.NATIVE_TABLE}
    assert [v.raw_number for v in charts] == [6259.98, 5700.91, 5736.36]
    assert cells[(2, 1)].raw_number == 5736.0 and cells[(2, 2)].raw_number == 5900.0
    assert not [i for i in out.issues if i["severity"] == "error"]


def test_merged_header_text_is_carried_across_the_span():
    out = _extract(_deck("columns"))
    cells = {(v.locator["row"], v.locator["col"]): v for v in out.values if v.method == ExtractionMethod.NATIVE_TABLE}
    assert cells[(2, 1)].context["col_header"] == "May 2026 / Actual"
    assert cells[(2, 2)].context["col_header"] == "May 2026 / Plan"


def test_one_failing_table_becomes_an_issue_and_charts_survive(monkeypatch):
    def broken(rec, run):
        raise KeyError("gridSpan")

    monkeypatch.setattr(pptx_route, "table_values", broken)
    out = _extract(_deck("columns"))
    assert [v.raw_number for v in out.values if v.method == ExtractionMethod.NATIVE_CHART] == [6259.98, 5700.91, 5736.36]
    failed = [i for i in out.issues if i["code"] == "OBJECT_EXTRACTION_FAILED"]
    assert len(failed) == 1 and failed[0]["page"] == 2 and failed[0]["severity"] == "warning"
    assert "KeyError" in failed[0]["detail"] and not out.failed
