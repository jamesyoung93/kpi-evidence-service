from __future__ import annotations

import io

from kpi_evidence.contracts import ExtractionMethod
from kpi_evidence.extract.pptx import extract_pptx_values
from kpi_evidence.ids import sha256_bytes


def _deck(chart_type: str, series: dict[str, list[float]], axis_format: str | None) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "TRx by payer channel"
    data = CategoryChartData()
    data.categories = ["Apr 2026", "May 2026"]
    for name, values in series.items():
        data.add_series(name, values)
    chart = slide.shapes.add_chart(getattr(XL_CHART_TYPE, chart_type), Inches(1), Inches(1.5), Inches(8), Inches(4),
                                   data).chart
    if axis_format:
        chart.value_axis.tick_labels.number_format = axis_format
        chart.value_axis.tick_labels.number_format_is_linked = False
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def _points(blob):
    out = extract_pptx_values(sha256_bytes(blob), blob, "deck.pptx")
    return [v for v in out.values if v.method == ExtractionMethod.NATIVE_CHART]


COUNTS = {"Commercial": [3200, 3300], "Medicare Part D": [1500, 1450]}
SHARES = {"Commercial": [0.68, 0.69], "Medicare Part D": [0.32, 0.31]}


def test_counts_in_a_100_percent_stack_ignore_the_percent_axis():
    points = _points(_deck("COLUMN_STACKED_100", COUNTS, "0%"))
    assert [v.raw_number for v in points] == [3200.0, 3300.0, 1500.0, 1450.0]
    assert all(v.number_format is None and "number_format_source" not in v.context for v in points)
    assert all("PERCENT_AXIS_NOT_APPLIED" in [f["code"] for f in v.flags] for v in points)


def test_shares_in_a_100_percent_stack_keep_the_percent_axis():
    points = _points(_deck("COLUMN_STACKED_100", SHARES, "0%"))
    assert all(v.number_format == "0%" and v.context["number_format_source"] == "axis" for v in points)
    assert not any("PERCENT_AXIS_NOT_APPLIED" in [f["code"] for f in v.flags] for v in points)


def test_other_charts_still_inherit_the_percent_axis():
    points = _points(_deck("COLUMN_CLUSTERED", {"Attainment": [1.6, 1.8]}, "0%"))
    assert all(v.number_format == "0%" and v.context["number_format_source"] == "axis" for v in points)
    points = _points(_deck("COLUMN_STACKED", COUNTS, "#,##0"))
    assert all(v.number_format == "#,##0" for v in points)
