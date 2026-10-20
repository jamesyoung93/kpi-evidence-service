from __future__ import annotations

import io
from datetime import datetime
from xml.etree import ElementTree as X

from kpi_evidence.contracts import ExtractionMethod
from kpi_evidence.extract._vendor import pptx_extract as px
from kpi_evidence.extract.pptx import _Run, chart_values, extract_pptx_values
from kpi_evidence.ids import sha256_bytes


def _deck() -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    deck = Presentation()
    layout = deck.slide_layouts[5]
    slide = deck.slides.add_slide(layout)
    slide.shapes.title.text = "TRx monthly trend"
    data = CategoryChartData(number_format="#,##0")
    values = []
    for year, months in (("2025", ["Nov", "Dec"]), ("2026", ["Jan", "Feb"])):
        category = data.add_category(year)
        for month in months:
            category.add_sub_category(month)
            values.append(5000 + 10 * len(values))
    data.add_series("TRx", values)
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(4), data)
    slide = deck.slides.add_slide(layout)
    slide.shapes.title.text = "TRx by region"
    data = CategoryChartData(number_format="#,##0")
    for region in ("West", "East"):
        category = data.add_category(region)
        for month in ("Apr 2026", "May 2026"):
            category.add_sub_category(month)
    data.add_series("TRx", [1210, 1265, 1500, 1520])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(4), data)
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def _points(out, page):
    return [v for v in out.values if v.method == ExtractionMethod.NATIVE_CHART and v.slide_or_page == page]


def test_year_and_region_levels_from_outer_to_inner():
    blob = _deck()
    out = extract_pptx_values(sha256_bytes(blob), blob, "deck.pptx")
    trend = _points(out, 1)
    assert [v.context["category"] for v in trend] == ["Nov", "Dec", "Jan", "Feb"]
    assert [v.context["category_levels"] for v in trend] == [["2025", "Nov"], ["2025", "Dec"], ["2026", "Jan"],
                                                             ["2026", "Feb"]]
    regions = _points(out, 2)
    assert [v.context["category_levels"] for v in regions] == [["West", "Apr 2026"], ["West", "May 2026"],
                                                               ["East", "Apr 2026"], ["East", "May 2026"]]
    assert all(v.context["category"] == v.context["category_levels"][-1] for v in trend + regions)
    assert all(isinstance(level, str) for v in trend + regions for level in v.context["category_levels"])


def test_single_level_categories_carry_no_levels():
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    data = CategoryChartData()
    data.categories = ["Apr 2026", "May 2026"]
    data.add_series("TRx", [1, 2])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(4), data)
    buf = io.BytesIO()
    deck.save(buf)
    blob = buf.getvalue()
    out = extract_pptx_values(sha256_bytes(blob), blob, "deck.pptx")
    assert all("category_levels" not in v.context for v in out.values)


def test_missing_inner_label_is_not_copied_from_the_previous_point():
    xml = f'''<c:chartSpace xmlns:c="{px.NS['c']}"><c:chart><c:plotArea><c:barChart><c:ser>
    <c:tx><c:v>TRx</c:v></c:tx><c:cat><c:multiLvlStrRef><c:f>Sheet1!A2:B4</c:f><c:multiLvlStrCache><c:ptCount val="3"/>
    <c:lvl><c:pt idx="0"><c:v>Apr</c:v></c:pt><c:pt idx="2"><c:v>Jan</c:v></c:pt></c:lvl>
    <c:lvl><c:pt idx="0"><c:v>2025</c:v></c:pt><c:pt idx="2"><c:v>2026</c:v></c:pt></c:lvl></c:multiLvlStrCache>
    </c:multiLvlStrRef></c:cat>
    <c:val><c:numLit><c:ptCount val="3"/><c:pt idx="0"><c:v>10</c:v></c:pt><c:pt idx="1"><c:v>20</c:v></c:pt>
    <c:pt idx="2"><c:v>30</c:v></c:pt></c:numLit></c:val>
    </c:ser></c:barChart></c:plotArea></c:chart></c:chartSpace>'''
    data, _warnings = px._chart_data(X.fromstring(xml), px.ExtractionConfig())
    record = {"kind": "chart", "record_id": "r", "slide_number": 1, "slide_title": "TRx", "source_part": "ppt/charts/c.xml",
              "shape_id": "2", "shape_name": "Chart 1", "chart_index_on_slide": 1, "bounding_box": None,
              "position_on_slide": "unknown", "description_evidence": {}, "uncertainty_reasons": [], "data": data}
    run = _Run(doc_id="d", run_id=None, created_at=datetime(2026, 1, 1), version="t", scans={}, slide_area=None)
    values = chart_values(record, run, {}, {})
    assert [v.context.get("category") for v in values] == ["Apr", None, "Jan"]
    assert [v.context["category_levels"] for v in values] == [["2025", "Apr"], ["2025", ""], ["2026", "Jan"]]
    assert "MISSING_CHART_POINTS" in [f["code"] for f in values[1].flags]
    assert "MISSING_CHART_POINTS" not in [f["code"] for f in values[0].flags + values[2].flags]
