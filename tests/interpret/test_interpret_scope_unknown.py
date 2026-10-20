from __future__ import annotations

import io

import pytest
from interp_builders import DIMS, chart_point, codes, make_doc, table_cell

from kpi_evidence.contracts import ExtractionMethod, ReasonCode
from kpi_evidence.reconcile import extraction_gate
from kpi_evidence.standardize import interpret

FOOTER = ["Data through May 2026"]


def run(values, *, registry=None, dims=DIMS):
    return interpret(make_doc(), values, [], registry, None, dims)


def one(value, **kw):
    [obs] = run([value], **kw)
    return obs


def scope_flag(obs):
    return next(f for f in obs.flags if f["code"] == "SCOPE_UNKNOWN")


def test_region_category_without_dimension_values_is_scope_unknown(registry):
    obs = one(chart_point(0.881, category="West", series_name="Call attainment", number_format="0.0%",
                          object_title="Call attainment by region", footnotes=FOOTER), registry=registry, dims=None)
    assert obs.scope == {}
    flag = scope_flag(obs)
    assert flag["dimension"] == "association" and "'West'" in flag["detail"]


def test_region_spelled_differently_from_the_platform_is_scope_unknown(registry):
    obs = one(chart_point(0.881, category="Western", series_name="Call attainment", number_format="0.0%",
                          footnotes=FOOTER), registry=registry)
    assert obs.scope == {} and "SCOPE_UNKNOWN" in codes(obs)


def test_region_column_header_without_dimension_values_is_scope_unknown(registry):
    obs = one(table_cell("88.1%", row_label="Call attainment", col_header="West", footnotes=FOOTER),
              registry=registry, dims=None)
    assert obs.label == "Call attainment" and obs.scope == {}
    assert "'West'" in scope_flag(obs)["detail"]


def test_unknown_second_row_label_is_scope_unknown(registry):
    obs = one(table_cell("1,265", row_label="TRx", row_labels=["TRx", "Retail"], col_header="May 2026"),
              registry=registry)
    assert obs.label == "TRx" and obs.scope == {}
    assert "'Retail'" in scope_flag(obs)["detail"]


def test_competitor_category_is_scope_unknown_but_own_brand_is_not(registry):
    own, other = run([chart_point(0.091, category="Kestrel", series_name="TRx share", number_format="0.0%"),
                      chart_point(0.262, category="Competitor A", series_name="TRx share", number_format="0.0%")],
                     registry=registry)
    assert "SCOPE_UNKNOWN" not in codes(own)
    assert "SCOPE_UNKNOWN" in codes(other)


@pytest.mark.parametrize("category", ["Total", "National", "All regions", "May 2026", "Q2 2026", "Actual", "Plan",
                                      "NRx", "Total US", "2025"])
def test_total_period_generic_and_kpi_categories_are_not_scope_unknown(registry, category):
    obs = one(chart_point(410.0, category=category, series_name="NRx", footnotes=FOOTER), registry=registry, dims=None)
    assert "SCOPE_UNKNOWN" not in codes(obs)


@pytest.mark.parametrize("header", ["Value", "Latest month", "vs PY %", "YTD", "May 2026 / Actual", "Prior year",
                                    "Current month", "Q1 2026", "Target", "TRx (000s)"])
def test_period_and_qualifier_column_headers_are_not_scope_unknown(registry, header):
    obs = one(table_cell("5,910", row_label="TRx", col_header=header, footnotes=FOOTER), registry=registry, dims=None)
    assert "SCOPE_UNKNOWN" not in codes(obs)


def test_metric_row_labels_are_not_scope_unknown(registry):
    for row in ("Calls (details delivered)", "Call attainment (writers only)", "IL-17 class share", "Total"):
        obs = one(table_cell("1,213", row_label=row, col_header="May 2026"), registry=registry, dims=None)
        assert "SCOPE_UNKNOWN" not in codes(obs), row


def test_known_region_category_sets_scope_without_flag(registry):
    obs = one(chart_point(0.772, category="West", series_name="Call attainment", number_format="0.0%"),
              registry=registry)
    assert obs.scope == {"region": "West"} and "SCOPE_UNKNOWN" not in codes(obs)


def test_breakdown_title_without_any_scope_is_scope_unknown(registry):
    obs = one(chart_point(1265.0, category="May 2026", series_name="Western", object_title="TRx by region"),
              registry=registry)
    assert obs.scope == {}
    assert "by region" in scope_flag(obs)["detail"]


def test_breakdown_title_with_an_explicit_total_category_is_not_flagged(registry):
    obs = one(chart_point(5736.0, category="Total", series_name="TRx", object_title="TRx by region"),
              registry=registry)
    assert obs.scope == {} and "SCOPE_UNKNOWN" not in codes(obs)


def test_scope_unknown_value_is_held_back_from_reconciliation(registry):
    obs = one(chart_point(0.881, category="West", series_name="Call attainment", number_format="0.0%",
                          footnotes=FOOTER), registry=registry, dims=None)
    assert extraction_gate(obs, ExtractionMethod.NATIVE_CHART) == ReasonCode.SCOPE_UNKNOWN


def _regional_deck(categories: list[str], values: list[float]) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Kestrel field performance - May 2026"
    data = CategoryChartData(number_format="0.0%")
    data.categories = categories
    data.add_series("Call attainment", values, number_format="0.0%")
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(8), Inches(4), data)
    frame.chart.has_title = True
    frame.chart.chart_title.text_frame.text = "Call attainment by region"
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def test_unknown_region_matching_the_national_value_is_not_verified():
    pytest.importorskip("pptx")
    from kpi_evidence import ReportCheck

    check = ReportCheck(load_demo_platform=True)
    check.add_reports([("kestrel_field_may.pptx", _regional_deck(["Western", "Eastern"], [0.880517, 0.9142]))],
                      brand="KESTREL", report_type="monthly_brand_review", period="May 2026")
    rows = check.run(materialize=False).comparisons
    rows = rows[rows["kpi_id"] == "call_attainment"]
    assert len(rows) == 2
    assert set(rows["status"]) == {"UNCERTAIN_EXTRACTION"}
    assert rows["reason"].str.contains("scope").all()
