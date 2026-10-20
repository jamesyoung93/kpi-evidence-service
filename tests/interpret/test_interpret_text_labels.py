from __future__ import annotations

import io
from datetime import date, datetime

import pytest
from interp_builders import DIMS, make_doc, text_value

from kpi_evidence.contracts import Document, DocumentFormat, DocumentStatus
from kpi_evidence.ids import sha256_bytes
from kpi_evidence.standardize import interpret


def one(value, *, registry=None, dims=DIMS):
    [obs] = interpret(make_doc(), [value], [], registry, None, dims)
    return obs


@pytest.mark.parametrize("line,raw,before,after,label", [
    ("NBRx 412 | New writers 12", "12", "New writers", None, "New writers"),
    ("Brand awareness 62% | Trial 2%", "2%", "Trial", None, "Trial"),
    ("Call attainment 84.4% | Reach 4.4%", "4.4%", "Reach", None, "Reach"),
    ("TRx 5,736 | Referrals 36", "36", "Referrals", None, "Referrals"),
    ("5,910 TRx | 412 NBRx | 62% brand awareness", "412", None, "NBRx", "NBRx"),
    ("5,910 TRx | 412 NBRx | 62% brand awareness", "62%", None, "brand awareness", "brand awareness"),
])
def test_text_value_uses_the_extraction_labels(line, raw, before, after, label):
    obs = one(text_value(raw, line, label_before=before, label_after=after, row_label=before or after))
    assert obs.label == label


@pytest.mark.parametrize("line,raw,label", [
    ("Call attainment 84.4% | Reach 4.4%", "4.4%", "Reach"),
    ("TRx 5,736 | Referrals 36", "36", "Referrals"),
    ("NBRx 412 | New writers 12", "12", "New writers"),
    ("Brand awareness 62% | Trial 2%", "2%", "Trial"),
    ("5,910 TRx | 412 NBRx", "412", "NBRx"),
    ("TRx 17,220 -31.6%", "-31.6%", "TRx"),
])
def test_text_value_without_extraction_labels_reads_its_own_segment(line, raw, label):
    assert one(text_value(raw, line)).label == label


def test_scope_comes_from_the_value_segment_and_metric_from_the_line():
    obs = one(text_value("3,910", "TRx West 13,910 | East 3,910", label_before="East", row_label="East"))
    assert obs.label == "TRx" and obs.scope == {"region": "East"}


def test_text_values_from_a_real_slide_keep_their_own_labels(registry):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    from kpi_evidence.extract import extract_document

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Highlights"
    box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(3)).text_frame
    box.text = "NBRx 412 | New writers 12"
    box.add_paragraph().text = "Call attainment 84.4% | Reach 4.4%"
    box.add_paragraph().text = "5,910 TRx | 412 NBRx"
    buf = io.BytesIO()
    deck.save(buf)
    blob = buf.getvalue()
    now = datetime(2026, 6, 5, 12)
    doc = Document(doc_id=sha256_bytes(blob), file_name="highlights.pptx", file_path="/tmp/highlights.pptx",
                   size_bytes=len(blob), format=DocumentFormat.PPTX, status=DocumentStatus.READY, registered_at=now,
                   updated_at=now, brand="KESTREL", report_type="monthly_brand_review", period_label="May 2026",
                   period_start=date(2026, 5, 1), period_end=date(2026, 5, 31), publication_date=date(2026, 6, 1))
    out = extract_document(None, doc, blob)
    observations = interpret(doc, out.values, out.pages, registry, None, DIMS)
    labels = {(v.context.get("text_line"), v.raw_text): o.label for v, o in zip(out.values, observations, strict=True)}
    assert labels[("NBRx 412 | New writers 12", "12")] == "New writers"
    assert labels[("Call attainment 84.4% | Reach 4.4%", "4.4%")] == "Reach"
    assert labels[("5,910 TRx | 412 NBRx", "412")] == "NBRx"
    assert labels[("5,910 TRx | 412 NBRx", "5,910")] == "TRx"
