from __future__ import annotations

import io
from datetime import date

import pandas as pd
import pytest

from kpi_evidence.contracts import ValidationStatus as S
from kpi_evidence.exports import (SCORECARD_COLUMNS, locator_text, scorecard_csv, scorecard_frame, scorecard_pptx,
                                  scorecard_xlsx)
from kpi_evidence.review import decide
from kpi_evidence.serving import add_months, month_end


@pytest.fixture()
def published(seeder, store, registry, settings):
    deck = seeder.doc("Kestrel_Monthly_Brand_Review_2026-06.pptx")
    trx = seeder.seed("trx", S.VERIFIED, doc=deck, page=2)
    share = seeder.seed("trx_share", S.VERIFIED, doc=deck, page=3, unit="fraction", report_value=0.092,
                        platform_value=0.0917, fingerprint="fp-share")
    samples = seeder.seed("samples", S.VARIANCE_REVIEW, doc=deck, page=4, report_value=1060, platform_value=1000,
                          fingerprint="fp-samples")
    aware = seeder.seed("brand_awareness", S.REPORT_ONLY, doc=deck, page=6, unit="fraction", report_value=0.62,
                        platform_value=None, fingerprint="fp-mr")
    decide(store, registry, settings, trx.result.result_id, "ACCEPT", "", "rev@example.com")
    decide(store, registry, settings, share.result.result_id, "ACCEPT", "", "rev@example.com")
    decide(store, registry, settings, samples.result.result_id, "ACCEPT", "Field counts include drop-offs",
           "rev@example.com")
    decide(store, registry, settings, aware.result.result_id, "ACCEPT", "Survey wave 2", "rev@example.com")
    return {"trx": trx, "share": share, "samples": samples, "aware": aware}


def test_scorecard_frame(published, store, registry):
    frame = scorecard_frame(store, registry, "KESTREL")
    assert list(frame.columns) == list(SCORECARD_COLUMNS) and len(frame) == 4
    assert list(frame["kpi_id"]) == ["trx", "trx_share", "samples", "brand_awareness"]
    trx = frame.iloc[0]
    assert trx["status"] == "Matches platform" and trx["scope"] == "Brand total" and trx["period"] == "May 2026"
    assert trx["evidence"].startswith("Kestrel_Monthly_Brand_Review_2026-06.pptx, slide 2, chart 'TRx by month'")
    assert "series 'TRx'" in trx["evidence"]
    assert trx["caveat"].startswith("Matches the platform value") and "0.5%" in trx["caveat"]
    samples = frame[frame["kpi_id"] == "samples"].iloc[0]
    assert samples["status"] == "Accepted variance"
    assert samples["caveat"] == ("Differs from the platform value (1,000, +6.0%); accepted by a reviewer: "
                                 "Field counts include drop-offs.")
    aware = frame[frame["kpi_id"] == "brand_awareness"].iloc[0]
    assert aware["status"] == "Report-only" and "not independently verified" in aware["caveat"]
    assert scorecard_frame(store, registry, "NOBODY").empty


def test_scorecard_csv_round_trips(published, store, registry):
    data = scorecard_csv(store, registry, "KESTREL")
    back = pd.read_csv(io.BytesIO(data), encoding="utf-8-sig")
    assert list(back.columns) == list(SCORECARD_COLUMNS) and len(back) == 4
    assert back.loc[back["kpi_id"] == "trx", "value"].iloc[0] == 5910.0


def test_scorecard_xlsx_is_formatted(published, store, registry):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(scorecard_xlsx(store, registry, "KESTREL")))
    assert wb.sheetnames == ["Scorecard", "Notes"]
    ws = wb["Scorecard"]
    header = [c.value for c in ws[1]]
    assert header[:4] == ["KPI", "Scope", "Period", "Value"] and ws["A1"].font.bold
    assert ws.max_row == 5 and ws.freeze_panes == "B2"
    value_col = header.index("Value") + 1
    rows = {ws.cell(row=r, column=header.index("KPI id") + 1).value: r for r in range(2, ws.max_row + 1)}
    assert ws.cell(row=rows["trx_share"], column=value_col).number_format == "0.0%"
    assert ws.cell(row=rows["trx_share"], column=value_col).value == pytest.approx(0.092)
    assert ws.cell(row=rows["trx"], column=value_col).number_format == "#,##0"
    variance = ws.cell(row=rows["samples"], column=header.index("Difference vs platform") + 1)
    assert variance.value == pytest.approx(0.06) and "%" in variance.number_format
    assert wb["Notes"]["B1"].value == "KESTREL"


def test_scorecard_pptx_has_title_and_table_slides(seeder, store, registry, settings):
    deck = seeder.doc("deck.pptx")
    for i in range(13):
        end = month_end(add_months(date(2025, 5, 1), i))
        s = seeder.seed("trx", S.VERIFIED, doc=deck, period_end=end, report_value=5000.0 + i,
                        platform_value=5000.0 + i)
        decide(store, registry, settings, s.result.result_id, "ACCEPT", "", "rev@example.com")
    from pptx import Presentation

    prs = Presentation(io.BytesIO(scorecard_pptx(store, registry, "KESTREL")))
    slides = list(prs.slides)
    assert len(slides) == 3
    assert slides[0].shapes.title.text == "Kestrel KPI scorecard"
    tables = [sh.table for sh in slides[1].shapes if sh.has_table]
    assert len(tables) == 1 and len(tables[0].rows) == 13
    assert tables[0].cell(0, 0).text == "KPI" and tables[0].cell(1, 4).text == "Matches platform"
    last = [sh.table for sh in slides[2].shapes if sh.has_table][0]
    assert len(last.rows) == 2 and slides[2].shapes.title.text == "Kestrel KPIs (2 of 2)"


def test_empty_exports_are_valid_files(store, registry):
    from openpyxl import load_workbook
    from pptx import Presentation

    prs = Presentation(io.BytesIO(scorecard_pptx(store, registry, "KESTREL")))
    assert len(prs.slides) == 1 and "No approved values" in prs.slides[0].placeholders[1].text
    ws = load_workbook(io.BytesIO(scorecard_xlsx(store, registry, "KESTREL")))["Scorecard"]
    assert ws.max_row == 1
    assert pd.read_csv(io.BytesIO(scorecard_csv(store, registry, "KESTREL")), encoding="utf-8-sig").empty


def test_locator_text_variants():
    assert locator_text("r.pdf", {"page": 3, "row": 2, "col": 4}) == "r.pdf, page 3, row 2, column 4"
    assert locator_text("d.pptx", '{"page": 5}', {"table_title": "KPIs", "row_label": "NRx", "col_header": "May"}) == \
        "d.pptx, slide 5, table 'KPIs', row 'NRx', column 'May'"
    assert locator_text(None, None) == ""
